import io
import asyncio
import copy
from contextlib import nullcontext
from email.message import EmailMessage
import hashlib
import json
import logging
import os
import re
import secrets
import smtplib
import ssl
import threading
import time
import uuid
from collections import OrderedDict
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError
from urllib.parse import quote, urlsplit

from fastapi.encoders import jsonable_encoder
from fastapi.middleware.gzip import GZipMiddleware
from dotenv import load_dotenv, set_key

# Load local configuration before importing app modules. app.database reads
# DATABASE_URL_* at import time, so loading dotenv below those imports left it
# using stale process-environment values (or the local fallback).
load_dotenv()

from fastapi import BackgroundTasks, FastAPI, HTTPException, Request
from fastapi.responses import (
    FileResponse,
    HTMLResponse,
    JSONResponse,
    PlainTextResponse,
    RedirectResponse,
    StreamingResponse,
)
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
from starlette.templating import Jinja2Templates
from starlette.concurrency import run_in_threadpool

from authlib.integrations.base_client import OAuthError
from authlib.integrations.starlette_client import OAuth

from werkzeug.security import generate_password_hash, check_password_hash
from app.database import (
    ODDI_LOCAL_ONLY_STORAGE,
    ODDI_BROWSER_LOCAL_CHATS,
    ODDI_CHAT_FILES_STORAGE,
    ODDI_DRIVE_CHAT_STORAGE,
    ODDI_RENDER_LAPTOP_REDIRECT,
    MEMORY_EXTRACTION_BATCH_SIZE,
    create_tables,
    create_user,
    get_user_by_email,
    create_conversation,
    import_local_conversation,
    get_conversations,
    get_conversation_summaries,
    get_conversation,
    update_conversation,
    delete_conversation,
    delete_all_conversations,
    get_deleted_conversations,
    get_deleted_conversation_summaries,
    update_conversation_metadata,
    append_conversation_message,
    update_conversation_message,
    delete_conversation_message,
    get_memory,
    save_memory,
    delete_memory,
    register_files,
    count_files,
    FileLimitExceeded,
    USER_FILE_LIMIT,
    purge_expired_deleted_conversations,
    purge_expired_conversation_history,
    update_file_record,
    get_files,
    get_file,
    delete_file_record,
    delete_user_record,
    get_user_settings,
    get_user_by_id,
    list_account_user_ids,
    save_user_settings,
    store_password_reset_token,
    get_password_reset_email,
    consume_password_reset_token,
    clear_password_reset_token,
    update_user_password,
    update_user_profile,
    clear_memory,
    get_vector_store_id,
    enqueue_memory_message,
    get_memory_extraction_status,
    clear_memory_extraction_queue,
    recover_stale_memory_claims,
    create_interview_session,
    list_interview_sessions,
    get_interview_session,
    update_interview_session,
    create_skills_learning_session,
    list_skills_learning_sessions,
    get_skills_learning_session,
    update_skills_learning_session,
)
from app.google_drive_storage import (
    DriveStorageError,
    backfill_conversation_summaries,
    user_id_for_email as drive_user_id_for_email,
)

from app.engine import process_message, PROVIDER_UNAVAILABLE_RESPONSE
from app.shared_workspaces import (
    cleanup_deleted_account as cleanup_shared_workspaces_for_account,
    router as shared_workspaces_router,
    shared_workspace_queue_worker,
)
from app.interview_studio import (
    InterviewStudioError,
    generate_alternate_question,
    generate_interview_answer,
    generate_interview_answer_set,
    generate_question_set,
    review_interview_answers,
)
from app.help_feedback import (
    FeedbackCooldownError,
    FeedbackValidationError,
    admin_feedback_overview,
    admin_feedback_reply,
    is_admin_email as is_help_feedback_admin,
    mark_replies_read as mark_help_feedback_replies_read,
    submit_feedback as save_help_feedback_submission,
    unread_feedback_replies as get_unread_help_feedback_replies,
    user_feedback_state as get_help_feedback_state,
)
from app.attachments import AttachmentReadError, MAX_UPLOAD_BYTES, read_attachments
from app.config import client
from app.fast_responses import fast_response
from identity.identity import create_identity
from app.memory_ai import (
    is_safe_memory_message,
    memory_provider_keys_configured,
    memory_user_disabled,
    memory_user_lock,
    memory_user_opted_out,
    set_memory_user_disabled,
    set_memory_user_opt_out,
    normalize_memory_settings,
    process_one_memory_batch,
)
from app.swarm import (
    AGENT_REGISTRY as SWARM_AGENT_REGISTRY,
    MAX_AGENT_CONCURRENCY as SWARM_MAX_AGENT_CONCURRENCY,
    MAX_AGENT_COUNT as SWARM_MAX_AGENT_COUNT,
    SPECIALISTS as SWARM_SPECIALISTS,
    SwarmFile,
    auto_route as swarm_auto_route,
    create_execution_plan as create_swarm_execution_plan,
    get_agent_registry as get_swarm_agent_registry,
    run_task as run_swarm_task,
    validate_agent_ids as validate_swarm_agent_ids,
)
from app.storage import (
    storage_router,
    StorageQuotaExceeded,
    FileNotFound as StorageFileNotFound,
)


logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("oddi.persistence")

# A browser Retry action resends the same generation ID and attachments. Keep
# a bounded in-memory idempotency cache so those exact files are not stored in
# the user's Library twice when the original request reached storage but its
# response failed on the way back to the browser.
_chat_upload_retry_lock = threading.Lock()
_chat_upload_retry_cache = OrderedDict()
_CHAT_UPLOAD_RETRY_CACHE_LIMIT = 512


def _uploaded_batch_fingerprint(file_payloads):
    digest = hashlib.sha256()
    for uploaded_file, file_bytes in file_payloads:
        for value in (
            uploaded_file.filename or "unnamed-file",
            uploaded_file.mimetype or "application/octet-stream",
            str(len(file_bytes)),
            hashlib.sha256(file_bytes).hexdigest(),
        ):
            digest.update(value.encode("utf-8", errors="replace"))
            digest.update(b"\0")
    return digest.hexdigest()


def _get_chat_upload_retry_fingerprint(user_id, generation_id):
    key = (str(user_id), str(generation_id))
    with _chat_upload_retry_lock:
        fingerprint = _chat_upload_retry_cache.get(key)
        if fingerprint is not None:
            _chat_upload_retry_cache.move_to_end(key)
        return fingerprint


def _remember_chat_upload_retry_fingerprint(user_id, generation_id, fingerprint):
    key = (str(user_id), str(generation_id))
    with _chat_upload_retry_lock:
        _chat_upload_retry_cache[key] = fingerprint
        _chat_upload_retry_cache.move_to_end(key)
        while len(_chat_upload_retry_cache) > _CHAT_UPLOAD_RETRY_CACHE_LIMIT:
            _chat_upload_retry_cache.popitem(last=False)


def _library_filename_key(filename):
    value = str(filename or "unnamed-file").replace("\\", "/").rsplit("/", 1)[-1]
    return value.strip().casefold() or "unnamed-file"


def _unique_library_filename(filename, occupied_names):
    """Keep older same-named Library files by giving each new version a label."""
    filename = str(filename or "unnamed-file")
    if _library_filename_key(filename) not in occupied_names:
        return filename

    stem, extension = os.path.splitext(filename)
    stem = stem or filename
    version = 2
    while True:
        candidate = f"{stem} ({version}){extension}"
        if _library_filename_key(candidate) not in occupied_names:
            return candidate
        version += 1


def _read_library_file_bytes(user_id, record):
    storage_key = record.get("storage_key")
    if storage_key:
        try:
            return storage_router.read_file(user_id, storage_key)
        except Exception:
            pass

    try:
        full_record = get_file(record.get("id"), user_id)
    except Exception:
        return None
    if not full_record:
        return None
    data = full_record.get("file_data")
    if isinstance(data, memoryview):
        data = data.tobytes()
    return data if isinstance(data, bytes) else None


def _library_record_content_details(user_id, record):
    cached_hash = str(record.get("content_hash") or "").strip().lower()
    if len(cached_hash) == 64:
        try:
            int(cached_hash, 16)
        except ValueError:
            cached_hash = ""

    storage_key = record.get("storage_key")
    if storage_key:
        try:
            path = storage_router.get_file_path(user_id, storage_key)
            physical_size = int(path.stat().st_size)
            if cached_hash:
                return cached_hash, physical_size

            digest = hashlib.sha256()
            with open(path, "rb") as stored_file:
                for chunk in iter(lambda: stored_file.read(1024 * 1024), b""):
                    digest.update(chunk)
            return digest.hexdigest(), physical_size
        except Exception:
            # A hash is only meaningful while its backing blob still exists.
            cached_hash = ""

    if cached_hash:
        return cached_hash, 0
    data = _read_library_file_bytes(user_id, record)
    if data is None:
        return None, 0
    return hashlib.sha256(data).hexdigest(), 0


def _plan_chat_library_uploads(user_id, file_payloads):
    """Plan content deduplication while preserving distinct same-named files."""
    library_records = get_files(user_id, include_content_hash=True)
    library_hashes = set()
    library_file_ids_by_hash = {}
    occupied_names = set()

    for record in library_records:
        digest, _ = _library_record_content_details(user_id, record)
        if digest:
            library_hashes.add(digest)
            library_file_ids_by_hash[digest] = record.get("id")
            if str(record.get("content_hash") or "").lower() != digest:
                try:
                    update_file_record(record.get("id"), user_id, content_hash=digest)
                except Exception:
                    logger.debug("Could not cache content hash for Library file %s", record.get("id"))
        occupied_names.add(_library_filename_key(record.get("filename")))

    planned_uploads = []
    already_exists_count = 0
    already_existing_files = []
    duplicate_uploads = []
    planned_upload_index_by_hash = {}
    seen_hashes = set(library_hashes)
    for upload_index, (uploaded_file, file_bytes) in enumerate(file_payloads):
        original_filename = uploaded_file.filename or "unnamed-file"
        digest = hashlib.sha256(file_bytes).hexdigest()
        if digest in seen_hashes:
            already_exists_count += 1
            if digest in library_hashes:
                existing_file_id = library_file_ids_by_hash.get(digest)
                if existing_file_id is not None:
                    already_existing_files.append({
                        "upload_index": upload_index,
                        "file_id": existing_file_id,
                    })
            else:
                planned_upload_index = planned_upload_index_by_hash.get(digest)
                if planned_upload_index is not None:
                    duplicate_uploads.append({
                        "upload_index": upload_index,
                        "planned_upload_index": planned_upload_index,
                    })
            continue

        seen_hashes.add(digest)
        filename = _unique_library_filename(original_filename, occupied_names)
        occupied_names.add(_library_filename_key(filename))
        planned_upload_index_by_hash[digest] = upload_index
        planned_uploads.append({
            "uploaded_file": uploaded_file,
            "file_bytes": file_bytes,
            "filename": filename,
            "digest": digest,
            "upload_index": upload_index,
        })

    return {
        "uploads": planned_uploads,
        "replace_file_ids": [],
        "replace_storage_keys": [],
        "quota_credit_bytes": 0,
        "already_exists_count": already_exists_count,
        "already_existing_files": already_existing_files,
        "duplicate_uploads": duplicate_uploads,
    }


def _persist_chat_library_uploads(user_id, conversation_id, plan):
    uploads = plan.get("uploads") or []
    if not uploads:
        return []

    stored_keys = []
    records = []
    try:
        for entry in uploads:
            uploaded_file = entry["uploaded_file"]
            file_bytes = entry["file_bytes"]
            stored = storage_router.save_file(
                user_id=user_id,
                data=file_bytes,
                original_filename=entry["filename"],
                quota_credit_bytes=plan.get("quota_credit_bytes", 0),
            )
            storage_key = stored["storage_key"]
            stored_keys.append(storage_key)
            records.append({
                "filename": entry["filename"],
                "mime_type": uploaded_file.mimetype,
                "size_bytes": len(file_bytes),
                "conversation_id": conversation_id,
                "storage_backend": "filesystem",
                "storage_key": storage_key,
                "external_file_id": None,
                "file_data": None,
                "content_hash": entry["digest"],
                "status": "processed",
            })

        new_record_ids = register_files(
            user_id,
            records,
            replace_file_ids=plan.get("replace_file_ids"),
        )
    except Exception:
        rollback_stored_files(user_id, stored_keys)
        raise

    # Distinct uploads are saved as separate records, even when their
    # original filenames match an older Library file.
    return new_record_ids

_chat_generation_lock = threading.Lock()
_active_chat_generations = {}


# =========================================================
# FASTAPI APPLICATION
# =========================================================

app = FastAPI(
    title="ODDI AI",
    version="1.0.0",
)
app.include_router(shared_workspaces_router)

_SUMMARY_BACKFILL_LOCK = threading.Lock()
_SUMMARY_BACKFILL_USERS = set()
_WORKSPACE_MUTATION_LOCK = threading.RLock()
_USER_SETTINGS_CACHE_LOCK = threading.RLock()
_USER_SETTINGS_CACHE = {}
_USER_SETTINGS_CACHE_TTL_SECONDS = 30
_LAPTOP_DRIVE_SYNC_STATE_LOCK = threading.RLock()
app.state.laptop_drive_sync = {
    "state": "idle",
    "last_attempt_at": None,
    "last_success_at": None,
    "last_error_code": None,
    "last_result": {},
}


def _drive_sync_timestamp():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _cache_user_settings(user_id, settings):
    with _USER_SETTINGS_CACHE_LOCK:
        _USER_SETTINGS_CACHE[str(user_id)] = (time.monotonic(), copy.deepcopy(settings or {}))


def _cached_user_settings(user_id):
    key = str(user_id)
    with _USER_SETTINGS_CACHE_LOCK:
        cached = _USER_SETTINGS_CACHE.get(key)
        if cached and time.monotonic() - cached[0] < _USER_SETTINGS_CACHE_TTL_SECONDS:
            return copy.deepcopy(cached[1])
    settings = get_user_settings(user_id)
    _cache_user_settings(user_id, settings)
    return settings


def _run_laptop_drive_sync():
    """Run one laptop mirror pass and keep a safe, user-visible status snapshot."""
    from app import drive_sync

    attempt_at = _drive_sync_timestamp()
    with _LAPTOP_DRIVE_SYNC_STATE_LOCK:
        previous_status = dict(app.state.laptop_drive_sync)
        status = dict(previous_status)
        status.update({"state": "syncing", "last_attempt_at": attempt_at, "last_error_code": None})
        app.state.laptop_drive_sync = status

    try:
        result = drive_sync.sync_laptop_once()
    except Exception as error:
        message = str(error).casefold()
        auth_rejected = any(token in message for token in ("refresh token", "invalid_grant", "oauth", "authorization rejected"))
        with _LAPTOP_DRIVE_SYNC_STATE_LOCK:
            status = dict(app.state.laptop_drive_sync)
            status.update({
                "state": "authorization_required" if auth_rejected else "retrying",
                "last_error_code": "authorization_required" if auth_rejected else "sync_failed",
                "last_error_at": _drive_sync_timestamp(),
            })
            app.state.laptop_drive_sync = status
        raise

    if result.get("skipped"):
        with _LAPTOP_DRIVE_SYNC_STATE_LOCK:
            current = dict(app.state.laptop_drive_sync)
            if current.get("last_attempt_at") == attempt_at:
                app.state.laptop_drive_sync = previous_status
    else:
        with _LAPTOP_DRIVE_SYNC_STATE_LOCK:
            status = dict(app.state.laptop_drive_sync)
            status.update({
                "state": "synced",
                "last_success_at": _drive_sync_timestamp(),
                "last_error_code": None,
                "last_result": {
                    key: int(value or 0)
                    for key, value in result.items()
                    if key in {"accounts", "chats", "uploaded_chats", "downloaded_chats", "memories"}
                },
            })
            status.pop("last_error_at", None)
            app.state.laptop_drive_sync = status
    return result


def _storage_sync_status_payload(user_id):
    from app import google_drive_storage, local_chat_files

    if ODDI_LOCAL_ONLY_STORAGE:
        pending = local_chat_files.get_sync_status(user_id)
        local_chats = len(local_chat_files.get_conversations(user_id))
        configured = google_drive_storage.is_configured()
        with _LAPTOP_DRIVE_SYNC_STATE_LOCK:
            runtime = dict(app.state.laptop_drive_sync)
        if not configured:
            state = "not_configured"
        elif runtime.get("state") in {"authorization_required", "retrying", "syncing"}:
            state = runtime["state"]
        elif pending["pending_chat_count"] or pending["pending_delete_count"]:
            state = "pending"
        elif runtime.get("state") == "idle":
            state = "checking"
        elif pending["initialized"]:
            state = "synced"
        else:
            state = "checking"
        return {
            "storage_mode": "laptop",
            "state": state,
            "local_chat_count": local_chats,
            "pending_known": pending["initialized"],
            "pending_chat_count": pending["pending_chat_count"],
            "pending_delete_count": pending["pending_delete_count"],
            "last_sync_at": pending["last_sync_at"],
            "last_sync_result": pending["last_sync_result"],
            "last_attempt_at": runtime.get("last_attempt_at"),
            "error_code": runtime.get("last_error_code"),
            "drive_configured": configured,
            "can_reconnect": bool(configured and oauth.create_client("google")),
        }

    chat_count = len(get_conversations(user_id))
    if ODDI_DRIVE_CHAT_STORAGE:
        return {
            "storage_mode": "drive",
            "state": "cloud",
            "local_chat_count": chat_count,
            "pending_known": True,
            "pending_chat_count": 0,
            "pending_delete_count": 0,
            "last_sync_at": None,
            "last_sync_result": {},
            "drive_configured": True,
            "can_reconnect": False,
        }
    return {
        "storage_mode": "server",
        "state": "server_storage",
        "local_chat_count": chat_count,
        "pending_known": False,
        "pending_chat_count": 0,
        "pending_delete_count": 0,
        "last_sync_at": None,
        "last_sync_result": {},
        "drive_configured": False,
        "can_reconnect": False,
    }


def _backfill_chat_summaries(user_id):
    try:
        backfill_conversation_summaries(user_id, 40)
    except Exception:
        logger.exception("Could not update legacy chat summary metadata for account %s", user_id)
    finally:
        with _SUMMARY_BACKFILL_LOCK:
            _SUMMARY_BACKFILL_USERS.discard(str(user_id))


def _enqueue_memory_message_if_enabled(user_id, event_id, message):
    """Queue a quick-reply turn after sending only if account consent allows it."""
    if (
        ODDI_BROWSER_LOCAL_CHATS
        or not memory_provider_keys_configured()
        or memory_user_disabled(user_id)
        or memory_user_opted_out(user_id)
    ):
        return
    try:
        settings = normalize_memory_settings(get_user_settings(user_id).get("memory"))
        if settings["enabled"] and settings["auto_extract"]:
            safe_message = message if is_safe_memory_message(message) else ""
            enqueue_memory_message(str(user_id), event_id, safe_message)
    except Exception:
        logger.exception("Could not queue a quick chat turn for memory extraction")
app.add_middleware(GZipMiddleware, minimum_size=1024)


@app.middleware("http")
async def identify_unhandled_chat_errors(request: Request, call_next):
    """Give unexpected /chat failures a support reference and a traceback log."""
    if request.url.path != "/chat":
        return await call_next(request)
    request_id = uuid.uuid4().hex[:12]
    request.state.chat_request_id = request_id
    try:
        response = await call_next(request)
    except Exception:
        logger.exception("Unhandled /chat request failure request_id=%s", request_id)
        response = JSONResponse(
            {
                "error": "The chat server could not complete this request. Please retry.",
                "request_id": request_id,
            },
            status_code=500,
        )
    response.headers["X-Oddi-Request-ID"] = request_id
    return response


@app.middleware("http")
async def require_drive_storage_configuration(request: Request, call_next):
    """Do not accept public accounts or chats before persistent Drive is ready."""
    if ODDI_RENDER_LAPTOP_REDIRECT:
        if request.url.path == "/healthz":
            return JSONResponse({"status": "drive-storage-not-configured"})
        return HTMLResponse(
            """<!doctype html>
<html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>ODDI storage is not configured</title>
<style>
*{box-sizing:border-box}body{margin:0;min-height:100vh;display:grid;place-items:center;padding:24px;background:#090a0a;color:#f4f4f4;font:16px/1.6 system-ui,-apple-system,Segoe UI,sans-serif}
main{width:min(520px,100%);padding:32px;border:1px solid #292b2b;border-radius:20px;background:#111313;box-shadow:0 24px 80px #0008}h1{margin:0 0 12px;font-size:24px;letter-spacing:-.03em}p{margin:10px 0;color:#b8bcbc}.badge{display:inline-flex;align-items:center;gap:8px;margin-bottom:20px;padding:6px 10px;border:1px solid #6c4d20;border-radius:999px;color:#ffd38a;font-size:12px}.dot{width:8px;height:8px;border-radius:50%;background:#efad4d}small{display:block;margin-top:24px;color:#858a8a}
</style>
<main><div class="badge"><span class="dot"></span> Drive storage is not connected</div>
<h1>ODDI needs its shared chat storage</h1>
<p>The public site uses Google Drive for account records and chat history while the laptop is offline.</p>
<p>No account or chat was saved on Render. Configure Drive access in Render's Environment settings, then redeploy.</p>
<small>Set <code>GOOGLE_CLIENT_ID</code>, <code>GOOGLE_CLIENT_SECRET</code>, <code>GOOGLE_REFRESH_TOKEN</code>, and <code>GOOGLE_DRIVE_FOLDER_ID</code> in the laptop and Render environments.</small>
</main></html>""",
            status_code=503,
            headers={"Cache-Control": "no-store", "Retry-After": "60"},
        )

    if ODDI_BROWSER_LOCAL_CHATS and request.url.path.startswith("/api/conversations"):
        if request.method == "GET":
            return JSONResponse([])
        return JSONResponse(
            {"error": "Chat history is saved in this browser on your device."},
            status_code=410,
        )
    return await call_next(request)


@app.on_event("startup")
async def start_bin_retention_cleanup():
    if ODDI_RENDER_LAPTOP_REDIRECT:
        return

    if not ODDI_BROWSER_LOCAL_CHATS:
        app.state.scheduled_tasks_task = asyncio.create_task(_scheduled_task_worker())

    async def cleanup_expired_bin_chats():
        # Do not compete with the first account/history requests when Render
        # starts. Drive-backed cleanup is user-scoped when the Bin is opened;
        # the all-account sweep only needs to run once per day.
        if ODDI_DRIVE_CHAT_STORAGE:
            await asyncio.sleep(5 * 60)
            cleanup_interval = 24 * 60 * 60
        else:
            cleanup_interval = 60 * 60
        while True:
            try:
                await run_in_threadpool(purge_expired_deleted_conversations)
            except Exception:
                logger.exception("Automatic Bin retention cleanup failed")
            await asyncio.sleep(cleanup_interval)

    app.state.bin_retention_task = asyncio.create_task(cleanup_expired_bin_chats())

    async def cleanup_expired_conversation_history():
        # Expired active/archived chats are hard-deleted. Keep this sweep
        # separate from the existing seven-day Bin cleanup and avoid loading
        # Drive account data during application startup.
        await asyncio.sleep(5 * 60 if ODDI_DRIVE_CHAT_STORAGE else 60)
        cleanup_interval = 24 * 60 * 60 if ODDI_DRIVE_CHAT_STORAGE else 60 * 60
        while True:
            try:
                removed = await run_in_threadpool(purge_expired_conversation_history)
                if removed:
                    logger.info("Conversation retention removed %s expired chats", removed)
            except Exception:
                logger.exception("Automatic conversation retention cleanup failed")
            await asyncio.sleep(cleanup_interval)

    app.state.conversation_retention_task = asyncio.create_task(cleanup_expired_conversation_history())

    if not ODDI_DRIVE_CHAT_STORAGE or os.getenv("ODDI_SHARED_WORKSPACES_DATABASE_URL", "").strip():
        app.state.shared_workspace_task = asyncio.create_task(shared_workspace_queue_worker())

    if not ODDI_BROWSER_LOCAL_CHATS and memory_provider_keys_configured():
        await run_in_threadpool(recover_stale_memory_claims)

        async def extract_queued_memories():
            # Let the app finish its first account and conversation requests.
            await asyncio.sleep(20)
            while True:
                try:
                    result = await run_in_threadpool(process_one_memory_batch)
                except Exception:
                    logger.exception("Memory extraction worker hit a storage error; it will retry")
                    result = "retry"
                await asyncio.sleep(3 if result == "processed" else 15)

        app.state.memory_extraction_task = asyncio.create_task(extract_queued_memories())

    if ODDI_LOCAL_ONLY_STORAGE:
        from app import google_drive_storage
        if google_drive_storage.is_configured():
            async def sync_laptop_chats():
                retry_delay = 30
                auth_failure_reported = False
                while True:
                    try:
                        await run_in_threadpool(_run_laptop_drive_sync)
                    except google_drive_storage.DriveStorageError as error:
                        message = str(error).casefold()
                        if "refresh token" in message or "oauth" in message:
                            if not auth_failure_reported:
                                logger.error(
                                    "Google Drive sync is paused because its authorization was rejected. "
                                    "Local chats remain saved on this laptop; reconnect at /auth/google/drive."
                                )
                                auth_failure_reported = True
                            await asyncio.sleep(retry_delay)
                            retry_delay = min(retry_delay * 2, 300)
                            continue
                        if "temporarily unreachable" in message or "temporarily unavailable" in message:
                            logger.warning(
                                "Laptop to Google Drive sync is waiting for network access; "
                                "local chats remain saved. Retrying in %s seconds.",
                                retry_delay,
                            )
                        else:
                            logger.error("Laptop to Google Drive chat sync failed: %s", error)
                    except Exception:
                        logger.exception("Laptop to Google Drive chat sync failed; it will retry")
                        await asyncio.sleep(retry_delay)
                        retry_delay = min(retry_delay * 2, 300)
                        continue
                    else:
                        auth_failure_reported = False
                        retry_delay = 30
                        await asyncio.sleep(30)
                        continue
                    await asyncio.sleep(retry_delay)
                    retry_delay = min(retry_delay * 2, 300)

            app.state.drive_sync_task = asyncio.create_task(sync_laptop_chats())


@app.on_event("shutdown")
async def stop_bin_retention_cleanup():
    for task_name in ("bin_retention_task", "conversation_retention_task", "drive_sync_task", "memory_extraction_task", "shared_workspace_task", "scheduled_tasks_task"):
        task = getattr(app.state, task_name, None)
        if task:
            task.cancel()
            try:
                await task
            except asyncio.CancelledError:
                pass

# Flask's signed session cookie is replaced by Starlette's signed session
# middleware. The 30-day lifetime and the important cookie protections are
# preserved from the old Flask server.
_IS_RENDER_SERVER = (
    os.getenv("RENDER", "").strip().lower() in {"1", "true", "yes", "on"}
    or bool(os.getenv("RENDER_SERVICE_ID", "").strip())
)
_IS_PRODUCTION_SERVER = (
    os.getenv("ENVIRONMENT", "").strip().lower() in {"production", "prod"}
    or _IS_RENDER_SERVER
)
_SECURE_SESSION_COOKIES = (
    _IS_PRODUCTION_SERVER
    or os.getenv("ODDI_SECURE_COOKIES", "").strip().lower() in {"1", "true", "yes", "on"}
)
_SESSION_SECRET = (
    os.getenv("FLASK_SECRET_KEY", "").strip()
    or os.getenv("ODDI_SESSION_SECRET", "").strip()
    or os.getenv("SECRET_KEY", "").strip()
)
if _SECURE_SESSION_COOKIES and ODDI_RENDER_LAPTOP_REDIRECT and len(_SESSION_SECRET) < 32:
    # Render never handles an account session in laptop-storage mode.
    _SESSION_SECRET = secrets.token_urlsafe(48)
elif _SECURE_SESSION_COOKIES and ODDI_BROWSER_LOCAL_CHATS and len(_SESSION_SECRET) < 32:
    # Browser-local mode has no persistent account sessions. Use a
    # cryptographically random process key when none is configured; the
    # temporary guest session can safely be recreated.
    _SESSION_SECRET = secrets.token_urlsafe(48)
elif _SECURE_SESSION_COOKIES and len(_SESSION_SECRET) < 32:
    # Render needs a stable signing key across restarts so account cookies keep
    # working. If the operator has not set a dedicated session key, derive one
    # from persistent storage credentials. Never log or expose this material.
    _database_secret_material = "\0".join(
        os.getenv(name, "").strip()
        for name in (
            "DATABASE_URL",
            "DATABASE_URL_CHAT",
            "DATABASE_URL_FILES",
            "DATABASE_URL_ARCHIVE_MEMORY",
        )
        if os.getenv(name, "").strip()
    )
    if _database_secret_material:
        _SESSION_SECRET = hashlib.sha256(
            b"oddi-session-cookie-signing-v1\0"
            + _database_secret_material.encode("utf-8")
        ).hexdigest()
    elif ODDI_DRIVE_CHAT_STORAGE:
        _drive_secret_material = "\0".join(
            os.getenv(name, "").strip()
            # Refresh tokens can be replaced during Drive reauthorization.
            # Keep existing sessions valid across that storage credential
            # rotation; the OAuth client secret remains the stable key source.
            for name in ("GOOGLE_CLIENT_ID", "GOOGLE_CLIENT_SECRET")
            if os.getenv(name, "").strip()
        )
        _SESSION_SECRET = hashlib.sha256(
            b"oddi-drive-session-cookie-signing-v1\0"
            + _drive_secret_material.encode("utf-8")
        ).hexdigest()
    else:
        if ODDI_CHAT_FILES_STORAGE:
            raise RuntimeError(
                "Set FLASK_SECRET_KEY to a persistent random value of at least 32 characters "
                "before exposing the laptop-hosted ODDI service."
            )
        raise RuntimeError(
            "Set FLASK_SECRET_KEY to a persistent random value of at least 32 characters, "
            "or configure a persistent PostgreSQL DATABASE_URL for production."
        )
if not _SESSION_SECRET:
    _SESSION_SECRET = "local-development-only-secret-change-before-deploy"

app.add_middleware(
    SessionMiddleware,
    secret_key=_SESSION_SECRET,
    max_age=int(timedelta(days=30).total_seconds()),
    same_site="lax",
    https_only=_SECURE_SESSION_COOKIES,
)


@app.exception_handler(DriveStorageError)
async def drive_storage_error_handler(request: Request, exc: DriveStorageError):
    message = str(exc)
    if request.url.path.startswith("/api/"):
        return JSONResponse({"success": False, "error": message}, status_code=503)
    safe_message = message.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return HTMLResponse(
        "<!doctype html><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
        "<title>ODDI storage setup</title>"
        "<style>body{margin:0;min-height:100vh;display:grid;place-items:center;background:#080909;color:#eee;"
        "font:16px/1.6 system-ui;padding:24px}main{max-width:620px;padding:28px;border:1px solid #333;"
        "border-radius:18px;background:#111}h1{font-size:23px}p{color:#bbb}</style>"
        f"<main><h1>ODDI storage needs setup</h1><p>{safe_message}</p></main>",
        status_code=503,
        headers={"Cache-Control": "no-store"},
    )

# Keep the existing static directory available under the same /static URL.
# This is important because the existing HTML references /static/... assets.
if os.path.isdir("static"):
    app.mount("/static", StaticFiles(directory="static"), name="static")

# Keep the existing Jinja templates.
templates = Jinja2Templates(directory="templates")


# =========================================================
# TEMPLATE URL COMPATIBILITY
# =========================================================

# Existing ODDI templates were written for Flask's url_for(), including the
# Flask-specific `filename=` argument for static files. This helper keeps the
# existing templates working during the framework migration instead of forcing
# a rewrite of index.html/login.html/signup.html right now.
def template_url_for(request: Request, endpoint: str, **values):
    if endpoint == "static":
        filename = values.pop("filename", values.pop("path", ""))
        return str(request.url_for("static", path=filename))

    return str(request.url_for(endpoint, **values))


def render_template(request: Request, template_name: str, **context):
    context["request"] = request
    context["url_for"] = lambda endpoint, **values: template_url_for(
        request,
        endpoint,
        **values,
    )
    return templates.TemplateResponse(
        request=request,
        name=template_name,
        context=context,
    )


# =========================================================
# GOOGLE OAUTH
# =========================================================

oauth = OAuth()

google_client_id = os.getenv("GOOGLE_CLIENT_ID")
google_client_secret = os.getenv("GOOGLE_CLIENT_SECRET")

if google_client_id and google_client_secret:
    oauth.register(
        name="google",
        server_metadata_url=(
            "https://accounts.google.com/"
            ".well-known/openid-configuration"
        ),
        client_id=google_client_id,
        client_secret=google_client_secret,
        client_kwargs={
            "scope": "openid profile email",
        },
    )
else:
    logger.warning(
        "Google OAuth is not configured. Set "
        "GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET in .env."
    )


# Initialize the existing database exactly as before. Bin retention cleanup is
# started in the application startup task below; keeping Drive I/O out of module
# import ensures a temporary credential/API problem cannot make Render crash
# before Uvicorn opens its port.
create_tables()


# =========================================================
# HELPERS
# =========================================================

async def safe_json(request: Request):
    """Flask-compatible equivalent of request.get_json(silent=True)."""
    try:
        data = await request.json()
    except Exception:
        return {}
    return data if isinstance(data, dict) else {}


async def sync_laptop_accounts_from_drive():
    """Load Drive-created accounts before local authentication/registration."""
    if not ODDI_LOCAL_ONLY_STORAGE:
        return
    from app import google_drive_storage
    if not google_drive_storage.is_configured():
        return
    try:
        await run_in_threadpool(_run_laptop_drive_sync)
    except Exception:
        logger.exception("Could not refresh laptop accounts from Google Drive")


def set_user_session(request: Request, user, auth_provider=None):
    """Store the same user identity fields used by the old Flask server."""
    request.session["user_id"] = user["id"]
    request.session["username"] = user["username"]
    request.session["email"] = user["email"]

    if auth_provider:
        request.session["auth_provider"] = auth_provider
    else:
        request.session.pop("auth_provider", None)


def _trusted_drive_session_user(request: Request):
    """Validate the signed account identity without a Drive round trip."""
    if not ODDI_DRIVE_CHAT_STORAGE:
        return None
    email = str(request.session.get("email") or "").strip().casefold()
    provider = str(request.session.get("auth_provider") or "").casefold()
    if not email or email.endswith("@local.oddi.invalid") or provider == "browser-local":
        return None
    try:
        session_id = int(request.session.get("user_id"))
        expected_id = drive_user_id_for_email(email)
    except (TypeError, ValueError, OverflowError):
        return None
    if session_id != expected_id:
        return None
    return {
        "id": session_id,
        "username": str(request.session.get("username") or email.split("@", 1)[0]),
        "email": email,
    }


def _get_session_user(request: Request):
    is_browser_local_identity = (
        request.session.get("auth_provider") == "browser-local"
        or str(request.session.get("email") or "").strip().casefold().endswith("@local.oddi.invalid")
    )
    if is_browser_local_identity and not ODDI_BROWSER_LOCAL_CHATS:
        return None
    user_id = request.session.get("user_id")
    if not user_id:
        return None

    # The session cookie is signed, and the Drive account ID is a stable hash
    # of its email. Rechecking Drive on every page/API request made regular
    # navigation depend on several remote calls and caused repeated sign-ins.
    trusted_drive_user = _trusted_drive_session_user(request)
    if trusted_drive_user:
        return trusted_drive_user

    user = get_user_by_id(user_id)
    session_email = str(request.session.get("email") or "").strip()
    if not session_email:
        return None
    user_email = str(user["email"] if user else "").strip()
    if user and user_email.casefold() == session_email.casefold():
        return user
    if not ODDI_LOCAL_ONLY_STORAGE:
        return None

    # A cloud session's numeric ID can collide with a different local account.
    # Match by email and rebind the cookie to the corresponding local ID.
    user = get_user_by_email(session_email)
    if user:
        set_user_session(
            request,
            user,
            auth_provider=request.session.get("auth_provider"),
        )
    return user


def _request_logging_enabled(request: Request):
    user_id = request.session.get("user_id")
    if not user_id:
        return True
    if ODDI_LOCAL_ONLY_STORAGE:
        user = _get_session_user(request)
        if not user:
            return False
        user_id = user["id"]
    try:
        privacy = get_user_settings(user_id).get("privacy", {})
        return privacy.get("request_logging", True) is not False
    except Exception:
        return True


def require_user_id(request: Request):
    user_id = request.session.get("user_id")
    if not user_id:
        raise HTTPException(status_code=401, detail="Not logged in.")

    trusted_drive_user = _trusted_drive_session_user(request)
    if trusted_drive_user:
        return trusted_drive_user["id"]

    try:
        user = _get_session_user(request)
    except DriveStorageError:
        # The session cookie is signed by this app and the Drive-backed account
        # ID is derived from its email. If Drive is temporarily unreachable,
        # allow an already-authenticated account to keep chatting; its pending
        # conversation remains browser-local until Drive recovers.
        email = str(request.session.get("email") or "").strip().casefold()
        provider = str(request.session.get("auth_provider") or "").casefold()
        try:
            session_id = int(user_id)
            expected_id = drive_user_id_for_email(email)
        except (TypeError, ValueError, OverflowError):
            session_id = None
            expected_id = None
        if (
            ODDI_DRIVE_CHAT_STORAGE
            and provider != "browser-local"
            and not email.endswith("@local.oddi.invalid")
            and session_id is not None
            and session_id == expected_id
        ):
            logger.warning("Google Drive account lookup failed; continuing with the signed-in session.")
            return session_id
        raise
    if not user:
        request.session.clear()
        raise HTTPException(status_code=401, detail="This account is no longer available. Sign in again.")
    return user["id"] if ODDI_LOCAL_ONLY_STORAGE else user_id


def is_temporary_browser_account(request: Request):
    """Return true for the throwaway browser identity that cannot sync across devices."""
    provider = str(request.session.get("auth_provider") or "").strip().casefold()
    email = str(request.session.get("email") or "").strip().casefold()
    return provider == "browser-local" or email.endswith("@local.oddi.invalid")


async def require_user_id_async(request: Request):
    """Resolve Drive-backed account records without blocking Uvicorn's loop."""
    return await run_in_threadpool(require_user_id, request)


async def request_logging_enabled_async(request: Request):
    return await run_in_threadpool(_request_logging_enabled, request)


def _uploaded_file_size(uploaded_file):
    """Return upload size without consuming the upload stream."""
    try:
        current = uploaded_file.stream.tell()
        uploaded_file.stream.seek(0, 2)
        size = uploaded_file.stream.tell()
        uploaded_file.stream.seek(current)
        return int(size or 0)
    except Exception:
        return int(getattr(uploaded_file, "content_length", 0) or 0)


class UploadedFileAdapter:
    """Small compatibility wrapper for code that previously received Flask FileStorage.

    FastAPI/Starlette uses UploadFile. ODDI's existing engine can continue to
    receive an object with the old attributes: filename, mimetype and stream.
    """

    def __init__(self, upload):
        self._upload = upload
        self.filename = upload.filename or "unnamed-file"
        self.mimetype = upload.content_type or "application/octet-stream"
        self.content_type = self.mimetype
        self.content_length = getattr(upload, "size", None)
        self.stream = upload.file

    def read(self, size=-1):
        return self.stream.read(size)

    def seek(self, offset, whence=0):
        return self.stream.seek(offset, whence)

    def tell(self):
        return self.stream.tell()

    def seekable(self):
        """Expose the underlying stream's seek support to zipfile/python-docx."""
        return self.stream.seekable()

    def readable(self):
        return self.stream.readable()

    def writable(self):
        return self.stream.writable()

    def fileno(self):
        return self.stream.fileno()

    def close(self):
        return self.stream.close()


def adapt_uploaded_files(uploaded_files):
    return [UploadedFileAdapter(upload) for upload in uploaded_files]


def rollback_stored_files(user_id, storage_keys, file_record_ids=None):
    """Best-effort rollback for a failed upload/chat transaction.

    Physical storage and metadata are kept in sync: if processing fails after
    the files were stored, remove both the filesystem objects and their DB
    records. This prevents orphaned files from consuming the user's quota.
    """
    file_record_ids = list(file_record_ids or [])

    for storage_key in list(storage_keys or []):
        try:
            storage_router.delete_file(user_id, storage_key)
        except StorageFileNotFound:
            pass
        except Exception as error:
            logger.exception(
                "Failed to roll back physical file %s for user %s: %s",
                storage_key,
                user_id,
                error,
            )

    for file_record_id in file_record_ids:
        try:
            delete_file_record(file_record_id, user_id)
        except Exception as error:
            logger.exception(
                "Failed to roll back file metadata %s for user %s: %s",
                file_record_id,
                user_id,
                error,
            )


def response_from_engine(reply):
    """Convert common engine return values into FastAPI responses.

    The old Flask endpoint returned `reply` directly. This compatibility layer
    lets existing engine code return text, dict/list JSON data, or an already
    constructed Starlette/FastAPI response without forcing an engine rewrite.
    """
    if isinstance(reply, (HTMLResponse, JSONResponse, PlainTextResponse, StreamingResponse, RedirectResponse)):
        return reply

    if isinstance(reply, (dict, list, tuple, int, float, bool)) or reply is None:
        return JSONResponse(content=reply)

    if isinstance(reply, str):
        return PlainTextResponse(reply)

    # Last-resort compatibility for JSON-serializable objects.
    try:
        return JSONResponse(content=reply)
    except Exception:
        return PlainTextResponse(str(reply))


# =========================================================
# PAGES / BASIC AUTH
# =========================================================

@app.get("/", response_class=HTMLResponse, name="home")
def home(request: Request):
    # Browser-local mode uses a throwaway account solely for request
    # authorization; history remains in that browser.
    if ODDI_BROWSER_LOCAL_CHATS and not _get_session_user(request):
        local_email = f"browser-{uuid.uuid4().hex}@local.oddi.invalid"
        create_user(
            "Local user",
            local_email,
            generate_password_hash(secrets.token_urlsafe(32)),
        )
        local_user = get_user_by_email(local_email)
        if local_user:
            set_user_session(request, local_user, auth_provider="browser-local")

    # ODDI uses real persisted accounts in sync mode. Discard anonymous
    # browser-local cookies before they can be mistaken for a real account.
    if not request.session.get("user_id") or not _get_session_user(request):
        request.session.clear()
        return RedirectResponse(url="/login", status_code=303)

    response = render_template(
        request,
        "index.html",
        logged_in=True,
        username=request.session.get("username"),
        user_id=request.session.get("user_id"),
        browser_local_chats=ODDI_BROWSER_LOCAL_CHATS,
        library_sync_enabled=not is_temporary_browser_account(request),
    )

    # Preserve Flask's no-cache behavior for the live conversation-sync page.
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


@app.get("/interview-studio", response_class=HTMLResponse, name="interview_studio")
async def interview_studio(request: Request):
    try:
        await require_user_id_async(request)
    except HTTPException as error:
        if error.status_code == 401:
            return RedirectResponse(url="/login", status_code=303)
        raise
    response = render_template(
        request,
        "interview_studio.html",
        interview_user_key=str(request.session.get("user_id") or "guest"),
    )
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    return response


@app.get("/skills-studio", response_class=HTMLResponse, name="skills_studio")
async def skills_studio(request: Request):
    try:
        await require_user_id_async(request)
    except HTTPException as error:
        if error.status_code == 401:
            return RedirectResponse(url="/login", status_code=303)
        raise
    response = render_template(
        request,
        "skills_studio.html",
        skills_user_key=str(request.session.get("user_id") or "guest"),
        skills_browser_local_chats=ODDI_BROWSER_LOCAL_CHATS,
    )
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    return response


@app.get("/swarm-intelligence", response_class=HTMLResponse, name="swarm_intelligence")
async def swarm_intelligence(request: Request):
    try:
        user_id = await require_user_id_async(request)
    except HTTPException as error:
        if error.status_code == 401:
            return RedirectResponse(url="/login", status_code=303)
        raise
    return _render_swarm_page(request, user_id, "home")


@app.get("/swarm-intelligence/chat", response_class=HTMLResponse, name="swarm_chat_workspace")
async def swarm_chat_workspace(request: Request):
    return await _swarm_page_or_login(request, "chat")


def _render_swarm_page(request: Request, user_id, page_mode="home", agent_id=""):
    response = render_template(
        request,
        "swarm_intelligence.html",
        swarm_user_key=str(user_id),
        swarm_user_id=str(user_id),
        swarm_username=str(request.session.get("username") or "ODDI User"),
        swarm_email=str(request.session.get("email") or ""),
        swarm_browser_local_chats=ODDI_BROWSER_LOCAL_CHATS,
        swarm_page_mode=page_mode,
        swarm_agent_id=agent_id,
    )
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    return response


async def _swarm_page_or_login(request: Request, page_mode="home", agent_id=""):
    try:
        user_id = await require_user_id_async(request)
    except HTTPException as error:
        if error.status_code == 401:
            return RedirectResponse(url="/login", status_code=303)
        raise
    return _render_swarm_page(request, user_id, page_mode, agent_id)


@app.get("/swarm-intelligence/central", response_class=HTMLResponse, name="swarm_central_workspace")
async def swarm_central_workspace(request: Request):
    return await _swarm_page_or_login(request, "central")


@app.get("/swarm-intelligence/runs", response_class=HTMLResponse, name="swarm_runs_page")
async def swarm_runs_page(request: Request):
    return await _swarm_page_or_login(request, "runs")


@app.get("/swarm-intelligence/runs/{run_id}", response_class=HTMLResponse, name="swarm_run_page")
async def swarm_run_page(request: Request, run_id: str):
    return await _swarm_page_or_login(request, "run", run_id)


@app.get("/swarm-intelligence/agents/{agent_id}", response_class=HTMLResponse, name="swarm_agent_page")
async def swarm_agent_page(request: Request, agent_id: str):
    if agent_id not in SWARM_AGENT_REGISTRY:
        raise HTTPException(status_code=404, detail="That Swarm specialist does not exist.")
    return await _swarm_page_or_login(request, "agent", agent_id)


SKILLS_STUDIO_PHASES = {"assess", "learn", "practice", "build", "prove", "improve"}


@app.get("/api/skills/sessions", name="api_skills_sessions")
async def api_skills_sessions(request: Request):
    user_id = await require_user_id_async(request)
    sessions = await run_in_threadpool(list_skills_learning_sessions, str(user_id))
    return JSONResponse({"sessions": sessions}, headers={"Cache-Control": "no-store"})


@app.post("/api/skills/sessions", name="api_create_skills_session")
async def api_create_skills_session(request: Request):
    user_id = await require_user_id_async(request)
    payload = await _read_interview_json(request)
    skill = _interview_text(payload.get("skill"), "Skill", 120)
    goal = _interview_text(payload.get("goal"), "Learning goal", 500, required=False)
    current_level = _interview_text(payload.get("current_level"), "Current experience", 80, required=False)
    conversation_title = f"Skills Studio • {skill}"[:180]
    conversation_id = await run_in_threadpool(create_conversation, user_id, conversation_title)
    try:
        session = await run_in_threadpool(
            create_skills_learning_session,
            str(user_id),
            conversation_id,
            skill,
            goal,
            current_level,
        )
    except Exception:
        logger.exception("Could not create Skills Studio session for user %s", user_id)
        await run_in_threadpool(delete_conversation, conversation_id, user_id)
        raise HTTPException(status_code=500, detail="Your learning room could not be saved. Please try again.")
    return JSONResponse({"session": session}, status_code=201, headers={"Cache-Control": "no-store"})


@app.get("/api/skills/sessions/{session_id}", name="api_get_skills_session")
async def api_get_skills_session(request: Request, session_id: str):
    user_id = await require_user_id_async(request)
    session = await run_in_threadpool(get_skills_learning_session, str(user_id), session_id)
    if not session:
        raise HTTPException(status_code=404, detail="That learning room could not be found.")
    try:
        conversation_id = int(session.get("conversation_id"))
    except (TypeError, ValueError):
        raise HTTPException(status_code=404, detail="That learning room has no saved chat.")
    conversation = await run_in_threadpool(get_conversation, conversation_id, user_id)
    if not conversation or conversation.get("deleted"):
        raise HTTPException(status_code=404, detail="That learning room's saved chat could not be found.")
    return JSONResponse(
        {"session": session, "messages": conversation.get("messages") or []},
        headers={"Cache-Control": "no-store"},
    )


@app.patch("/api/skills/sessions/{session_id}", name="api_update_skills_session")
async def api_update_skills_session(request: Request, session_id: str):
    user_id = await require_user_id_async(request)
    payload = await _read_interview_json(request)
    existing = await run_in_threadpool(get_skills_learning_session, str(user_id), session_id)
    if not existing:
        raise HTTPException(status_code=404, detail="That learning room could not be found.")
    patch = {}
    if "phase_progress" in payload:
        progress = payload.get("phase_progress")
        if not isinstance(progress, dict) or any(str(key) not in SKILLS_STUDIO_PHASES for key in progress):
            raise HTTPException(status_code=400, detail="The phase progress is invalid.")
        clean_progress = {phase: 0 for phase in SKILLS_STUDIO_PHASES}
        clean_progress.update(existing.get("phase_progress") if isinstance(existing.get("phase_progress"), dict) else {})
        for phase, value in progress.items():
            if type(value) is not int or value < 0 or value > 100:
                raise HTTPException(status_code=400, detail="Phase progress must be between 0 and 100.")
            clean_progress[phase] = value
        patch["phase_progress"] = clean_progress
    if "active_phase" in payload:
        phase = str(payload.get("active_phase") or "").strip().casefold()
        if phase not in SKILLS_STUDIO_PHASES:
            raise HTTPException(status_code=400, detail="The active learning phase is invalid.")
        patch["active_phase"] = phase
    if "status" in payload:
        status = str(payload.get("status") or "").strip().casefold()
        if status not in {"in_progress", "completed"}:
            raise HTTPException(status_code=400, detail="The learning room status is invalid.")
        patch["status"] = status
    if "goal" in payload:
        patch["goal"] = _interview_text(payload.get("goal"), "Learning goal", 500, required=False)
    if "current_level" in payload:
        patch["current_level"] = _interview_text(payload.get("current_level"), "Current experience", 80, required=False)
    session = await run_in_threadpool(update_skills_learning_session, str(user_id), session_id, patch)
    if not session:
        raise HTTPException(status_code=404, detail="That learning room could not be saved.")
    return JSONResponse({"session": session}, headers={"Cache-Control": "no-store"})


async def _read_interview_json(request: Request):
    try:
        payload = await request.json()
    except Exception as error:
        raise HTTPException(status_code=400, detail="A valid JSON request is required.") from error
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="A valid JSON request is required.")
    return payload


def _interview_text(value, label: str, maximum: int, *, required: bool = True) -> str:
    text = str(value or "").strip()
    if required and not text:
        raise HTTPException(status_code=400, detail=f"{label} is required.")
    if len(text) > maximum:
        raise HTTPException(status_code=400, detail=f"{label} is too long.")
    return text


def _clean_interview_questions(value):
    if not isinstance(value, list) or len(value) not in {5, 10}:
        raise HTTPException(status_code=400, detail="A practice round must contain 5 or 10 questions.")
    questions = []
    for item in value:
        if not isinstance(item, dict):
            raise HTTPException(status_code=400, detail="An interview question is invalid.")
        question_type = str(item.get("type") or "written").strip().casefold()
        if question_type not in {"written", "multiple_choice"}:
            raise HTTPException(status_code=400, detail="An interview question type is invalid.")
        question = _interview_text(item.get("question"), "Question", 700)
        options = item.get("options") if isinstance(item.get("options"), list) else []
        if question_type == "multiple_choice":
            if len(options) != 4:
                raise HTTPException(status_code=400, detail="Each multiple-choice question needs four options.")
            options = [_interview_text(option, "Choice", 240) for option in options]
        else:
            options = []
        questions.append({"type": question_type, "question": question, "options": options})
    return questions


def _interview_session_summary(session):
    questions = session.get("questions") if isinstance(session.get("questions"), list) else []
    answers = session.get("answers") if isinstance(session.get("answers"), list) else []
    sample_answers = session.get("question_sample_answers") if isinstance(session.get("question_sample_answers"), list) else []
    sources = session.get("answer_sources") if isinstance(session.get("answer_sources"), list) else []
    answered_count = sum(
        bool(str(answer or "").strip()) or (
            index < len(sample_answers)
            and bool(str(sample_answers[index] or "").strip())
            and index < len(sources)
            and sources[index] == "oddi_sample"
        )
        for index, answer in enumerate(answers)
    )
    return {
        "id": str(session.get("id") or ""),
        "profession": str(session.get("profession") or "Interview practice"),
        "status": str(session.get("status") or "in_progress"),
        "has_review": bool(session.get("review")),
        "answered_count": answered_count,
        "question_count": len(questions),
        "created_at": session.get("created_at"),
        "updated_at": session.get("updated_at"),
    }


@app.get("/api/interview/sessions", name="api_interview_sessions")
async def api_interview_sessions(request: Request):
    user_id = await require_user_id_async(request)
    sessions = await run_in_threadpool(list_interview_sessions, str(user_id))
    return JSONResponse(
        {"sessions": [_interview_session_summary(session) for session in sessions]},
        headers={"Cache-Control": "no-store"},
    )


@app.get("/api/interview/sessions/{session_id}", name="api_interview_session")
async def api_interview_session(request: Request, session_id: str):
    user_id = await require_user_id_async(request)
    session = await run_in_threadpool(get_interview_session, str(user_id), session_id)
    if not session:
        raise HTTPException(status_code=404, detail="That practice session could not be found.")
    return JSONResponse({"session": session}, headers={"Cache-Control": "no-store"})


@app.patch("/api/interview/sessions/{session_id}", name="api_update_interview_session")
async def api_update_interview_session(request: Request, session_id: str):
    user_id = await require_user_id_async(request)
    payload = await _read_interview_json(request)
    if len(session_id) > 100:
        raise HTTPException(status_code=400, detail="Practice session ID is invalid.")
    questions = _clean_interview_questions(payload.get("questions"))
    raw_answers = payload.get("answers")
    if not isinstance(raw_answers, list) or len(raw_answers) != len(questions):
        raise HTTPException(status_code=400, detail="Every question needs an answer slot.")
    answers = [_interview_text(answer, "Answer", 10000, required=False) for answer in raw_answers]
    if sum(map(len, answers)) > 100000:
        raise HTTPException(status_code=400, detail="The combined answers are too long to save.")
    existing_session = await run_in_threadpool(get_interview_session, str(user_id), session_id)
    if not existing_session:
        raise HTTPException(status_code=404, detail="That practice session could not be found.")
    raw_samples = payload.get("question_sample_answers", existing_session.get("question_sample_answers", []))
    if not isinstance(raw_samples, list):
        raw_samples = []
    if not raw_samples:
        raw_samples = [""] * len(questions)
    if len(raw_samples) != len(questions):
        raise HTTPException(status_code=400, detail="Model answers must match the practice round.")
    sample_answers = [_interview_text(answer, "Model answer", 2500, required=False) for answer in raw_samples]
    raw_sources = payload.get("answer_sources", existing_session.get("answer_sources", []))
    if not isinstance(raw_sources, list):
        raw_sources = []
    if not raw_sources:
        raw_sources = ["candidate" if answers[index] else ("oddi_sample" if sample_answers[index] else "unanswered") for index in range(len(questions))]
    if len(raw_sources) != len(questions):
        raise HTTPException(status_code=400, detail="Answer sources must match the practice round.")
    answer_sources = []
    for index, source in enumerate(raw_sources):
        source = str(source or "").strip().casefold()
        if source not in {"unanswered", "candidate", "oddi_sample"}:
            raise HTTPException(status_code=400, detail="An answer source is invalid.")
        if answers[index]:
            source = "candidate"
        elif sample_answers[index]:
            source = "oddi_sample"
        else:
            source = "unanswered"
        answer_sources.append(source)
    session = await run_in_threadpool(
        update_interview_session,
        str(user_id),
        session_id,
        {
            "questions": questions,
            "answers": answers,
            "question_sample_answers": sample_answers,
            "answer_sources": answer_sources,
            "status": "in_progress",
            "formatted_answers": "",
            "review": "",
        },
    )
    if not session:
        raise HTTPException(status_code=404, detail="That practice session could not be found.")
    return JSONResponse({"session": _interview_session_summary(session)}, headers={"Cache-Control": "no-store"})


@app.post("/api/interview/questions", name="api_interview_questions")
async def api_interview_questions(request: Request):
    user_id = await require_user_id_async(request)
    payload = await _read_interview_json(request)
    profession = _interview_text(payload.get("profession"), "Profession", 160)
    context = _interview_text(payload.get("context"), "Role details", 500, required=False)
    count = payload.get("count", 10)
    if type(count) is not int or count not in {5, 10}:
        raise HTTPException(status_code=400, detail="Choose either 5 or 10 questions for a practice round.")
    question_mode = str(payload.get("question_mode") or "mixed").strip().casefold()
    if question_mode not in {"mixed", "written", "multiple_choice"}:
        raise HTTPException(status_code=400, detail="The requested question format is invalid.")
    difficulty = str(payload.get("difficulty") or "standard").strip().casefold()
    if difficulty not in {"standard", "easier", "challenging"}:
        raise HTTPException(status_code=400, detail="The requested difficulty is invalid.")
    raw_existing_questions = payload.get("existing_questions")
    existing_questions = []
    if isinstance(raw_existing_questions, list):
        existing_questions = [str(item).strip()[:700] for item in raw_existing_questions[:20] if str(item).strip()]
    email = str(request.session.get("email") or "")
    try:
        result = await run_in_threadpool(
            generate_question_set,
            str(user_id),
            email,
            profession,
            context,
            count,
            question_mode,
            difficulty,
            existing_questions,
        )
        if isinstance(result.get("questions"), list) and len(result["questions"]) == count:
            try:
                session = await run_in_threadpool(
                    create_interview_session,
                    str(user_id),
                    profession,
                    context,
                    result["questions"],
                )
                result["session_id"] = session["id"]
            except Exception:
                logger.exception("Interview Studio generated questions but could not save the practice session.")
                result["session_error"] = "Questions were created, but this practice could not be saved."
        return JSONResponse(result, headers={"Cache-Control": "no-store"})
    except HTTPException:
        raise
    except InterviewStudioError as error:
        return JSONResponse({"error": str(error)}, status_code=502, headers={"Cache-Control": "no-store"})
    except Exception:
        logger.exception("Interview Studio question generation failed.")
        return JSONResponse(
            {"error": "ODDI couldn't generate this interview right now. Please try again."},
            status_code=503,
            headers={"Cache-Control": "no-store"},
        )


@app.post("/api/interview/question", name="api_interview_alternate_question")
async def api_interview_alternate_question(request: Request):
    user_id = await require_user_id_async(request)
    payload = await _read_interview_json(request)
    profession = _interview_text(payload.get("profession"), "Profession", 160)
    context = _interview_text(payload.get("context"), "Role details", 500, required=False)
    excluded = payload.get("existing_questions")
    if not isinstance(excluded, list):
        excluded = []
    excluded = [str(item).strip()[:700] for item in excluded[:20] if str(item).strip()]
    email = str(request.session.get("email") or "")
    try:
        question = await run_in_threadpool(
            generate_alternate_question,
            str(user_id),
            email,
            profession,
            context,
            excluded,
        )
        return JSONResponse({"question": question}, headers={"Cache-Control": "no-store"})
    except HTTPException:
        raise
    except InterviewStudioError as error:
        return JSONResponse({"error": str(error)}, status_code=502, headers={"Cache-Control": "no-store"})
    except Exception:
        logger.exception("Interview Studio alternate question generation failed.")
        return JSONResponse(
            {"error": "ODDI couldn't create a different question right now. Please try again."},
            status_code=503,
            headers={"Cache-Control": "no-store"},
        )


@app.post("/api/interview/answer", name="api_interview_answer")
async def api_interview_answer(request: Request):
    user_id = await require_user_id_async(request)
    payload = await _read_interview_json(request)
    profession = _interview_text(payload.get("profession"), "Profession", 160)
    context = _interview_text(payload.get("context"), "Role details", 500, required=False)
    question = _interview_text(payload.get("question"), "Question", 700)
    session_id = _interview_text(payload.get("session_id"), "Practice session", 100, required=False)
    question_index = payload.get("question_index")
    question_type = str(payload.get("question_type") or "written").strip().casefold()
    if question_type not in {"written", "multiple_choice"}:
        raise HTTPException(status_code=400, detail="Question type is invalid.")
    raw_options = payload.get("options")
    options = []
    if question_type == "multiple_choice":
        if not isinstance(raw_options, list) or len(raw_options) != 4:
            raise HTTPException(status_code=400, detail="A multiple-choice question needs four options.")
        options = [_interview_text(option, "Choice", 240) for option in raw_options]

    email = str(request.session.get("email") or "")
    try:
        answer = await run_in_threadpool(
            generate_interview_answer,
            str(user_id),
            email,
            profession,
            context,
            question,
            question_type,
            options,
        )
        if session_id:
            if type(question_index) is not int or question_index < 0:
                raise HTTPException(status_code=400, detail="The question number is invalid.")
            session = await run_in_threadpool(get_interview_session, str(user_id), session_id)
            if not session or question_index >= len(session.get("questions") or []):
                raise HTTPException(status_code=404, detail="That question could not be found in this practice.")
            if str(session["questions"][question_index].get("question") or "") != question:
                raise HTTPException(status_code=409, detail="This question has changed. Refresh and try again.")
            samples = session.get("question_sample_answers")
            sources = session.get("answer_sources")
            samples = list(samples) if isinstance(samples, list) else []
            sources = list(sources) if isinstance(sources, list) else []
            while len(samples) < len(session["questions"]): samples.append("")
            while len(sources) < len(session["questions"]): sources.append("unanswered")
            samples[question_index] = answer
            candidates = session.get("answers") if isinstance(session.get("answers"), list) else []
            sources[question_index] = "candidate" if question_index < len(candidates) and str(candidates[question_index] or "").strip() else "oddi_sample"
            updated = await run_in_threadpool(
                update_interview_session,
                str(user_id),
                session_id,
                {"question_sample_answers": samples, "answer_sources": sources},
            )
            if not updated:
                raise HTTPException(status_code=404, detail="That practice session could not be found.")
        return JSONResponse({"answer": answer}, headers={"Cache-Control": "no-store"})
    except HTTPException:
        raise
    except InterviewStudioError as error:
        return JSONResponse({"error": str(error)}, status_code=502, headers={"Cache-Control": "no-store"})
    except Exception:
        logger.exception("Interview Studio sample answer generation failed.")
        return JSONResponse(
            {"error": "ODDI couldn't prepare this sample answer right now. Please try again."},
            status_code=503,
            headers={"Cache-Control": "no-store"},
        )


@app.post("/api/interview/answers", name="api_interview_answers")
async def api_interview_answers(request: Request):
    user_id = await require_user_id_async(request)
    payload = await _read_interview_json(request)
    profession = _interview_text(payload.get("profession"), "Profession", 160)
    context = _interview_text(payload.get("context"), "Role details", 500, required=False)
    questions = _clean_interview_questions(payload.get("questions"))
    session_id = _interview_text(payload.get("session_id"), "Practice session", 100, required=False)
    email = str(request.session.get("email") or "")
    try:
        answers = await run_in_threadpool(
            generate_interview_answer_set,
            str(user_id),
            email,
            profession,
            context,
            questions,
        )
        if session_id:
            session = await run_in_threadpool(get_interview_session, str(user_id), session_id)
            if not session or len(session.get("questions") or []) != len(questions):
                raise HTTPException(status_code=404, detail="That practice session could not be found.")
            if [str(item.get("question") or "") for item in session["questions"]] != [item["question"] for item in questions]:
                raise HTTPException(status_code=409, detail="This question set has changed. Refresh and try again.")
            candidates = session.get("answers") if isinstance(session.get("answers"), list) else []
            sources = ["candidate" if index < len(candidates) and str(candidates[index] or "").strip() else "oddi_sample" for index in range(len(questions))]
            updated = await run_in_threadpool(
                update_interview_session,
                str(user_id),
                session_id,
                {"question_sample_answers": answers, "answer_sources": sources},
            )
            if not updated:
                raise HTTPException(status_code=404, detail="That practice session could not be found.")
        return JSONResponse(
            {"answers": answers},
            headers={"Cache-Control": "no-store"},
        )
    except HTTPException:
        raise
    except InterviewStudioError as error:
        return JSONResponse({"error": str(error)}, status_code=502, headers={"Cache-Control": "no-store"})
    except Exception:
        logger.exception("Interview Studio full sample answer generation failed.")
        return JSONResponse(
            {"error": "ODDI couldn't prepare model answers for this practice round. Please try again."},
            status_code=503,
            headers={"Cache-Control": "no-store"},
        )


@app.post("/api/interview/review", name="api_interview_review")
async def api_interview_review(request: Request):
    user_id = await require_user_id_async(request)
    payload = await _read_interview_json(request)
    profession = _interview_text(payload.get("profession"), "Profession", 160)
    session_id = _interview_text(payload.get("session_id"), "Practice session", 100, required=False)
    existing_session = None
    if session_id:
        existing_session = await run_in_threadpool(get_interview_session, str(user_id), session_id)
        if not existing_session:
            raise HTTPException(status_code=404, detail="That practice session could not be found.")
    raw_answers = payload.get("answers")
    expected_count = len(existing_session.get("questions") or []) if existing_session else None
    if not isinstance(raw_answers, list) or len(raw_answers) not in {5, 10}:
        raise HTTPException(status_code=400, detail="Answer all questions in this 5- or 10-question round before submitting.")
    if expected_count is not None and len(raw_answers) != expected_count:
        raise HTTPException(status_code=400, detail="The answers do not match this practice round.")

    saved_samples = existing_session.get("question_sample_answers") if existing_session else []
    saved_sources = existing_session.get("answer_sources") if existing_session else []
    answers = []
    candidate_answers = []
    normalized_samples = []
    answer_sources = []
    for index, answer in enumerate(raw_answers):
        if not isinstance(answer, dict):
            raise HTTPException(status_code=400, detail="Every interview answer must include its question and answer.")
        question = _interview_text(answer.get("question"), "Question", 700)
        if existing_session and str(existing_session["questions"][index].get("question") or "") != question:
            raise HTTPException(status_code=409, detail="This question set has changed. Refresh and try again.")
        response = _interview_text(answer.get("answer"), "Answer", 10000, required=False)
        sample = _interview_text(
            answer.get("sample_answer") or (saved_samples[index] if index < len(saved_samples) else ""),
            "Model answer",
            2500,
            required=False,
        )
        requested_source = str(answer.get("answer_source") or (saved_sources[index] if index < len(saved_sources) else "") or "").strip().casefold()
        if requested_source and requested_source not in {"candidate", "oddi_sample", "unanswered"}:
            raise HTTPException(status_code=400, detail="An answer source is invalid.")
        source = "candidate" if response else ("oddi_sample" if sample else "unanswered")
        if not response and source != "oddi_sample":
            raise HTTPException(status_code=400, detail="Answer each question yourself or ask ODDI for its sample answer first.")
        effective_response = response if response else sample
        answers.append({
            "question": question,
            "answer": effective_response,
            "answer_source": source,
            "sample_answer": sample,
        })
        candidate_answers.append(response)
        normalized_samples.append(sample)
        answer_sources.append(source)

    formatted_size = sum(
        len(item["question"])
        + len(item["answer"])
        + (len(item.get("sample_answer") or "") if item.get("answer_source") == "candidate" else 0)
        for item in answers
    )
    if formatted_size > 50000:
        raise HTTPException(status_code=400, detail="The combined answers are too long to review.")

    email = str(request.session.get("email") or "")
    try:
        formatted_answers, review = await run_in_threadpool(
            review_interview_answers,
            str(user_id),
            email,
            profession,
            answers,
        )
        if session_id:
            updated_session = await run_in_threadpool(
                update_interview_session,
                str(user_id),
                session_id,
                {
                    "answers": candidate_answers,
                    "question_sample_answers": normalized_samples,
                    "answer_sources": answer_sources,
                    "formatted_answers": formatted_answers,
                    "review": review,
                    "status": "completed",
                },
            )
            if not updated_session:
                raise HTTPException(status_code=404, detail="That practice session could not be found.")
        return JSONResponse(
            {"session_id": session_id or None, "formatted_answers": formatted_answers, "review": review},
            headers={"Cache-Control": "no-store"},
        )
    except HTTPException:
        raise
    except Exception:
        logger.exception("Interview Studio answer review failed.")
        return JSONResponse(
            {"error": "ODDI couldn't review these answers right now. Please try again."},
            status_code=503,
            headers={"Cache-Control": "no-store"},
        )


@app.get("/login", response_class=HTMLResponse, name="login_page")
def login_page(request: Request):
    return render_template(request, "login.html")


LEGAL_TERMS_VERSION = "2026-10-22"
LEGAL_PRIVACY_VERSION = "2026-10-22"
_AUTH_EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")
_PASSWORD_RESET_RATE_LOCK = threading.Lock()
_PASSWORD_RESET_ATTEMPTS = {}


def _valid_email(value):
    value = str(value or "").strip()
    return len(value) <= 254 and bool(_AUTH_EMAIL_RE.fullmatch(value))


def _terms_accepted(value):
    return value is True or str(value or "").strip().lower() in {"1", "true", "yes", "on", "accepted"}


def _validate_signup(username, email, password, accepted_terms):
    username = str(username or "").strip()
    email = str(email or "").strip().casefold()
    if len(username) < 2:
        return username, email, "Name must be at least 2 characters."
    if len(username) > 80:
        return username, email, "Name must be 80 characters or fewer."
    if not _valid_email(email):
        return username, email, "Enter a valid email address."
    if not isinstance(password, str) or len(password) < 8:
        return username, email, "Password must be at least 8 characters."
    if len(password) > 128:
        return username, email, "Password must be 128 characters or fewer."
    if not _terms_accepted(accepted_terms):
        return username, email, "Accept the Terms of Service and Privacy Policy to create an account."
    return username, email, None


def _signup_legal_settings():
    return {
        "legal": {
            "terms_version": LEGAL_TERMS_VERSION,
            "privacy_version": LEGAL_PRIVACY_VERSION,
            "accepted_at": datetime.now(timezone.utc).isoformat(),
        }
    }


def _password_reset_base_url(request):
    configured = os.getenv("ODDI_PUBLIC_BASE_URL", "").strip().rstrip("/")
    if configured:
        try:
            parsed = urlsplit(configured)
            parsed.port
        except ValueError:
            return ""
        local_http = parsed.hostname in {"localhost", "127.0.0.1", "::1"} and parsed.scheme == "http"
        if (parsed.scheme != "https" and not local_http) or not parsed.hostname or not parsed.netloc or parsed.username or parsed.password or parsed.query or parsed.fragment:
            return ""
        return configured
    if _IS_PRODUCTION_SERVER:
        return ""
    return str(request.base_url).rstrip("/") if request else ""


def _password_reset_email_configured(request):
    return bool(
        os.getenv("ODDI_SMTP_HOST", "").strip()
        and (os.getenv("ODDI_SMTP_FROM", "").strip() or os.getenv("ODDI_SMTP_USER", "").strip())
        and _password_reset_base_url(request)
    )


def _send_password_reset_email(email, reset_url):
    host = os.getenv("ODDI_SMTP_HOST", "").strip()
    sender = os.getenv("ODDI_SMTP_FROM", "").strip() or os.getenv("ODDI_SMTP_USER", "").strip()
    username = os.getenv("ODDI_SMTP_USER", "").strip()
    password = os.getenv("ODDI_SMTP_PASSWORD", "")
    port = int(os.getenv("ODDI_SMTP_PORT", "587").strip() or "587")
    use_ssl = os.getenv("ODDI_SMTP_USE_SSL", "").strip().lower() in {"1", "true", "yes", "on"}
    use_starttls = os.getenv("ODDI_SMTP_STARTTLS", "true").strip().lower() in {"1", "true", "yes", "on"}

    message = EmailMessage()
    message["Subject"] = "Reset your ODDI password"
    message["From"] = sender
    message["To"] = email
    message.set_content(
        "We received a request to reset the password for your ODDI account.\n\n"
        f"Use this one-time link within 30 minutes:\n{reset_url}\n\n"
        "If you did not request this, you can ignore this message."
    )

    context = ssl.create_default_context()
    smtp = (
        smtplib.SMTP_SSL(host, port, timeout=20, context=context)
        if use_ssl
        else smtplib.SMTP(host, port, timeout=20)
    )
    try:
        smtp.ehlo()
        if not use_ssl and use_starttls:
            smtp.starttls(context=context)
            smtp.ehlo()
        if username:
            smtp.login(username, password)
        smtp.send_message(message)
    finally:
        try:
            smtp.quit()
        except Exception:
            smtp.close()


def _rate_limit_password_reset(request, email):
    now = time.monotonic()
    client_host = request.client.host if request.client else "unknown"
    keys = (
        ("ip:" + hashlib.sha256(client_host.encode("utf-8")).hexdigest(), 20),
        ("email:" + hashlib.sha256(email.encode("utf-8")).hexdigest(), 4),
    )
    with _PASSWORD_RESET_RATE_LOCK:
        for key, _ in keys:
            recent = [stamp for stamp in _PASSWORD_RESET_ATTEMPTS.get(key, []) if now - stamp < 3600]
            _PASSWORD_RESET_ATTEMPTS[key] = recent
            limit = 20 if key.startswith("ip:") else 4
            if len(recent) >= limit:
                return False
        for key, _ in keys:
            _PASSWORD_RESET_ATTEMPTS[key].append(now)
        if len(_PASSWORD_RESET_ATTEMPTS) > 5000:
            for key in list(_PASSWORD_RESET_ATTEMPTS):
                recent = [stamp for stamp in _PASSWORD_RESET_ATTEMPTS[key] if now - stamp < 3600]
                if recent:
                    _PASSWORD_RESET_ATTEMPTS[key] = recent
                else:
                    _PASSWORD_RESET_ATTEMPTS.pop(key, None)
        return True


def _issue_password_reset(email, base_url):
    user = get_user_by_email(email)
    if not user or not str(user["password_hash"] or ""):
        return
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    expires_at = (datetime.now(timezone.utc) + timedelta(minutes=30)).isoformat()
    reset_url = f"{base_url.rstrip('/')}/reset-password?token={quote(token, safe='')}"
    try:
        store_password_reset_token(email, token_hash, expires_at)
        _send_password_reset_email(email, reset_url)
    except Exception:
        try:
            clear_password_reset_token(email, token_hash)
        except Exception:
            logger.exception("Could not remove an undelivered password reset token")
        logger.exception("Password reset request could not be completed")


def _no_store(response):
    response.headers["Cache-Control"] = "no-store"
    response.headers["Referrer-Policy"] = "no-referrer"
    return response


@app.get("/forgot-password", response_class=HTMLResponse, name="forgot_password_page")
def forgot_password_page(request: Request):
    return _no_store(render_template(
        request,
        "forgot_password.html",
        email_configured=_password_reset_email_configured(request),
        submitted=False,
        message="",
        is_error=False,
        email_value="",
    ))


@app.post("/forgot-password", response_class=HTMLResponse, name="request_password_reset")
async def request_password_reset(request: Request, background_tasks: BackgroundTasks):
    form = await request.form()
    email = str(form.get("email") or "").strip().casefold()
    configured = _password_reset_email_configured(request)
    message = ""
    is_error = False
    if not _valid_email(email):
        message = "Enter a valid email address."
        is_error = True
    elif not configured:
        message = "Email recovery is not configured on this server yet. If you normally use Google, continue with Google sign-in."
        is_error = True
    elif not _rate_limit_password_reset(request, email):
        message = "Too many reset requests. Wait a while before trying again."
        is_error = True
    else:
        base_url = _password_reset_base_url(request)
        background_tasks.add_task(_issue_password_reset, email, base_url)
        message = "If this email belongs to an account with password sign-in and delivery is available, a reset link will arrive shortly. If it doesn’t, try Google sign-in when linked to your account."

    return _no_store(render_template(
        request,
        "forgot_password.html",
        email_configured=configured,
        submitted=True,
        message=message,
        is_error=is_error,
        email_value=email,
    ))


@app.get("/reset-password", response_class=HTMLResponse, name="reset_password_page")
def reset_password_page(request: Request):
    token = str(request.query_params.get("token") or "").strip()
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest() if 20 <= len(token) <= 256 else ""
    valid = bool(token_hash and get_password_reset_email(token_hash))
    return _no_store(render_template(request, "reset_password.html", token=token if valid else "", valid=valid, message="", is_error=False))


@app.post("/reset-password", response_class=HTMLResponse, name="reset_password")
async def reset_password(request: Request):
    form = await request.form()
    token = str(form.get("token") or "").strip()
    password = str(form.get("password") or "")
    confirm = str(form.get("confirm_password") or "")
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest() if 20 <= len(token) <= 256 else ""
    valid = bool(token_hash and get_password_reset_email(token_hash))
    message = ""
    is_error = False

    if not valid:
        message = "This reset link is invalid or expired. Request a new one to continue."
        is_error = True
    elif len(password) < 8 or len(password) > 128:
        message = "Choose a password between 8 and 128 characters."
        is_error = True
    elif password != confirm:
        message = "The passwords do not match."
        is_error = True
    else:
        email = consume_password_reset_token(token_hash)
        if not email:
            valid = False
            message = "This reset link has already been used or expired. Request a new one to continue."
            is_error = True
        else:
            updated = await run_in_threadpool(
                update_user_password,
                email,
                generate_password_hash(password),
            )
            if updated:
                return RedirectResponse(url="/login?password_reset=1", status_code=303)
            valid = False
            message = "The account could not be updated. Request a fresh reset link and try again."
            is_error = True

    return _no_store(render_template(request, "reset_password.html", token=token if valid else "", valid=valid, message=message, is_error=is_error))


@app.get("/terms", response_class=HTMLResponse, name="terms_page")
def terms_page(request: Request):
    return render_template(request, "legal.html", legal_page="terms", version=LEGAL_TERMS_VERSION)


@app.get("/privacy", response_class=HTMLResponse, name="privacy_page")
def privacy_page(request: Request):
    return render_template(request, "legal.html", legal_page="privacy", version=LEGAL_PRIVACY_VERSION)


@app.post("/login", name="login")
async def login(request: Request):
    form = await request.form()

    email = str(form.get("email", "")).strip()
    password = str(form.get("password", ""))

    def login_error(code):
        return RedirectResponse(
            url=f"/login?error={code}&email={quote(email, safe='')}",
            status_code=303,
        )

    if not email or not password:
        return login_error("missing_fields")

    user = await run_in_threadpool(get_user_by_email, email)
    if user is None:
        await sync_laptop_accounts_from_drive()
        user = await run_in_threadpool(get_user_by_email, email)

    if not user:
        return login_error("account_not_found")

    password_hash = str(user["password_hash"] or "")
    if not password_hash:
        return login_error("google_only")

    if check_password_hash(password_hash, password):
        set_user_session(request, user)
        return RedirectResponse(url="/", status_code=303)

    return login_error("incorrect_password")


@app.get("/logout", name="logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse(url="/login", status_code=303)


@app.get("/profile", name="profile")
def profile(request: Request):
    if "user_id" not in request.session:
        return RedirectResponse(url="/login", status_code=303)
    return RedirectResponse(url="/?openProfile=1", status_code=303)


@app.get("/api/account/stats", name="api_account_stats")
def api_account_stats(request: Request):
    user_id = require_user_id(request)
    conversations = get_conversations(user_id)
    lifetime_tokens = 0
    peak_tokens = 0
    longest_task_seconds = 0
    activity_timestamps = []

    def parse_timestamp(value):
        if not value:
            return None
        if isinstance(value, datetime):
            parsed = value
        else:
            try:
                parsed = datetime.fromisoformat(str(value).strip().replace("Z", "+00:00"))
            except (TypeError, ValueError):
                return None
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed

    for conversation in conversations if isinstance(conversations, list) else []:
        if not isinstance(conversation, dict) or conversation.get("deleted") or conversation.get("deleted_at"):
            continue
        messages = conversation.get("messages") if isinstance(conversation.get("messages"), list) else []
        conversation_tokens = 0
        conversation_times = []
        for message in messages:
            if not isinstance(message, dict):
                continue
            role = str(message.get("role") or "").lower()
            text = str(message.get("text") or message.get("content") or message.get("message") or "")
            if role == "assistant":
                response_meta = message.get("response_meta") if isinstance(message.get("response_meta"), dict) else {}
                try:
                    token_count = max(0, int(response_meta.get("token_estimate") or 0))
                except (TypeError, ValueError, OverflowError):
                    token_count = 0
                conversation_tokens += token_count or ((len(text) + 3) // 4)
                try:
                    elapsed_ms = max(0, int(response_meta.get("elapsed_ms") or response_meta.get("latency_ms") or 0))
                except (TypeError, ValueError, OverflowError):
                    elapsed_ms = 0
                longest_task_seconds = max(longest_task_seconds, elapsed_ms // 1000)
            elif role == "user":
                conversation_tokens += (len(text) + 3) // 4

            for key in ("completed_at", "created_at", "timestamp", "sent_at"):
                parsed = parse_timestamp(message.get(key))
                if parsed:
                    conversation_times.append(parsed)
                    activity_timestamps.append(parsed.isoformat())
                    break

        lifetime_tokens += conversation_tokens
        peak_tokens = max(peak_tokens, conversation_tokens)

        if not conversation_times:
            for key in ("created_at", "updated_at"):
                parsed = parse_timestamp(conversation.get(key))
                if parsed:
                    conversation_times.append(parsed)
                    activity_timestamps.append(parsed.isoformat())
        if len(conversation_times) > 1:
            span_seconds = max(0, int((max(conversation_times) - min(conversation_times)).total_seconds()))
            longest_task_seconds = max(longest_task_seconds, span_seconds)

    return JSONResponse({
        "lifetime_tokens": lifetime_tokens,
        "peak_tokens": peak_tokens,
        "longest_task_seconds": longest_task_seconds,
        "activity_timestamps": sorted(set(activity_timestamps)),
        "token_estimates": True,
    })


@app.get("/signup", response_class=HTMLResponse, name="signup_page")
def signup_page(request: Request):
    email = str(request.query_params.get("email") or "").strip()
    return render_template(request, "signup.html", signup_email=email, signup_error="")


@app.post("/signup", name="signup")
async def signup(request: Request):
    form = await request.form()

    username, email, validation_error = _validate_signup(
        form.get("username"),
        form.get("email"),
        form.get("password"),
        form.get("accept_terms"),
    )
    password = form.get("password") if isinstance(form.get("password"), str) else ""
    if validation_error:
        return render_template(
            request,
            "signup.html",
            signup_email=email,
            signup_error=validation_error,
        )

    await sync_laptop_accounts_from_drive()

    if get_user_by_email(email):
        return render_template(
            request,
            "signup.html",
            signup_email=email,
            signup_error="An account with this email already exists. Sign in instead, or use Google if that is how you normally access ODDI.",
        )

    password_hash = generate_password_hash(password)
    create_user(username, email, password_hash, initial_settings=_signup_legal_settings())

    user = get_user_by_email(email)
    if user is None:
        return PlainTextResponse("Account was created but could not be loaded.", status_code=500)

    set_user_session(request, user)
    return RedirectResponse(url="/", status_code=303)


# =========================================================
# AUTH API
# =========================================================

@app.post("/api/auth/login", name="api_login")
async def api_login(request: Request):
    data = await safe_json(request)

    email = str(data.get("email", "")).strip()
    password = str(data.get("password") or "")

    if not email or not password:
        return JSONResponse(
            {"success": False, "error": "Email and password are required."},
            status_code=400,
        )

    user = await run_in_threadpool(get_user_by_email, email)
    if user is None:
        await sync_laptop_accounts_from_drive()
        user = await run_in_threadpool(get_user_by_email, email)

    if user is None:
        return JSONResponse(
            {"success": False, "error": "Email not found."},
            status_code=401,
        )

    password_hash = str(user["password_hash"] or "")
    if not password_hash:
        return JSONResponse(
            {"success": False, "error": "This account uses Google sign-in."},
            status_code=401,
        )

    if not check_password_hash(password_hash, password):
        return JSONResponse(
            {"success": False, "error": "Incorrect password."},
            status_code=401,
        )

    set_user_session(request, user)

    return JSONResponse(
        {
            "success": True,
            "user": {
                "id": user["id"],
                "name": user["username"],
                "email": user["email"],
            },
        }
    )


@app.post("/api/auth/register", name="api_register")
async def api_register(request: Request):
    data = await safe_json(request)

    username, email, validation_error = _validate_signup(
        data.get("name"),
        data.get("email"),
        data.get("password"),
        data.get("accept_terms", data.get("terms_accepted")),
    )
    password = data.get("password") if isinstance(data.get("password"), str) else ""

    if validation_error:
        return JSONResponse(
            {
                "success": False,
                "error": validation_error,
            },
            status_code=400,
        )

    await sync_laptop_accounts_from_drive()

    existing_user = get_user_by_email(email)

    if existing_user:
        return JSONResponse(
            {
                "success": False,
                "error": "An account with this email already exists.",
            },
            status_code=409,
        )

    password_hash = generate_password_hash(password)
    create_user(username, email, password_hash, initial_settings=_signup_legal_settings())

    user = get_user_by_email(email)

    if user is None:
        return JSONResponse(
            {
                "success": False,
                "error": "Account was created but could not be loaded.",
            },
            status_code=500,
        )

    set_user_session(request, user)

    return JSONResponse(
        {
            "success": True,
            "user": {
                "id": user["id"],
                "name": user["username"],
                "email": user["email"],
            },
        }
    )


# =========================================================
# GOOGLE OAUTH
# =========================================================

@app.get("/auth/google", name="google_login")
async def google_login(request: Request):
    if not google_client_id or not google_client_secret:
        logger.error("Google OAuth requested but credentials are missing.")
        return PlainTextResponse(
            "Google sign-in is not configured. "
            "Set GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET in .env.",
            status_code=503,
        )

    redirect_uri = str(request.url_for("google_callback"))

    return await oauth.google.authorize_redirect(
        request,
        redirect_uri,
    )


@app.get("/auth/google/drive", name="google_drive_reauthorize")
async def google_drive_reauthorize(request: Request):
    """Refresh the laptop's Drive grant without exposing its token."""
    if not ODDI_LOCAL_ONLY_STORAGE:
        return PlainTextResponse("Drive reconnection is available from the laptop only.", status_code=404)
    if not _get_session_user(request):
        return RedirectResponse(url="/login", status_code=303)
    if not google_client_id or not google_client_secret:
        return PlainTextResponse("Google sign-in is not configured on this laptop.", status_code=503)

    current_user = _get_session_user(request)
    request.session["drive_reauthorization_pending"] = True
    redirect_uri = str(request.url_for("google_callback"))
    return await oauth.google.authorize_redirect(
        request,
        redirect_uri,
        scope="openid profile email https://www.googleapis.com/auth/drive",
        access_type="offline",
        prompt="consent",
        login_hint=str(current_user["email"] or ""),
    )


@app.get("/auth/google/callback", name="google_callback")
async def google_callback(request: Request):
    drive_reauthorization = bool(request.session.pop("drive_reauthorization_pending", False))
    try:
        token = await oauth.google.authorize_access_token(request)
    except OAuthError as error:
        logger.warning("Google OAuth failed: %s", error)
        if drive_reauthorization:
            return HTMLResponse(
                "<main><h1>Drive was not reconnected</h1><p>Your chats are still saved on this laptop. "
                "Please retry Google Drive reconnection from this laptop.</p></main>",
                status_code=400,
            )
        return RedirectResponse(url="/login", status_code=303)

    userinfo = token.get("userinfo")

    if not userinfo:
        try:
            userinfo = await oauth.google.userinfo(token=token)
        except Exception as error:
            logger.exception(
                "Unable to retrieve Google user information: %s",
                error,
            )
            return PlainTextResponse(
                "Google account information could not be retrieved.",
                status_code=502,
            )

    if drive_reauthorization:
        if not ODDI_LOCAL_ONLY_STORAGE or not _get_session_user(request):
            return PlainTextResponse("Drive reconnection must start from a signed-in laptop session.", status_code=403)
        refresh_token = str(token.get("refresh_token") or "").strip()
        if not refresh_token:
            return HTMLResponse(
                "<main><h1>Google did not issue a Drive refresh token</h1><p>Your chats are still saved on this laptop. "
                "Retry Drive reconnection and approve offline Drive access.</p></main>",
                status_code=400,
            )
        granted_scopes = set(str(token.get("scope") or "").split())
        drive_scope = "https://www.googleapis.com/auth/drive"
        if granted_scopes and drive_scope not in granted_scopes:
            return HTMLResponse(
                "<main><h1>Drive access was not granted</h1><p>Your chats are still saved on this laptop. "
                "Retry and approve the requested Google Drive access.</p></main>",
                status_code=403,
            )

        from app import google_drive_storage

        try:
            await run_in_threadpool(
                google_drive_storage.validate_access_token,
                token.get("access_token"),
            )
            dotenv_path = os.path.join(os.path.dirname(__file__), ".env")
            saved, _, _ = set_key(
                dotenv_path,
                "GOOGLE_REFRESH_TOKEN",
                refresh_token,
                quote_mode="always",
            )
            if not saved:
                raise OSError("Could not save the refreshed Google Drive credential.")
            os.environ["GOOGLE_REFRESH_TOKEN"] = refresh_token
        except Exception as error:
            logger.warning(
                "Could not save the renewed local Google Drive grant (error_type=%s)",
                type(error).__name__,
            )
            message = (
                str(error)
                if isinstance(error, DriveStorageError)
                else "The renewed Google Drive credential could not be saved on this laptop."
            )
            return HTMLResponse(
                f"<main><h1>Drive reconnection could not finish</h1><p>{message}</p>"
                "<p>Your existing laptop chats are untouched.</p></main>",
                status_code=503,
            )

        try:
            from app.drive_sync import sync_laptop_once

            sync_result = await run_in_threadpool(sync_laptop_once)
        except Exception:
            logger.exception("Initial laptop-to-Drive sync failed after Drive reauthorization; retrying in background")
            return HTMLResponse(
                "<main><h1>Drive reconnected on this laptop</h1><p>The credential was saved securely. "
                "Chat sync will retry automatically; your laptop copies remain safe.</p>"
                "<p>To restore the public website, update GOOGLE_REFRESH_TOKEN in the ODDI web service's "
                "Render Environment settings with the renewed value from this laptop's .env, then redeploy.</p></main>",
                status_code=503,
            )

        checked_chats = int(sync_result.get("chats", 0) or 0)
        return HTMLResponse(
            "<main><h1>Drive reconnected on this laptop</h1>"
            f"<p>Initial sync completed; {checked_chats} chat records were checked.</p>"
            "<p>To restore the public website, update GOOGLE_REFRESH_TOKEN in the ODDI web service's "
            "Render Environment settings with the renewed value from this laptop's .env, then redeploy. "
            "The token was saved in this laptop's local .env file and was not shown on this page.</p>"
            "<p><a href='/'>Return to ODDI</a></p></main>"
        )

    email = str(userinfo.get("email") or "").strip().lower()
    name = str(userinfo.get("name") or "").strip()
    email_verified = userinfo.get("email_verified", False)

    if not email:
        return PlainTextResponse("Google did not provide an email address.", status_code=400)

    if email_verified is not True:
        return PlainTextResponse("Google email is not verified.", status_code=403)

    # Preserve the current ODDI behavior: Gmail accounts only.
    if not email.endswith("@gmail.com"):
        return PlainTextResponse(
            "Please use a Gmail account to sign in to ODDI.",
            status_code=403,
        )

    if not name:
        name = email.split("@")[0]

    user = get_user_by_email(email)
    if user is None:
        await sync_laptop_accounts_from_drive()
        user = get_user_by_email(email)

    if user is None:
        # Google-only accounts do not have a local password. Store an empty
        # hash so password sign-in can explain that Google is the right path.
        create_user(name, email, "")
        user = get_user_by_email(email)

    if user is None:
        return PlainTextResponse(
            "Google login succeeded, but the ODDI account could not be created.",
            status_code=500,
        )

    set_user_session(request, user, auth_provider="google")

    return RedirectResponse(url="/", status_code=303)


@app.get("/api/auth/me", name="api_auth_me")
def api_auth_me(request: Request):
    user_id = request.session.get("user_id")

    if not user_id:
        return JSONResponse({"authenticated": False}, status_code=401)

    user = _get_session_user(request)
    if not user:
        request.session.clear()
        return JSONResponse({"authenticated": False}, status_code=401)

    return JSONResponse(
        {
            "authenticated": True,
            "user": {
                "id": user["id"],
                "name": user["username"],
                "email": user["email"],
            },
            "can_submit_feedback": True,
            "auth_provider": request.session.get("auth_provider", "password"),
        }
    )


def _help_feedback_session_user(request: Request, *, allow_guest=False):
    user = _get_session_user(request)
    if not user:
        if allow_guest:
            return None
        raise HTTPException(status_code=401, detail="Sign in to use feedback.")
    # Local account lookups return sqlite3.Row, while trusted Drive sessions
    # return a dict. Feedback handlers use .get(), so normalize both shapes.
    keys = getattr(user, "keys", None)
    if callable(keys):
        return {key: user[key] for key in keys()}
    return user


def _require_help_feedback_admin(request: Request):
    user = _help_feedback_session_user(request)
    if not is_help_feedback_admin(user.get("email")):
        raise HTTPException(status_code=403, detail="The feedback inbox is only available to the ODDI host account.")
    return user


@app.get("/api/help-feedback/state", name="api_help_feedback_state")
async def api_help_feedback_state(request: Request):
    user = await run_in_threadpool(_help_feedback_session_user, request, allow_guest=True)
    if not user:
        return JSONResponse({"authenticated": False, "can_submit": False, "is_admin": False, "submissions": [], "unread_replies": []})
    browser_local = is_temporary_browser_account(request)
    payload = {
        "authenticated": True,
        "can_submit": True,
        "is_admin": is_help_feedback_admin(user.get("email")),
        "user_email": str(user.get("email") or ""),
        "submissions": [],
        "cooldown_until": None,
        "unread_replies": [],
    }
    if browser_local:
        payload["account_message"] = "Feedback and replies are linked to this browser-only account. Sign in to a personal ODDI account to keep them across devices."
    state = await run_in_threadpool(get_help_feedback_state, user["id"])
    payload.update(state)
    return JSONResponse(payload)


@app.get("/api/help-feedback/notifications", name="api_help_feedback_notifications")
async def api_help_feedback_notifications(request: Request):
    user = await run_in_threadpool(_help_feedback_session_user, request, allow_guest=True)
    if not user:
        return JSONResponse({"unread_replies": []})
    replies = await run_in_threadpool(get_unread_help_feedback_replies, user["id"])
    return JSONResponse({"unread_replies": replies})


@app.post("/api/help-feedback/submissions", name="api_help_feedback_submit")
async def api_help_feedback_submit(request: Request):
    user = await run_in_threadpool(_help_feedback_session_user, request)
    data = await safe_json(request)
    try:
        submission = await run_in_threadpool(
            save_help_feedback_submission,
            user["id"],
            user.get("username"),
            user.get("email"),
            data.get("ratings"),
            data.get("reasons"),
            data.get("overall_feedback"),
        )
    except FeedbackCooldownError as exc:
        return JSONResponse(
            {"error": str(exc), "cooldown_until": exc.cooldown_until},
            status_code=429,
        )
    except FeedbackValidationError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    except DriveStorageError as exc:
        logger.exception("ODDI feedback could not be saved to Google Drive")
        return JSONResponse({"error": str(exc)}, status_code=503)
    except Exception:
        logger.exception("ODDI feedback submission failed")
        return JSONResponse(
            {"error": "ODDI couldn't save your feedback. Your answers are still on this screen; please try again."},
            status_code=500,
        )
    return JSONResponse({"submission": submission}, status_code=201)


@app.get("/api/help-feedback/admin", name="api_help_feedback_admin")
async def api_help_feedback_admin(request: Request):
    await run_in_threadpool(_require_help_feedback_admin, request)
    overview = await run_in_threadpool(admin_feedback_overview)
    return JSONResponse(overview)


@app.post("/api/help-feedback/admin/{submission_id}/replies", name="api_help_feedback_admin_reply")
async def api_help_feedback_admin_reply(request: Request, submission_id: str):
    host = await run_in_threadpool(_require_help_feedback_admin, request)
    data = await safe_json(request)
    try:
        reply = await run_in_threadpool(
            admin_feedback_reply,
            submission_id,
            data.get("target_key"),
            data.get("body"),
            host["id"],
            host.get("email"),
        )
    except FeedbackValidationError as exc:
        return JSONResponse({"error": str(exc)}, status_code=400)
    except LookupError as exc:
        return JSONResponse({"error": str(exc)}, status_code=404)
    return JSONResponse({"reply": reply}, status_code=201)


@app.post("/api/help-feedback/replies/read", name="api_help_feedback_replies_read")
async def api_help_feedback_replies_read(request: Request):
    user = await run_in_threadpool(_help_feedback_session_user, request)
    data = await safe_json(request)
    reply_ids = data.get("reply_ids") if isinstance(data.get("reply_ids"), list) else []
    updated = await run_in_threadpool(mark_help_feedback_replies_read, user["id"], reply_ids[:100])
    return JSONResponse({"updated": updated})


@app.get("/api/settings", name="api_get_settings")
def api_get_settings(request: Request):
    user_id = require_user_id(request)
    settings = get_user_settings(user_id)
    return JSONResponse({"settings": settings})


@app.put("/api/settings", name="api_update_settings")
async def api_update_settings(request: Request):
    user_id = require_user_id(request)
    data = await safe_json(request)
    incoming = data.get("settings")
    if not isinstance(incoming, dict):
        return JSONResponse({"error": "Settings must be an object."}, status_code=400)
    privacy_update = incoming.get("privacy")
    if isinstance(privacy_update, dict) and "retention_days" in privacy_update:
        retention_value = str(privacy_update.get("retention_days") or "").strip().lower()
        if retention_value not in {"30", "90", "365", "forever"}:
            return JSONResponse({"error": "Choose a valid conversation retention period."}, status_code=400)
    memory_update = incoming.get("memory")
    requested_memory = (
        normalize_memory_settings(memory_update)
        if isinstance(memory_update, dict)
        else None
    )
    if requested_memory:
        # Publish the opt-out before waiting for any in-flight extraction to
        # finish, so it cannot start a fallback provider request meanwhile.
        if not requested_memory["enabled"]:
            set_memory_user_disabled(user_id, True)
        if not requested_memory["enabled"] or not requested_memory["auto_extract"]:
            set_memory_user_opt_out(user_id, True)
    personalization_update = incoming.get("personalization")
    if isinstance(personalization_update, dict):
        instruction_value = personalization_update.get("instructions")
        if instruction_value is not None and not isinstance(instruction_value, str):
            return JSONResponse({"error": "Personal Instructions must be text."}, status_code=400)
        if isinstance(instruction_value, str) and len(instruction_value) > 3000:
            return JSONResponse({"error": "Personal Instructions must be 3,000 characters or fewer."}, status_code=400)
        enabled_value = personalization_update.get("enabled")
        if enabled_value is not None and not isinstance(enabled_value, bool):
            return JSONResponse({"error": "Personal Instructions enabled must be true or false."}, status_code=400)

    def update_settings():
        # Turning extraction off waits for any already-running provider request,
        # saves the opt-out, then drops queued messages before the API returns.
        preference_lock = (
            memory_user_lock(user_id)
            if requested_memory and (not requested_memory["enabled"] or not requested_memory["auto_extract"])
            else nullcontext()
        )
        with preference_lock:
            with _WORKSPACE_MUTATION_LOCK:
                current = get_user_settings(user_id)
                # Only settings used by the UI are accepted. Never store arbitrary payloads.
                if isinstance(current.get("chat_preferences"), dict):
                    current["chat_preferences"].pop("default_provider", None)
                allowed = {"privacy", "chat_preferences", "profile", "memory", "personalization"}
                for key in allowed:
                    value = incoming.get(key)
                    if isinstance(value, dict):
                        if key == "privacy":
                            merged_privacy = current.get("privacy") if isinstance(current.get("privacy"), dict) else {}
                            for preference in ("analytics", "training", "request_logging", "crash_reporting"):
                                if isinstance(value.get(preference), bool):
                                    merged_privacy[preference] = value[preference]
                            if "retention_days" in value:
                                merged_privacy["retention_days"] = str(value["retention_days"]).strip().lower()
                            value = merged_privacy
                        elif key == "chat_preferences":
                            value = {name: item for name, item in value.items() if name != "default_provider"}
                        elif key == "memory":
                            value = normalize_memory_settings(value)
                        elif key == "personalization":
                            existing = current.get("personalization") if isinstance(current.get("personalization"), dict) else {}
                            if isinstance(value.get("instructions"), str):
                                existing["instructions"] = value["instructions"].strip()[:3000]
                            if isinstance(value.get("enabled"), bool):
                                existing["enabled"] = value["enabled"]
                            value = {"instructions": str(existing.get("instructions") or "")[:3000], "enabled": existing.get("enabled", True) is not False}
                        current[key] = value
                save_user_settings(user_id, current)
                if requested_memory:
                    set_memory_user_disabled(user_id, not requested_memory["enabled"])
                    set_memory_user_opt_out(
                        user_id,
                        not requested_memory["enabled"] or not requested_memory["auto_extract"],
                    )
                    if not requested_memory["enabled"] or not requested_memory["auto_extract"]:
                        clear_memory_extraction_queue(user_id)
                return current

    current = await run_in_threadpool(update_settings)
    _cache_user_settings(user_id, current)
    if isinstance(privacy_update, dict) and "retention_days" in privacy_update:
        try:
            await run_in_threadpool(purge_expired_conversation_history, user_id=user_id)
        except Exception:
            logger.exception("Could not immediately apply the saved conversation retention preference")
    return JSONResponse({"success": True, "settings": current})


_SCHEDULED_TASK_LIMIT = 25
_SCHEDULED_TASK_PROMPT_LIMIT = 3000
_SCHEDULED_TASK_TITLE_LIMIT = 80


def _schedule_parse_datetime(value):
    try:
        parsed = datetime.fromisoformat(str(value or "").strip().replace("Z", "+00:00"))
    except (TypeError, ValueError):
        raise ValueError("Choose a valid date and time.")
    if parsed.tzinfo is None:
        raise ValueError("The scheduled date and time must include a timezone.")
    return parsed.astimezone(timezone.utc)


def _schedule_utc_iso(value):
    return value.astimezone(timezone.utc).isoformat(timespec="seconds").replace("+00:00", "Z")


def _schedule_local_occurrence(day, clock, zone):
    """Resolve a local wall-clock time, moving forward across DST gaps."""
    candidate = datetime.combine(day, clock).replace(tzinfo=zone)
    for _ in range(181):
        local_roundtrip = candidate.astimezone(timezone.utc).astimezone(zone)
        if local_roundtrip.replace(tzinfo=None) == candidate.replace(tzinfo=None):
            return candidate.astimezone(timezone.utc)
        candidate += timedelta(minutes=1)
    raise ValueError("That local time could not be resolved in this timezone.")


def _schedule_next_occurrence(recurrence, timezone_name, time_value, weekday=None, now=None):
    try:
        zone = ZoneInfo(str(timezone_name or "UTC").strip())
    except (ZoneInfoNotFoundError, ValueError, TypeError):
        raise ValueError("Choose a valid timezone.")
    try:
        clock = datetime.strptime(str(time_value or "").strip(), "%H:%M").time()
    except (TypeError, ValueError):
        raise ValueError("Choose a valid time.")
    local_now = (now or datetime.now(timezone.utc)).astimezone(zone)
    day = local_now.date()
    if recurrence == "weekly":
        try:
            wanted_weekday = int(weekday)
        except (TypeError, ValueError):
            raise ValueError("Choose a weekday for a weekly task.")
        if wanted_weekday < 0 or wanted_weekday > 6:
            raise ValueError("Choose a valid weekday.")
        day += timedelta(days=(wanted_weekday - day.weekday()) % 7)
    candidate = _schedule_local_occurrence(day, clock, zone)
    if candidate <= (now or datetime.now(timezone.utc)):
        day += timedelta(days=7 if recurrence == "weekly" else 1)
        candidate = _schedule_local_occurrence(day, clock, zone)
    return candidate


def _scheduled_tasks_read(user_id):
    with _WORKSPACE_MUTATION_LOCK:
        settings = get_user_settings(user_id)
        tasks = settings.get("scheduled_tasks") if isinstance(settings, dict) else []
        return [dict(task) for task in tasks if isinstance(task, dict)] if isinstance(tasks, list) else []


def _scheduled_tasks_mutate(user_id, mutation):
    with _WORKSPACE_MUTATION_LOCK:
        settings = get_user_settings(user_id)
        settings = settings if isinstance(settings, dict) else {}
        tasks = settings.get("scheduled_tasks")
        tasks = [dict(task) for task in tasks if isinstance(task, dict)] if isinstance(tasks, list) else []
        result, changed = mutation(tasks)
        if changed:
            settings["scheduled_tasks"] = tasks
            save_user_settings(user_id, settings)
            _cache_user_settings(user_id, settings)
        return result


def _scheduled_tasks_require_persistent_account(request):
    return (
        ODDI_RENDER_LAPTOP_REDIRECT
        or ODDI_BROWSER_LOCAL_CHATS
        or is_temporary_browser_account(request)
    )


def _scheduled_task_create(user_id, payload):
    title = str(payload.get("title") or "").strip()[:_SCHEDULED_TASK_TITLE_LIMIT]
    prompt = str(payload.get("prompt") or "").strip()[:_SCHEDULED_TASK_PROMPT_LIMIT]
    if not title:
        raise ValueError("Give this scheduled task a name.")
    if not prompt:
        raise ValueError("Write what ODDI should do when the task runs.")
    recurrence = str(payload.get("recurrence") or "once").strip().casefold()
    if recurrence not in {"once", "daily", "weekly"}:
        raise ValueError("Choose once, daily, or weekly.")
    timezone_name = str(payload.get("timezone") or "UTC").strip()[:80]
    try:
        ZoneInfo(timezone_name)
    except (ZoneInfoNotFoundError, ValueError, TypeError):
        raise ValueError("Choose a valid timezone.")
    now = datetime.now(timezone.utc)
    if recurrence == "once":
        next_run = _schedule_parse_datetime(payload.get("run_at"))
        if next_run <= now + timedelta(seconds=10):
            raise ValueError("Choose a future time at least 10 seconds from now.")
        time_value = next_run.astimezone(ZoneInfo(timezone_name)).strftime("%H:%M")
        weekday = None
    else:
        time_value = str(payload.get("time_of_day") or "").strip()
        weekday = None
        if recurrence == "weekly":
            try:
                weekday = int(payload.get("weekday"))
            except (TypeError, ValueError):
                raise ValueError("Choose a weekday for a weekly task.")
        next_run = _schedule_next_occurrence(recurrence, timezone_name, time_value, weekday, now)
    task = {
        "id": uuid.uuid4().hex,
        "title": title,
        "prompt": prompt,
        "recurrence": recurrence,
        "timezone": timezone_name,
        "time_of_day": time_value,
        "weekday": weekday,
        "next_run_at": _schedule_utc_iso(next_run),
        "enabled": True,
        "is_running": False,
        "last_status": "",
        "last_error": "",
        "last_run_at": "",
        "last_result_conversation_id": None,
        "created_at": _schedule_utc_iso(now),
    }

    def mutation(tasks):
        if len(tasks) >= _SCHEDULED_TASK_LIMIT:
            raise ValueError(f"You can keep up to {_SCHEDULED_TASK_LIMIT} scheduled tasks.")
        tasks.insert(0, task)
        return task, True

    return _scheduled_tasks_mutate(user_id, mutation)


def _scheduled_task_update(user_id, task_id, payload):
    enabled = payload.get("enabled")
    if not isinstance(enabled, bool):
        raise ValueError("Choose whether to pause or resume this task.")

    def mutation(tasks):
        task = next((item for item in tasks if str(item.get("id")) == str(task_id)), None)
        if not task:
            return None, False
        if task.get("is_running") and enabled:
            raise ValueError("This task is already running. Resume it after the current run finishes.")
        if enabled and task.get("recurrence") == "once" and task.get("last_status") == "completed":
            raise ValueError("This one-time task has already completed. Create a new schedule to run it again.")
        task["enabled"] = enabled
        return task, True

    return _scheduled_tasks_mutate(user_id, mutation)


def _scheduled_task_delete(user_id, task_id):
    def mutation(tasks):
        index = next((i for i, item in enumerate(tasks) if str(item.get("id")) == str(task_id)), -1)
        if index < 0:
            return False, False
        if tasks[index].get("is_running"):
            raise ValueError("This task is running now. Try deleting it after the run finishes.")
        tasks.pop(index)
        return True, True

    return bool(_scheduled_tasks_mutate(user_id, mutation))


def _claim_due_scheduled_tasks():
    now = datetime.now(timezone.utc)
    claimed = []
    for user_id in list_account_user_ids():
        user_claims = []
        try:
            with _WORKSPACE_MUTATION_LOCK:
                settings = get_user_settings(user_id)
                settings = settings if isinstance(settings, dict) else {}
                tasks = settings.get("scheduled_tasks")
                tasks = [dict(task) for task in tasks if isinstance(task, dict)] if isinstance(tasks, list) else []
                changed = False
                for task in tasks:
                    if task.get("is_running"):
                        try:
                            started = _schedule_parse_datetime(task.get("last_started_at"))
                        except ValueError:
                            started = now - timedelta(hours=1)
                        if started <= now - timedelta(minutes=30):
                            task["is_running"] = False
                            task["last_status"] = "interrupted"
                            task["last_error"] = "ODDI restarted before this run finished. The next scheduled run is still queued."
                            changed = True
                    if not task.get("enabled") or task.get("is_running"):
                        continue
                    try:
                        next_run = _schedule_parse_datetime(task.get("next_run_at"))
                    except ValueError:
                        continue
                    if next_run > now:
                        continue
                    task["is_running"] = True
                    task["last_status"] = "running"
                    task["last_started_at"] = _schedule_utc_iso(now)
                    task["last_error"] = ""
                    task["scheduled_for"] = _schedule_utc_iso(next_run)
                    if task.get("recurrence") == "once":
                        task["enabled"] = False
                    else:
                        task["next_run_at"] = _schedule_utc_iso(_schedule_next_occurrence(
                            str(task.get("recurrence")),
                            str(task.get("timezone") or "UTC"),
                            str(task.get("time_of_day") or ""),
                            task.get("weekday"),
                            now,
                        ))
                    user_claims.append({"user_id": user_id, **task})
                    changed = True
                if changed:
                    settings["scheduled_tasks"] = tasks
                    save_user_settings(user_id, settings)
                    _cache_user_settings(user_id, settings)
                    claimed.extend(user_claims)
        except Exception:
            logger.exception("Could not inspect scheduled tasks for user %s", user_id)
    return claimed


def _finish_scheduled_task(user_id, task_id, status, conversation_id=None, error=""):
    finished_at = _schedule_utc_iso(datetime.now(timezone.utc))

    def mutation(tasks):
        task = next((item for item in tasks if str(item.get("id")) == str(task_id)), None)
        if not task:
            return False, False
        task["is_running"] = False
        task["last_status"] = status
        task["last_run_at"] = finished_at
        task["last_error"] = str(error or "")[:300]
        if conversation_id is not None:
            task["last_result_conversation_id"] = conversation_id
        return True, True

    _scheduled_tasks_mutate(user_id, mutation)


async def _execute_scheduled_task(job):
    user_id = job.get("user_id")
    task_id = str(job.get("id") or "")
    conversation_id = None
    conversation_saved = False
    try:
        user = await run_in_threadpool(get_user_by_id, user_id)
        if not user:
            raise RuntimeError("ODDI could not find the task owner account.")
        # SQLite returns sqlite3.Row objects, which support indexing but not .get().
        user = dict(user)
        settings = await run_in_threadpool(get_user_settings, user_id)
        personalization = settings.get("personalization") if isinstance(settings, dict) else {}
        personal_instructions = ""
        if isinstance(personalization, dict) and personalization.get("enabled") is not False:
            personal_instructions = str(personalization.get("instructions") or "").strip()[:3000]
        memory = normalize_memory_settings(settings.get("memory") if isinstance(settings, dict) else {})
        identity = create_identity(user_id=str(user_id), email=str(user.get("email") or ""))
        reply = await run_in_threadpool(
            process_message,
            str(job.get("prompt") or ""),
            [],
            [],
            user_id=user_id,
            identity=identity,
            long_term_memory_enabled=memory["enabled"],
            personal_instructions=personal_instructions,
        )
        answer = str(reply or "").strip()
        if not answer or answer == str(PROVIDER_UNAVAILABLE_RESPONSE).strip():
            raise RuntimeError("No AI provider completed this scheduled run.")
        completed_at = _schedule_utc_iso(datetime.now(timezone.utc))
        title = f"Scheduled · {str(job.get('title') or 'ODDI task')[:80]} · {datetime.now(timezone.utc).astimezone(ZoneInfo(str(job.get('timezone') or 'UTC'))).strftime('%Y-%m-%d %H:%M')}"
        conversation_id = await run_in_threadpool(create_conversation, user_id, title)
        messages = [
            {
                "id": str(uuid.uuid4()),
                "role": "user",
                "text": str(job.get("prompt") or ""),
                "created_at": completed_at,
                "scheduled_task_id": task_id,
            },
            {
                "id": str(uuid.uuid4()),
                "role": "assistant",
                "text": answer,
                "sources": list(getattr(reply, "sources", []) or []),
                "pinned": False,
                "feedback": None,
                "background_completion": True,
                "completed_at": completed_at,
                "scheduled_task_id": task_id,
                "response_meta": {"provider": "Auto routing", "source": "Scheduled task"},
            },
        ]
        saved = await run_in_threadpool(
            update_conversation,
            conversation_id,
            user_id,
            title,
            messages,
            expected_revision=0,
        )
        if not saved.get("ok"):
            raise RuntimeError("ODDI could not save the scheduled result chat.")
        conversation_saved = True
        await run_in_threadpool(_finish_scheduled_task, user_id, task_id, "completed", conversation_id, "")
        logger.info("Scheduled task completed user=%s task=%s conversation=%s", user_id, task_id, conversation_id)
    except Exception:
        logger.exception("Scheduled task failed user=%s task=%s", user_id, task_id)
        if conversation_id is not None and not conversation_saved:
            try:
                await run_in_threadpool(delete_conversation, conversation_id, user_id)
            except Exception:
                logger.exception("Could not clean up incomplete scheduled result chat=%s", conversation_id)
        try:
            await run_in_threadpool(
                _finish_scheduled_task,
                user_id,
                task_id,
                "failed",
                conversation_id,
                "ODDI could not complete this run. Check provider availability and the task prompt.",
            )
        except Exception:
            logger.exception("Could not save scheduled task failure state user=%s task=%s", user_id, task_id)


async def _scheduled_task_worker():
    await asyncio.sleep(20)
    active_runs = set()
    semaphore = asyncio.Semaphore(2)

    async def execute_with_limit(job):
        async with semaphore:
            await _execute_scheduled_task(job)

    while True:
        try:
            due = await run_in_threadpool(_claim_due_scheduled_tasks)
            for job in due:
                active = asyncio.create_task(execute_with_limit(job))
                active_runs.add(active)
                active.add_done_callback(active_runs.discard)
        except asyncio.CancelledError:
            for active in list(active_runs):
                active.cancel()
            if active_runs:
                await asyncio.gather(*active_runs, return_exceptions=True)
            raise
        except Exception:
            logger.exception("Scheduled task worker could not scan tasks; it will retry")
        await asyncio.sleep(30)


@app.get("/api/scheduled-tasks", name="api_get_scheduled_tasks")
async def api_get_scheduled_tasks(request: Request):
    if _scheduled_tasks_require_persistent_account(request):
        return JSONResponse({"error": "Scheduled Tasks require a regular account with persistent ODDI storage."}, status_code=409)
    user_id = await require_user_id_async(request)
    tasks = await run_in_threadpool(_scheduled_tasks_read, user_id)
    return JSONResponse({"tasks": tasks})


@app.post("/api/scheduled-tasks", name="api_create_scheduled_task")
async def api_create_scheduled_task(request: Request):
    if _scheduled_tasks_require_persistent_account(request):
        return JSONResponse({"error": "Scheduled Tasks require a regular account with persistent ODDI storage."}, status_code=409)
    user_id = await require_user_id_async(request)
    payload = await safe_json(request)
    try:
        task = await run_in_threadpool(_scheduled_task_create, user_id, payload if isinstance(payload, dict) else {})
    except ValueError as error:
        return JSONResponse({"error": str(error)}, status_code=400)
    except Exception:
        logger.exception("Could not create scheduled task for user %s", user_id)
        return JSONResponse({"error": "The scheduled task could not be saved."}, status_code=500)
    return JSONResponse({"success": True, "task": task}, status_code=201)


@app.patch("/api/scheduled-tasks/{task_id}", name="api_update_scheduled_task")
async def api_update_scheduled_task(request: Request, task_id: str):
    if _scheduled_tasks_require_persistent_account(request):
        return JSONResponse({"error": "Scheduled Tasks require a regular account with persistent ODDI storage."}, status_code=409)
    user_id = await require_user_id_async(request)
    payload = await safe_json(request)
    try:
        task = await run_in_threadpool(_scheduled_task_update, user_id, task_id, payload if isinstance(payload, dict) else {})
    except ValueError as error:
        return JSONResponse({"error": str(error)}, status_code=400)
    if not task:
        return JSONResponse({"error": "Scheduled task not found."}, status_code=404)
    return JSONResponse({"success": True, "task": task})


@app.delete("/api/scheduled-tasks/{task_id}", name="api_delete_scheduled_task")
async def api_delete_scheduled_task(request: Request, task_id: str):
    if _scheduled_tasks_require_persistent_account(request):
        return JSONResponse({"error": "Scheduled Tasks require a regular account with persistent ODDI storage."}, status_code=409)
    user_id = await require_user_id_async(request)
    try:
        deleted = await run_in_threadpool(_scheduled_task_delete, user_id, task_id)
    except ValueError as error:
        return JSONResponse({"error": str(error)}, status_code=409)
    if not deleted:
        return JSONResponse({"error": "Scheduled task not found."}, status_code=404)
    return JSONResponse({"success": True})


def _workspace_operation_result(user_id, operation):
    """Apply one idempotent workspace edit to the latest account settings."""
    with _WORKSPACE_MUTATION_LOCK:
        settings = get_user_settings(user_id)
        if not isinstance(settings, dict):
            settings = {}
        workspace = settings.get("workspace")
        if not isinstance(workspace, dict):
            workspace = {"projects": [], "tasks": [], "prompts": [], "hiddenHomeCards": [], "responseStyle": ""}
        workspace = json.loads(json.dumps(workspace, ensure_ascii=False))
        for key in ("projects", "tasks", "prompts", "hiddenHomeCards"):
            if not isinstance(workspace.get(key), list):
                workspace[key] = []

        op_id = str(operation.get("id") or "").strip()
        if not op_id or len(op_id) > 128:
            return None, "Workspace operation ID is invalid.", 400
        applied_ids = settings.get("workspace_operation_ids")
        if not isinstance(applied_ids, list):
            applied_ids = []
        if op_id in applied_ids:
            return {"workspace": workspace, "revision": int(settings.get("workspace_revision") or 0)}, None, 200

        kind = str(operation.get("type") or "")
        item = operation.get("item") if isinstance(operation.get("item"), dict) else {}
        item_id = str(operation.get("item_id") or item.get("id") or "").strip()

        def upsert(collection, value):
            value_id = str(value.get("id") or "").strip()
            if not value_id or len(value_id) > 128:
                raise ValueError("Workspace item ID is invalid.")
            values = workspace[collection]
            for index, existing in enumerate(values):
                if isinstance(existing, dict) and str(existing.get("id")) == value_id:
                    if collection == "projects":
                        value = {**existing, **value}
                        # Project memory was removed from the feature; purge its legacy field on edits.
                        value.pop("context", None)
                        if not isinstance(value.get("conversationIds"), list):
                            value["conversationIds"] = existing.get("conversationIds", [])
                    values[index] = value
                    break
            else:
                if collection == "projects":
                    value.setdefault("conversationIds", [])
                    value.setdefault("type", "general")
                if collection == "tasks":
                    values.insert(0, value)
                else:
                    values.append(value)
            if collection == "tasks":
                del values[300:]
            else:
                del values[:-{"projects": 100, "prompts": 100}[collection]]

        try:
            if kind in {"project.upsert", "task.upsert", "prompt.upsert"}:
                collection = kind.split(".", 1)[0] + "s"
                limits = {"project": ("name", 80), "task": ("text", 240), "prompt": ("name", 60)}
                field, maximum = limits[kind.split(".", 1)[0]]
                value = {"id": str(item.get("id") or ""), field: str(item.get(field) or "").strip()[:maximum]}
                if kind == "project.upsert":
                    if isinstance(item.get("conversationIds"), list):
                        value["conversationIds"] = [str(x) for x in item.get("conversationIds", []) if x][:500]
                    if "type" in item:
                        project_type = str(item.get("type") or "general").strip().casefold()
                        if project_type in {"general", "job_search", "study", "build", "personal"}:
                            value["type"] = project_type
                elif kind == "task.upsert":
                    value.update({"done": bool(item.get("done")), "chatTitle": str(item.get("chatTitle") or "")[:120]})
                else:
                    value["text"] = str(item.get("text") or "")[:1200]
                if not value.get(field):
                    return None, f"{field.capitalize()} cannot be empty.", 400
                upsert(collection, value)
            elif kind in {"project.delete", "task.delete", "prompt.delete"}:
                collection = kind.split(".", 1)[0] + "s"
                if not item_id:
                    return None, "Workspace item ID is required.", 400
                workspace[collection] = [x for x in workspace[collection] if not isinstance(x, dict) or str(x.get("id")) != item_id]
            elif kind == "project.link":
                chat_id = str(operation.get("conversation_id") or "").strip()
                project = next((x for x in workspace["projects"] if isinstance(x, dict) and str(x.get("id")) == item_id), None)
                if not project or not chat_id:
                    return None, "Project or conversation could not be found.", 404
                linked = project.get("conversationIds") if isinstance(project.get("conversationIds"), list) else []
                project["conversationIds"] = list(dict.fromkeys([*map(str, linked), chat_id]))[-500:]
            elif kind == "project.unlink":
                chat_id = str(operation.get("conversation_id") or "").strip()
                project = next((x for x in workspace["projects"] if isinstance(x, dict) and str(x.get("id")) == item_id), None)
                if not project or not chat_id:
                    return None, "Project or conversation could not be found.", 404
                project["conversationIds"] = [str(x) for x in project.get("conversationIds", []) if str(x) != chat_id]
            elif kind == "home.set_card":
                card_id = str(operation.get("card_id") or "")
                hidden = operation.get("hidden")
                if not card_id.startswith("quick-") or len(card_id) > 80 or not isinstance(hidden, bool):
                    return None, "Home card preferences are invalid.", 400
                cards = list(dict.fromkeys(x for x in workspace["hiddenHomeCards"] if isinstance(x, str)))
                if hidden and card_id not in cards:
                    cards.append(card_id)
                elif not hidden:
                    cards = [x for x in cards if x != card_id]
                workspace["hiddenHomeCards"] = cards[-20:]
            elif kind == "style.set":
                style = str(operation.get("style") or "")
                if style not in {"", "concise", "step_by_step", "table", "detailed"}:
                    return None, "Answer style is invalid.", 400
                workspace["responseStyle"] = style
            else:
                return None, "Workspace operation is not supported.", 400
        except (TypeError, ValueError) as exc:
            return None, str(exc), 400

        try:
            size = len(json.dumps(workspace, ensure_ascii=False).encode("utf-8"))
        except (TypeError, ValueError):
            return None, "Workspace data must be valid JSON.", 400
        if size > 256 * 1024:
            return None, "Workspace data is too large. Remove old projects or prompts and try again.", 413

        revision = int(settings.get("workspace_revision") or 0) + 1
        settings["workspace"] = workspace
        settings["workspace_revision"] = revision
        settings["workspace_operation_ids"] = [*applied_ids, op_id][-1000:]
        save_user_settings(user_id, settings)
        return {"workspace": workspace, "revision": revision}, None, 200


@app.get("/api/workspace", name="api_get_workspace")
def api_get_workspace(request: Request):
    user_id = require_user_id(request)
    settings = get_user_settings(user_id)
    settings = settings if isinstance(settings, dict) else {}
    workspace = settings.get("workspace")
    return JSONResponse({
        "workspace": workspace if isinstance(workspace, dict) else {},
        "revision": int(settings.get("workspace_revision") or 0),
        "account_id": str(user_id),
    })


@app.post("/api/workspace/operations", name="api_apply_workspace_operation")
async def api_apply_workspace_operation(request: Request):
    user_id = await require_user_id_async(request)
    data = await safe_json(request)
    operation = data.get("operation")
    if not isinstance(operation, dict):
        return JSONResponse({"error": "Workspace operation is required."}, status_code=400)
    result, error, status = await run_in_threadpool(_workspace_operation_result, user_id, operation)
    if error:
        return JSONResponse({"error": error}, status_code=status)
    return JSONResponse({"success": True, **result})


@app.put("/api/account/profile", name="api_update_profile")
async def api_update_profile(request: Request):
    user_id = require_user_id(request)
    data = await safe_json(request)
    name = str(data.get("name") or "").strip()
    bio = str(data.get("bio") or "").strip()
    avatar = str(data.get("avatar") or "").strip()
    if not name or len(name) > 80:
        return JSONResponse({"error": "Enter a name up to 80 characters."}, status_code=400)
    if len(bio) > 500:
        return JSONResponse({"error": "Bio must be 500 characters or less."}, status_code=400)
    allowed_avatar_prefixes = (
        "data:image/png;base64,", "data:image/jpeg;base64,",
        "data:image/gif;base64,", "data:image/webp;base64,",
    )
    if avatar and (len(avatar) > 2_800_000 or not avatar.startswith(allowed_avatar_prefixes)):
        return JSONResponse({"error": "Choose an image smaller than 2 MB."}, status_code=400)
    with _WORKSPACE_MUTATION_LOCK:
        update_user_profile(user_id, name)
        request.session["username"] = name
        settings = get_user_settings(user_id)
        profile = settings.get("profile", {})
        profile.update({"bio": bio, "avatar": avatar})
        settings["profile"] = profile
        save_user_settings(user_id, settings)
    return JSONResponse({"success": True, "user": {"name": name, "email": request.session.get("email", "")}, "profile": profile})


@app.get("/api/account/export", name="api_export_account_data")
def api_export_account_data(request: Request):
    user_id = require_user_id(request)
    payload = {
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "user": {"id": str(user_id), "name": request.session.get("username"), "email": request.session.get("email")},
        "conversations": get_conversations(user_id),
        "deleted_conversations": get_deleted_conversations(user_id),
        "memories": get_memory(user_id, sync_drive=ODDI_LOCAL_ONLY_STORAGE) or {},
        "files": get_files(user_id),
        "settings": get_user_settings(user_id),
    }
    return JSONResponse(
        content=jsonable_encoder(payload),
        headers={"Content-Disposition": 'attachment; filename="oddi-data-export.json"'},
    )


@app.delete("/api/account", name="api_delete_account")
async def api_delete_account(request: Request):
    user_id = require_user_id(request)
    email = str(request.session.get("email") or "")
    data = await safe_json(request)
    if not email or str(data.get("email") or "").strip().casefold() != email.strip().casefold():
        return JSONResponse({"error": "Enter the exact email address on your account to confirm deletion."}, status_code=400)

    # Remove the user's remote document index before erasing its identifier.
    vector_store_id = get_vector_store_id(user_id) if not ODDI_LOCAL_ONLY_STORAGE else None
    if vector_store_id:
        try:
            client.vector_stores.delete(vector_store_id)
        except Exception:
            logger.exception("Could not delete account vector store for user %s", user_id)
            return JSONResponse({"error": "Your remote document index could not be deleted; the account was left intact."}, status_code=502)

    for file_record in get_files(user_id):
        if file_record.get("storage_backend") == "filesystem" and file_record.get("storage_key"):
            try:
                storage_router.delete_file(user_id, file_record["storage_key"])
            except Exception:
                logger.exception("Could not delete account file %s", file_record.get("id"))
                return JSONResponse({"error": "An uploaded file could not be deleted; the account was left intact."}, status_code=502)
        delete_file_record(file_record.get("id"), user_id)

    delete_all_conversations(user_id)
    clear_memory(user_id)
    try:
        cleanup_shared_workspaces_for_account(int(user_id))
    except (TypeError, ValueError, OverflowError):
        logger.warning("Could not resolve shared workspace account ID during account deletion")
    delete_user_record(user_id)
    request.session.clear()
    return JSONResponse({"success": True, "redirect": "/login"})


@app.post("/api/auth/logout", name="api_logout")
def api_logout(request: Request):
    request.session.clear()
    return JSONResponse({"success": True})


# =========================================================
# CONVERSATIONS
# =========================================================

@app.get("/api/conversations", name="api_get_conversations")
def api_get_conversations(request: Request):
    # Looking up a Drive-backed privacy setting for each periodic history
    # refresh adds several extra Drive round trips. Keep those refreshes quiet.
    if not ODDI_DRIVE_CHAT_STORAGE and _request_logging_enabled(request):
        logger.info(
            "HTTP GET /api/conversations pid=%s user=%s",
            os.getpid(),
            request.session.get("user_id"),
        )
    user_id = require_user_id(request)
    return JSONResponse(content=jsonable_encoder(get_conversations(user_id)))


@app.get("/api/conversations/index", name="api_get_conversation_index")
def api_get_conversation_index(request: Request, background_tasks: BackgroundTasks):
    """Load the chat sidebar without downloading every chat's full history."""
    user_id = require_user_id(request)
    summaries = get_conversation_summaries(user_id)
    needs_backfill = ODDI_DRIVE_CHAT_STORAGE and any(item.get("__summary_missing") for item in summaries)
    schedule_backfill = False
    if needs_backfill:
        with _SUMMARY_BACKFILL_LOCK:
            key = str(user_id)
            if key not in _SUMMARY_BACKFILL_USERS:
                _SUMMARY_BACKFILL_USERS.add(key)
                schedule_backfill = True
    if schedule_backfill:
        # Existing Drive files predate the compact metadata index. Resolve up
        # to 40 titles after sending the fast response; new writes carry these
        # properties immediately. Remaining files are picked up on next visit.
        background_tasks.add_task(_backfill_chat_summaries, user_id)
    return JSONResponse(content=jsonable_encoder(summaries))


@app.post("/api/conversations", name="api_create_conversation")
async def api_create_conversation(request: Request):
    if not ODDI_BROWSER_LOCAL_CHATS and await request_logging_enabled_async(request):
        logger.info(
            "HTTP POST /api/conversations pid=%s user=%s",
            os.getpid(),
            request.session.get("user_id"),
        )
    user_id = await require_user_id_async(request)
    data = await safe_json(request)
    title = data.get("title", "New Chat")

    conversation_id = await run_in_threadpool(create_conversation, user_id, title)
    if ODDI_DRIVE_CHAT_STORAGE:
        # The newly written chat is already known to be empty. Reading it back
        # from Drive adds extra list/download requests to the send request.
        created_at = datetime.now(timezone.utc).isoformat()
        conversation = {
            "id": conversation_id,
            "title": str(title or "New Chat").strip()[:500] or "New Chat",
            "messages": [],
            "revision": 0,
            "created_at": created_at,
            "updated_at": created_at,
        }
    else:
        conversation = await run_in_threadpool(get_conversation, conversation_id, user_id)

    return JSONResponse(
        content=jsonable_encoder({
            "success": True,
            **(
                conversation
                or {
                    "id": conversation_id,
                    "title": title,
                    "messages": [],
                    "revision": 0,
                }
            ),
        }),
        status_code=201,
    )


@app.post("/api/conversations/import-local", name="api_import_local_conversation")
async def api_import_local_conversation(request: Request):
    user_id = await require_user_id_async(request)
    data = await safe_json(request)
    source_id = str(data.get("source_id") or "").strip()
    conversation = data.get("conversation")
    if not isinstance(conversation, dict):
        return JSONResponse({"error": "Conversation data is required."}, status_code=400)
    if not isinstance(conversation.get("messages", []), list):
        return JSONResponse({"error": "Conversation messages must be a list."}, status_code=400)
    try:
        payload_size = len(json.dumps(conversation, ensure_ascii=False).encode("utf-8"))
    except (TypeError, ValueError):
        return JSONResponse({"error": "Conversation data is not valid JSON."}, status_code=400)
    if payload_size > 6 * 1024 * 1024:
        return JSONResponse({"error": "This chat is too large to sync in one request."}, status_code=413)
    try:
        imported = await run_in_threadpool(
            import_local_conversation, user_id, source_id, conversation
        )
    except ValueError as error:
        return JSONResponse({"error": str(error)}, status_code=400)
    if not imported:
        return JSONResponse({"error": "The chat could not be imported."}, status_code=500)
    return JSONResponse(
        {"success": True, "conversation": jsonable_encoder(imported)},
        status_code=200,
    )


@app.get("/api/conversations/bin", name="api_get_deleted_conversations")
def api_get_deleted_conversations(request: Request):
    user_id = require_user_id(request)
    return JSONResponse(content=jsonable_encoder(get_deleted_conversation_summaries(user_id)))


@app.put("/api/conversations/{conversation_id}/metadata", name="api_update_conversation_metadata")
async def api_update_conversation_metadata(request: Request, conversation_id: int):
    if await request_logging_enabled_async(request):
        logger.info(
            "HTTP PUT /api/conversations/%s/metadata pid=%s user=%s",
            conversation_id,
            os.getpid(),
            request.session.get("user_id"),
        )
    user_id = await require_user_id_async(request)
    data = await safe_json(request)

    allowed = {"pinned", "archived", "deleted", "title"}
    patch = {key: data[key] for key in allowed if key in data}

    if not patch:
        return JSONResponse(
            {"error": "No metadata changes supplied."},
            status_code=400,
        )

    updated = await run_in_threadpool(
        update_conversation_metadata,
        conversation_id,
        user_id,
        pinned=patch.get("pinned"),
        archived=patch.get("archived"),
        deleted=patch.get("deleted"),
        title=patch.get("title"),
    )

    if not updated:
        return JSONResponse({"error": "Conversation not found."}, status_code=404)

    metadata_fields = (
        "id", "user_id", "title", "created_at", "updated_at", "revision",
        "pinned", "archived", "deleted", "deleted_at",
    )
    metadata = {key: updated[key] for key in metadata_fields if key in updated}
    return JSONResponse(content=jsonable_encoder({"success": True, "conversation": metadata}))


@app.get("/api/conversations/{conversation_id}", name="api_get_conversation")
def api_get_conversation(request: Request, conversation_id: int):
    user_id = require_user_id(request)
    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        return JSONResponse({"error": "Conversation not found."}, status_code=404)

    return JSONResponse(content=jsonable_encoder(conversation))


@app.post("/api/conversations/{conversation_id}/branch", name="api_branch_conversation")
async def api_branch_conversation(request: Request, conversation_id: int):
    user_id = await require_user_id_async(request)
    payload = await safe_json(request)
    message_id = str(payload.get("message_id") or "").strip() if isinstance(payload, dict) else ""
    mode = str(payload.get("mode") or "message").strip().lower() if isinstance(payload, dict) else "message"
    if not message_id or len(message_id) > 128:
        return JSONResponse({"error": "Choose a message to branch from."}, status_code=400)
    if mode not in {"full", "message"}:
        return JSONResponse({"error": "Choose either the full chat or this message."}, status_code=400)
    source = await run_in_threadpool(get_conversation, conversation_id, user_id)
    if not source or source.get("deleted"):
        return JSONResponse({"error": "Conversation not found."}, status_code=404)
    messages = source.get("messages") if isinstance(source.get("messages"), list) else []
    branch_index = next(
        (index for index, item in enumerate(messages)
         if isinstance(item, dict) and str(item.get("id") or "") == message_id),
        -1,
    )
    if branch_index < 0:
        return JSONResponse({"error": "That message is no longer in this conversation."}, status_code=404)
    if str(messages[branch_index].get("role") or "") != "assistant":
        return JSONResponse({"error": "Choose an ODDI response to branch from."}, status_code=400)

    source_title = str(source.get("title") or "Conversation").strip()[:450] or "Conversation"
    title = f"Branch · {source_title}"[:500]
    branch_messages = []
    messages_to_copy = messages if mode == "full" else [messages[branch_index]]
    for message in messages_to_copy:
        if not isinstance(message, dict) or message.get("role") not in {"user", "assistant"}:
            continue
        copied = copy.deepcopy(message)
        copied["id"] = str(uuid.uuid4())
        branch_messages.append(copied)
    if not branch_messages:
        return JSONResponse({"error": "There are no messages available to branch."}, status_code=400)
    branch_messages[-1]["branch_origin"] = {
        "mode": mode,
        "source_conversation_id": str(conversation_id),
        "source_title": source_title[:500],
        "source_message_id": message_id,
        "inherited_message_count": len(branch_messages),
    }
    try:
        branch_id = await run_in_threadpool(create_conversation, user_id, title)
        saved = await run_in_threadpool(
            update_conversation,
            branch_id,
            user_id,
            title,
            branch_messages,
            expected_revision=0,
        )
        if not saved.get("ok") or not saved.get("conversation"):
            raise RuntimeError("The new branch conversation could not be saved.")
    except Exception:
        logger.exception("Could not branch conversation=%s user=%s", conversation_id, user_id)
        if "branch_id" in locals():
            try:
                await run_in_threadpool(delete_conversation, branch_id, user_id)
            except Exception:
                logger.exception("Could not clean up incomplete branch conversation=%s", branch_id)
        return JSONResponse({"error": "ODDI could not create this branch. Please try again."}, status_code=500)
    return JSONResponse(content=jsonable_encoder({
        "success": True,
        "conversation": saved["conversation"],
        "branched_from": {
            "conversation_id": conversation_id,
            "conversation_title": source_title,
            "message_id": message_id,
            "mode": mode,
        },
    }), status_code=201)


@app.put("/api/conversations/{conversation_id}", name="api_update_conversation")
async def api_update_conversation(request: Request, conversation_id: int):
    if await request_logging_enabled_async(request):
        logger.info(
            "HTTP PUT /api/conversations/%s pid=%s user=%s",
            conversation_id,
            os.getpid(),
            request.session.get("user_id"),
        )
    user_id = await require_user_id_async(request)
    data = await safe_json(request)

    title = data.get("title", "New Chat")
    messages = data.get("messages", [])
    expected_revision = data.get("expected_revision")

    result = await run_in_threadpool(
        update_conversation,
        conversation_id,
        user_id,
        title,
        messages,
        expected_revision=expected_revision,
    )

    if not result.get("found"):
        return JSONResponse({"error": "Conversation not found."}, status_code=404)

    if result.get("conflict"):
        return JSONResponse(
            {
                "success": False,
                "conflict": True,
                "conversation": result["conversation"],
            },
            status_code=409,
        )

    return JSONResponse(
        content=jsonable_encoder({
            "success": True,
            "conversation": result["conversation"],
        })
    )


@app.post("/api/conversations/{conversation_id}/messages", name="api_append_conversation_message")
async def api_append_conversation_message(request: Request, conversation_id: int):
    if await request_logging_enabled_async(request):
        logger.info(
            "HTTP POST /api/conversations/%s/messages pid=%s user=%s",
            conversation_id,
            os.getpid(),
            request.session.get("user_id"),
        )
    user_id = await require_user_id_async(request)
    data = await safe_json(request)
    message = data.get("message")

    if not isinstance(message, dict):
        return JSONResponse(
            {"error": "Message object is required."},
            status_code=400,
        )

    conversation = await run_in_threadpool(
        append_conversation_message,
        conversation_id,
        user_id,
        message,
    )

    if not conversation:
        return JSONResponse({"error": "Conversation not found."}, status_code=404)

    return JSONResponse(content=jsonable_encoder({"success": True, "conversation": conversation}))


@app.put("/api/conversations/{conversation_id}/messages/{message_id}", name="api_update_conversation_message")
async def api_update_conversation_message(
    request: Request,
    conversation_id: int,
    message_id: str,
):
    if await request_logging_enabled_async(request):
        logger.info(
            "HTTP PUT /api/conversations/%s/messages/%s pid=%s user=%s",
            conversation_id,
            message_id,
            os.getpid(),
            request.session.get("user_id"),
        )
    user_id = await require_user_id_async(request)
    data = await safe_json(request)
    patch = data.get("patch", data)

    if not isinstance(patch, dict):
        return JSONResponse(
            {"error": "Message patch must be an object."},
            status_code=400,
        )

    conversation = await run_in_threadpool(
        update_conversation_message,
        conversation_id,
        user_id,
        message_id,
        patch,
    )

    if not conversation:
        return JSONResponse(
            {"error": "Conversation or message not found."},
            status_code=404,
        )

    return JSONResponse(content=jsonable_encoder({"success": True, "conversation": conversation}))


@app.delete("/api/conversations/{conversation_id}/messages/{message_id}", name="api_delete_conversation_message")
def api_delete_conversation_message(
    request: Request,
    conversation_id: int,
    message_id: str,
):
    if _request_logging_enabled(request):
        logger.warning(
            "HTTP DELETE /api/conversations/%s/messages/%s pid=%s user=%s",
            conversation_id,
            message_id,
            os.getpid(),
            request.session.get("user_id"),
        )
    user_id = require_user_id(request)

    conversation = delete_conversation_message(
        conversation_id,
        user_id,
        message_id,
    )

    if not conversation:
        return JSONResponse(
            {"error": "Conversation or message not found."},
            status_code=404,
        )

    return JSONResponse(content=jsonable_encoder({"success": True, "conversation": conversation}))


@app.delete("/api/conversations/clear", name="api_clear_conversations")
def api_clear_conversations(request: Request):
    if _request_logging_enabled(request):
        logger.warning(
            "HTTP DELETE /api/conversations/clear pid=%s user=%s",
            os.getpid(),
            request.session.get("user_id"),
        )
    user_id = require_user_id(request)
    delete_all_conversations(user_id)
    return JSONResponse({"success": True})


@app.delete("/api/conversations/{conversation_id}", name="api_delete_conversation")
def api_delete_conversation(request: Request, conversation_id: int):
    if _request_logging_enabled(request):
        logger.warning(
            "HTTP DELETE /api/conversations/%s pid=%s user=%s",
            conversation_id,
            os.getpid(),
            request.session.get("user_id"),
        )
    user_id = require_user_id(request)

    conversation = get_conversation(conversation_id, user_id)

    if not conversation:
        return JSONResponse({"error": "Conversation not found."}, status_code=404)

    delete_conversation(conversation_id, user_id)
    return JSONResponse({"success": True})


# =========================================================
# MEMORY CONTROLS
# =========================================================

MEMORY_CATEGORY_LABELS = {
    "goal": "Goals",
    "project": "Projects",
    "preference": "Preferences",
    "interest": "Interests",
    "person": "People",
    "skill": "Learning",
}


def _memory_category_label(category, all_categories):
    label = MEMORY_CATEGORY_LABELS.get(str(category).strip().casefold())
    if not label:
        return category

    identity = label.casefold()
    for other_category in all_categories:
        if other_category == category:
            continue
        other = str(other_category).strip().casefold()
        other_label = MEMORY_CATEGORY_LABELS.get(other, str(other_category)).casefold()
        if other == identity or other_label == identity:
            return category

    return label


@app.get("/api/memory", name="api_get_memory")
def api_get_memory(request: Request):
    user_id = require_user_id(request)
    memories = get_memory(user_id, sync_drive=ODDI_LOCAL_ONLY_STORAGE) or {}
    try:
        memory_settings = normalize_memory_settings(get_user_settings(user_id).get("memory"))
    except Exception as error:
        logger.warning("Memory preference lookup failed while opening Memory (%s).", type(error).__name__)
        memory_settings = {"enabled": False, "auto_extract": False}

    try:
        queue_status = get_memory_extraction_status(user_id)
    except Exception as error:
        logger.warning("Memory queue status lookup failed (%s).", type(error).__name__)
        queue_status = {
            "queued_count": 0,
            "processing_count": 0,
            "next_available_at": None,
            "last_activity": None,
            "lookup_failed": True,
        }
    provider_configured = memory_provider_keys_configured()
    memory_available = (
        not ODDI_BROWSER_LOCAL_CHATS
        and memory_settings.get("enabled") is True
        and memory_settings.get("auto_extract") is True
        and provider_configured
    )

    return JSONResponse(
        {
            "memory_settings": memory_settings,
            "auto_memory": {
                "available": memory_available,
                "batch_size": MEMORY_EXTRACTION_BATCH_SIZE,
                "browser_local_chats": ODDI_BROWSER_LOCAL_CHATS,
                "provider_configured": provider_configured,
                "queue": queue_status,
            },
            "memories": [
                {
                    "key": key,
                    "category": _memory_category_label(key, memories),
                    "memory": value,
                }
                for key, value in memories.items()
            ]
        }
    )


@app.put("/api/memory", name="api_update_memory")
async def api_update_memory(request: Request):
    user_id = require_user_id(request)
    data = await safe_json(request)

    key = str(data.get("key") or data.get("category") or "").strip()
    memory = str(data.get("memory") or data.get("value") or "").strip()

    if not key:
        return JSONResponse({"error": "Memory key is required."}, status_code=400)
    if not memory:
        return JSONResponse({"error": "Memory cannot be empty."}, status_code=400)

    def save_manual_memory():
        with memory_user_lock(user_id):
            save_memory(user_id, key, memory)

    await run_in_threadpool(save_manual_memory)

    return JSONResponse(
        {
            "success": True,
            "category": key,
            "memory": memory,
        }
    )


@app.delete("/api/memory", name="api_delete_memory")
async def api_delete_memory(request: Request):
    user_id = require_user_id(request)
    data = await safe_json(request)

    key = str(data.get("key") or data.get("category") or "").strip()

    if not key:
        return JSONResponse({"error": "Memory key is required."}, status_code=400)

    def delete_manual_memory():
        with memory_user_lock(user_id):
            existing = get_memory(user_id, key)
            if existing is None:
                return False
            delete_memory(user_id, key)
            return True

    deleted = await run_in_threadpool(delete_manual_memory)
    if not deleted:
        return JSONResponse({"error": "Memory not found."}, status_code=404)
    return JSONResponse({"success": True, "deleted": key})


# =========================================================
# FILES
# =========================================================

_PREVIEW_MAX_UPLOAD_BYTES = 20 * 1024 * 1024
_PREVIEW_TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".log", ".json", ".xml", ".html", ".htm",
    ".yaml", ".yml", ".py", ".js", ".ts", ".tsx", ".jsx", ".css", ".sql",
    ".ini", ".cfg", ".toml", ".rtf", ".csv", ".tsv",
}


def _preview_cell_value(value):
    if value is None:
        return ""
    if isinstance(value, (str, int, float, bool)):
        return str(value)[:1000]
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat()
        except (TypeError, ValueError):
            pass
    return str(value)[:1000]


def _extract_upload_preview(filename, data, content_type=""):
    """Read common office and text formats for the transient file-preview UI."""
    import csv
    import io
    import os
    import re
    import zipfile

    extension = os.path.splitext(filename)[1].casefold()
    office_extensions = {
        ".docx", ".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".pptx",
    }

    if extension in office_extensions - {".xls"}:
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                expanded_bytes = sum(member.file_size for member in members)
                if len(members) > 5000 or expanded_bytes > 120 * 1024 * 1024:
                    raise ValueError("This office file is too large to preview safely.")
        except zipfile.BadZipFile as exc:
            raise ValueError("This office file is damaged or has an unsupported format.") from exc

    if extension == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        blocks = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = str(getattr(paragraph.style, "name", ""))
            if style_name.lower().startswith("heading"):
                level = re.search(r"(\d+)$", style_name)
                text = "#" * min(6, int(level.group(1)) if level else 2) + " " + text
            blocks.append(text)
        for table_index, table in enumerate(document.tables, start=1):
            rows = []
            for row_index, row in enumerate(table.rows):
                if row_index >= 300:
                    break
                rows.append("\t".join(cell.text.replace("\n", " ").strip()[:500] for cell in row.cells[:30]))
            if rows:
                blocks.append(f"Table {table_index}\n" + "\n".join(rows))
        content = "\n\n".join(blocks) or "This Word document has no readable text."
        truncated = len(content) > 1_000_000
        return {"kind": "text", "content": content[:1_000_000], "truncated": truncated}

    if extension in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheets = []
        try:
            for worksheet in workbook.worksheets[:12]:
                max_row = min(max(worksheet.max_row or 1, 1), 500)
                max_column = min(max(worksheet.max_column or 1, 1), 30)
                rows = [
                    [_preview_cell_value(value) for value in row]
                    for row in worksheet.iter_rows(
                        max_row=max_row,
                        max_col=max_column,
                        values_only=True,
                    )
                ]
                sheets.append({
                    "name": worksheet.title,
                    "rows": rows,
                    "truncated": (worksheet.max_row or 0) > 500 or (worksheet.max_column or 0) > 30,
                })
        finally:
            workbook.close()
        return {"kind": "spreadsheet", "sheets": sheets}

    if extension == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        sheets = []
        try:
            for worksheet in workbook.sheets()[:12]:
                row_limit = min(worksheet.nrows, 500)
                column_limit = min(worksheet.ncols, 30)
                rows = [
                    [_preview_cell_value(value) for value in worksheet.row_values(row_index, 0, column_limit)]
                    for row_index in range(row_limit)
                ]
                sheets.append({
                    "name": worksheet.name,
                    "rows": rows,
                    "truncated": worksheet.nrows > 500 or worksheet.ncols > 30,
                })
        finally:
            workbook.release_resources()
        return {"kind": "spreadsheet", "sheets": sheets}

    if extension == ".pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(data))
        slides = []
        for number, slide in enumerate(presentation.slides, start=1):
            if number > 100:
                break
            lines = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = "\n".join(paragraph.text for paragraph in shape.text_frame.paragraphs).strip()
                    if text:
                        lines.append(text)
                if getattr(shape, "has_table", False):
                    for row_index, row in enumerate(shape.table.rows):
                        if row_index >= 100:
                            break
                        lines.append("\t".join(cell.text.strip() for cell in row.cells[:30]))
            slides.append(f"Slide {number}\n" + "\n".join(lines))
        content = "\n\n".join(slides) or "This presentation has no readable slide text."
        truncated = len(content) > 1_000_000 or len(presentation.slides) > 100
        return {"kind": "text", "content": content[:1_000_000], "truncated": truncated}

    if extension in _PREVIEW_TEXT_EXTENSIONS or str(content_type).casefold().startswith("text/"):
        text = data[:2 * 1024 * 1024].decode("utf-8-sig", errors="replace")
        truncated = len(data) > 2 * 1024 * 1024
        if extension in {".csv", ".tsv"}:
            delimiter = "\t" if extension == ".tsv" else ","
            if extension == ".csv":
                try:
                    delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
                except csv.Error:
                    pass
            rows = []
            for row in csv.reader(io.StringIO(text), delimiter=delimiter):
                rows.append([str(cell)[:1000] for cell in row[:30]])
                if len(rows) >= 500:
                    truncated = True
                    break
            return {"kind": "spreadsheet", "sheets": [{"name": filename, "rows": rows, "truncated": truncated}]}
        if extension == ".rtf":
            text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
            text = re.sub(r"\\(?:par|line)\b ?", "\n", text)
            text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text).replace("{", "").replace("}", "")
        return {"kind": "text", "content": text[:1_000_000], "truncated": truncated or len(text) > 1_000_000}

    raise ValueError("This file type cannot be previewed. It can still be attached to your message.")


@app.post("/api/preview-file", name="api_preview_file")
async def api_preview_file(request: Request):
    require_user_id(request)
    try:
        content_length = int(request.headers.get("content-length") or 0)
    except ValueError:
        content_length = 0
    if content_length > _PREVIEW_MAX_UPLOAD_BYTES + 256 * 1024:
        return JSONResponse({"error": "Preview is limited to files up to 20 MB."}, status_code=413)
    try:
        form = await request.form()
    except Exception:
        return JSONResponse({"error": "The preview upload could not be read."}, status_code=400)
    upload = form.get("file")
    if upload is None or not hasattr(upload, "read"):
        return JSONResponse({"error": "Choose a file to preview."}, status_code=400)

    filename = os.path.basename(str(getattr(upload, "filename", "") or "attachment"))
    content_type = str(getattr(upload, "content_type", "") or "")
    try:
        data = await upload.read(_PREVIEW_MAX_UPLOAD_BYTES + 1)
        if len(data) > _PREVIEW_MAX_UPLOAD_BYTES:
            return JSONResponse({"error": "Preview is limited to files up to 20 MB."}, status_code=413)
        preview = await run_in_threadpool(_extract_upload_preview, filename, data, content_type)
        return JSONResponse(preview)
    except ValueError as exc:
        return JSONResponse({"error": str(exc)}, status_code=415)
    except ImportError as exc:
        logger.exception("File preview dependency is missing for %s", filename)
        return JSONResponse({"error": f"Preview support for this format is unavailable: {exc.name or 'missing reader'}"}, status_code=503)
    except Exception:
        logger.exception("File preview failed for %s", filename)
        return JSONResponse({"error": "This file could not be previewed. It can still be attached to your message."}, status_code=422)
    finally:
        await upload.close()

@app.get("/api/files", name="api_get_files")
def api_get_files(request: Request):
    user_id = require_user_id(request)
    raw_conversation_id = request.query_params.get("conversation_id")

    try:
        conversation_id = (
            int(raw_conversation_id)
            if raw_conversation_id
            else None
        )
    except (TypeError, ValueError):
        conversation_id = None

    return JSONResponse(
        content=jsonable_encoder(
            get_files(
                user_id,
                conversation_id=conversation_id,
            )
        )
    )


@app.post("/api/files/import", name="api_import_library_file")
async def api_import_library_file(request: Request):
    user_id = await require_user_id_async(request)
    if is_temporary_browser_account(request):
        return JSONResponse(
            {"error": "Sign in to an account to sync Library files across devices."},
            status_code=403,
        )

    try:
        form = await request.form()
    except Exception:
        return JSONResponse({"error": "The uploaded file could not be read."}, status_code=400)
    upload = form.get("file")
    if not hasattr(upload, "file") or not hasattr(upload, "filename"):
        return JSONResponse({"error": "Choose a file to add to your Library."}, status_code=400)

    adapted_upload = adapt_uploaded_files([upload])[0]
    adapted_upload.filename = os.path.basename(str(adapted_upload.filename or "").replace("\\", "/")) or "uploaded-file"
    file_bytes = await upload.read(MAX_UPLOAD_BYTES + 1)
    if len(file_bytes) > MAX_UPLOAD_BYTES:
        await upload.close()
        return JSONResponse(
            {"error": "This file is larger than the 64 MB upload limit."},
            status_code=413,
        )
    await upload.close()

    plan = await run_in_threadpool(
        _plan_chat_library_uploads,
        user_id,
        [(adapted_upload, file_bytes)],
    )
    existing_file_ids = [
        int(entry["file_id"])
        for entry in plan.get("already_existing_files") or []
        if entry.get("file_id") is not None
    ]
    if not plan.get("uploads"):
        existing_record = (
            await run_in_threadpool(get_file, existing_file_ids[0], user_id)
            if existing_file_ids else None
        )
        return JSONResponse({
            "success": True,
            "file_ids": existing_file_ids,
            "filename": (existing_record or {}).get("filename"),
            "added_count": 0,
        })

    try:
        current_count = await run_in_threadpool(count_files, user_id)
        replacement_count = len(plan.get("replace_file_ids") or [])
        planned_count = len(plan.get("uploads") or [])
        if current_count - replacement_count + planned_count > USER_FILE_LIMIT:
            return JSONResponse(
                {"error": "Your Library is full (30 files). Delete an older file from Library before uploading more."},
                status_code=413,
            )

        planned_bytes = sum(len(entry.get("file_bytes") or b"") for entry in plan["uploads"])
        if planned_bytes:
            quota = await run_in_threadpool(storage_router.get_file_quota, user_id)
            if planned_bytes > int(quota.get("remaining_bytes", 0) or 0) + int(plan.get("quota_credit_bytes") or 0):
                return JSONResponse(
                    {"error": "File storage quota exceeded.", "storage": quota},
                    status_code=413,
                )

        file_ids = await run_in_threadpool(
            _persist_chat_library_uploads,
            user_id,
            None,
            plan,
        )
    except FileLimitExceeded as error:
        return JSONResponse({"error": str(error)}, status_code=413)
    except StorageQuotaExceeded as error:
        logger.info("Local Library file import exceeded quota for user %s: %s", user_id, error)
        return JSONResponse(
            {"error": "File storage quota exceeded.", "storage": storage_router.get_file_quota(user_id)},
            status_code=413,
        )
    except Exception:
        logger.exception("Local Library file import failed for user %s", user_id)
        return JSONResponse({"error": "This file could not be synced to your Library."}, status_code=500)

    synced_file_ids = [int(file_id) for file_id in file_ids] + existing_file_ids
    synced_record = (
        await run_in_threadpool(get_file, synced_file_ids[0], user_id)
        if synced_file_ids else None
    )
    return JSONResponse({
        "success": True,
        "file_ids": synced_file_ids,
        "filename": (synced_record or {}).get("filename"),
        "added_count": len(file_ids),
    })


@app.get("/api/storage", name="api_storage_info")
def api_storage_info(request: Request):
    """
    Return user-visible storage information.

    Only persistent user file storage is exposed here.
    ODDI short-term memory is intentionally not exposed.
    """
    user_id = require_user_id(request)

    return JSONResponse(
        content=jsonable_encoder(
            storage_router.get_user_storage_info(user_id)
        )
    )


@app.get("/api/storage/sync-status", name="api_storage_sync_status")
async def api_storage_sync_status(request: Request):
    user_id = await require_user_id_async(request)
    payload = await run_in_threadpool(_storage_sync_status_payload, user_id)
    return JSONResponse(content=jsonable_encoder(payload))


@app.post("/api/storage/sync-now", name="api_storage_sync_now")
async def api_storage_sync_now(request: Request):
    user_id = await require_user_id_async(request)
    if not ODDI_LOCAL_ONLY_STORAGE:
        return JSONResponse(
            {"success": False, "error": "Manual laptop sync is available in the laptop app."},
            status_code=409,
        )

    from app import google_drive_storage
    if not google_drive_storage.is_configured():
        return JSONResponse(
            {"success": False, "error": "Google Drive sync is not configured on this laptop."},
            status_code=503,
        )

    try:
        await run_in_threadpool(_run_laptop_drive_sync)
    except DriveStorageError:
        logger.warning("Manual Google Drive sync failed for account %s", user_id)
        status = await run_in_threadpool(_storage_sync_status_payload, user_id)
        return JSONResponse({"success": False, "sync": status}, status_code=503)
    except Exception:
        logger.exception("Manual Google Drive sync failed for account %s", user_id)
        status = await run_in_threadpool(_storage_sync_status_payload, user_id)
        return JSONResponse({"success": False, "sync": status}, status_code=503)

    status = await run_in_threadpool(_storage_sync_status_payload, user_id)
    return JSONResponse({"success": True, "sync": status})


@app.get("/api/files/{file_id}", name="api_download_file")
def api_download_file(request: Request, file_id: int):
    user_id = require_user_id(request)
    record = get_file(file_id, user_id)

    if not record:
        return JSONResponse(
            {"error": "File not found."},
            status_code=404,
        )

    filename = record.get("filename") or "download"
    media_type = record.get("mime_type") or "application/octet-stream"
    storage_backend = record.get("storage_backend") or "database"

    # New files live in the filesystem storage layer.
    if storage_backend == "filesystem":
        storage_key = record.get("storage_key")

        if not storage_key:
            return JSONResponse(
                {"error": "File storage reference is missing."},
                status_code=500,
            )

        try:
            path = storage_router.get_file_path(
                user_id,
                storage_key,
            )
        except StorageFileNotFound:
            return JSONResponse(
                {"error": "Stored file content is missing."},
                status_code=404,
            )

        response = FileResponse(
            path=str(path),
            media_type=media_type,
            filename=filename,
        )

        response.headers["Content-Disposition"] = (
            f"inline; filename*=UTF-8''{quote(filename)}"
        )

        return response

    # Backward compatibility for legacy database-backed files.
    data = record.get("file_data")

    if data is None:
        return JSONResponse(
            {"error": "File content is not stored for this record."},
            status_code=404,
        )

    response = StreamingResponse(
        io.BytesIO(bytes(data)),
        media_type=media_type,
    )
    response.headers["Content-Disposition"] = (
        f"inline; filename*=UTF-8''{quote(filename)}"
    )

    return response


@app.delete("/api/files/{file_id}", name="api_delete_file")
def api_delete_file(request: Request, file_id: int):
    user_id = require_user_id(request)
    record = get_file(file_id, user_id)

    if not record:
        return JSONResponse(
            {"error": "File not found."},
            status_code=404,
        )

    storage_backend = record.get("storage_backend") or "database"

    if storage_backend == "filesystem":
        storage_key = record.get("storage_key")

        if storage_key:
            try:
                storage_router.delete_file(
                    user_id,
                    storage_key,
                )
            except Exception as storage_error:
                logger.exception(
                    "Physical file deletion failed for file %s: %s",
                    file_id,
                    storage_error,
                )
                return JSONResponse(
                    {"error": "File content could not be deleted."},
                    status_code=500,
                )

    deleted = delete_file_record(file_id, user_id)

    if not deleted:
        return JSONResponse(
            {"error": "File not found."},
            status_code=404,
        )

    return JSONResponse(
        {
            "success": True,
            "deleted": file_id,
        }
    )


# =========================================================
# CHAT / FILE UPLOAD
# =========================================================

def _persist_completed_chat_message(conversation_id, user_id, message):
    """Retry a failed generated-answer save after returning the response."""
    try:
        saved = append_conversation_message(conversation_id, user_id, message)
        if not saved:
            logger.warning(
                "Completed chat response could not be saved conversation=%s user=%s",
                conversation_id,
                user_id,
            )
    except Exception:
        logger.exception(
            "Background save failed for completed chat conversation=%s user=%s",
            conversation_id,
            user_id,
        )

@app.post("/api/chat-generations/{generation_id}/cancel", name="cancel_chat_generation")
async def cancel_chat_generation(request: Request, generation_id: str):
    user_id = str(require_user_id(request))
    with _chat_generation_lock:
        generation = _active_chat_generations.get(generation_id)
        if generation and generation["user_id"] == user_id:
            generation["cancelled"].set()
            return JSONResponse({"success": True, "cancelled": True})
    return JSONResponse({"success": True, "cancelled": False})


_SWARM_RUNS_LOCK = threading.RLock()
_SWARM_ACTIVE_RUNS = {}
_SWARM_RUN_HISTORY_LIMIT = 20
_SWARM_SETTINGS_DEFAULTS = {"auto_route": True, "agents": [], "output_format": "answer", "max_concurrency": 3}


def _swarm_flag(value, default=False):
    if isinstance(value, bool):
        return value
    if value is None:
        return bool(default)
    return str(value).strip().casefold() in {"1", "true", "yes", "on"}


def _swarm_cache_headers():
    return {"Cache-Control": "no-store, no-cache, must-revalidate, max-age=0"}


def _swarm_settings(user_id):
    settings = get_user_settings(user_id)
    return settings if isinstance(settings, dict) else {}


def _swarm_runs_from(settings):
    runs = settings.get("swarm_runs", []) if isinstance(settings, dict) else []
    return runs if isinstance(runs, list) else []


def _swarm_find_run(user_id, run_id):
    settings = _swarm_settings(user_id)
    for item in _swarm_runs_from(settings):
        if isinstance(item, dict) and str(item.get("id")) == str(run_id):
            return copy.deepcopy(item)
    return None


def _swarm_save_run(user_id, record):
    with _SWARM_RUNS_LOCK:
        settings = _swarm_settings(user_id)
        runs = [item for item in _swarm_runs_from(settings) if isinstance(item, dict) and item.get("id") != record.get("id")]
        runs.insert(0, copy.deepcopy(record))
        settings["swarm_runs"] = sorted(runs, key=lambda item: str(item.get("created_at") or ""), reverse=True)[:_SWARM_RUN_HISTORY_LIMIT]
        save_user_settings(user_id, settings)


def _swarm_mutate_run(user_id, run_id, mutation):
    with _SWARM_RUNS_LOCK:
        settings = _swarm_settings(user_id)
        runs = _swarm_runs_from(settings)
        for index, item in enumerate(runs):
            if isinstance(item, dict) and str(item.get("id")) == str(run_id):
                record = copy.deepcopy(item)
                mutation(record)
                record["updated_at"] = datetime.now(timezone.utc).isoformat()
                runs[index] = record
                settings["swarm_runs"] = runs
                save_user_settings(user_id, settings)
                return copy.deepcopy(record)
    return None


def _swarm_public_run(record, *, summary=False):
    value = copy.deepcopy(record)
    value.pop("idempotency_key", None)
    if summary:
        value.pop("task", None)
        value.pop("result", None)
        value.pop("reports", None)
        value.pop("plan", None)
        value.pop("attachments", None)
        value.pop("artifacts", None)
    else:
        for report in value.get("reports", []):
            if isinstance(report, dict):
                for key in list(report):
                    if key.startswith("_"):
                        report.pop(key, None)
    return value


def _swarm_updated_task(user_id, run_id, task_record):
    def update(record):
        for entry in record.get("plan", []):
            if entry.get("id") == task_record.get("id"):
                entry.update({key: task_record.get(key) for key in ("status", "attempts", "output", "error", "retryable")})
                break
        record["status"] = "cancelling" if record.get("cancel_requested") else "running"
        record["current_task"] = task_record.get("name") if task_record.get("status") in {"running", "retrying"} else ""
    try:
        _swarm_mutate_run(user_id, run_id, update)
    except Exception as error:
        logger.warning("Swarm progress could not be saved for user %s (%s).", user_id, type(error).__name__)


def _swarm_load_run_files(user_id, attachments):
    blobs = []
    for item in attachments or []:
        key = str(item.get("storage_key") or "")
        if not key:
            continue
        data = storage_router.read_file(user_id, key)
        blobs.append((str(item.get("name") or "attachment"), str(item.get("mime_type") or "application/octet-stream"), data))
    return blobs


MAX_SWARM_LIBRARY_SEARCH_BYTES = 64 * 1024 * 1024
MAX_SWARM_LIBRARY_CONTEXT_CHARS = 300_000


def _chat_requests_library_content(message, history=None):
    """Return whether this turn asks to inspect Library file contents."""
    current = str(message or "").casefold()
    recent = []
    for item in (history or [])[-12:]:
        if isinstance(item, dict):
            role = str(item.get("role") or "").casefold()
            content = str(item.get("text") or item.get("content") or "")
            if role and role not in {"user", "assistant"}:
                continue
        else:
            content = str(item or "")
        if content.strip():
            recent.append(content.casefold())

    library_pattern = r"\blibrary\b|\bmy\s+(?:uploaded\s+)?(?:files|documents|uploads)\b"
    mentions_library = bool(re.search(library_pattern, current))
    mentions_library_in_history = any(re.search(library_pattern, text) for text in recent)
    content_action = bool(re.search(
        r"\b(check|read|open|search|find|scan|review|compare|analy[sz]e|summari[sz]e|extract|inspect|look\s+through|go\s+through)\b",
        current,
    ))
    content_target = bool(re.search(
        r"\b(company|companies|employer|employers|name|names|contents?|details?|skills?|experience|education|job|jobs|application|applications|summary|information)\b",
        current,
    ))
    all_files = bool(re.search(
        r"\b(?:all|every|each|entire|whole)\b.{0,60}\b(?:library\s+)?(?:files?|documents?|uploads?)\b|"
        r"\blibrary\s+(?:files?|documents?|uploads?)\b.{0,40}\b(?:all|every|each)\b",
        current,
    ))
    inventory_only = bool(re.search(r"\b(how\s+many|count|list|show|which|available)\b", current)) and not content_action and not content_target
    prior_content_request = any(
        re.search(library_pattern, text)
        and re.search(r"\b(check|read|open|search|find|scan|review|compare|analy[sz]e|summari[sz]e|extract|inspect|company|companies|employer|names?)\b", text)
        for text in recent
    )

    if inventory_only:
        return False
    return (
        (mentions_library and (content_action or content_target))
        or (all_files and mentions_library_in_history and prior_content_request)
    )


def _swarm_search_library(user_id, query):
    """Search the authenticated user's Library locally and return bounded excerpts."""
    stop_words = {
        "about", "after", "also", "and", "are", "because", "before", "between", "could", "does", "from", "have", "into",
        "any", "can", "get", "how", "just", "many", "me", "more", "most", "only", "please", "search", "should", "tell", "that", "their", "them", "there", "these", "this", "will",
        "through", "under", "using", "what", "when", "where", "which", "with", "would", "your", "मेरे", "मेरा", "मेरी",
        "में", "का", "की", "के", "और", "है", "हैं", "क्या", "कौन", "किस", "मुझे", "करो", "करना", "बताओ",
    }
    query_terms = list(dict.fromkeys(
        normalized
        for token in re.findall(r"[\w'-]{2,}", str(query or "").casefold())
        if token not in stop_words
        for normalized in (
            token[:-3] + "y" if token.endswith("ies") and len(token) > 4
            else token[:-1] if token.endswith("s") and len(token) > 4
            else token,
        )
    ))[:80]
    document_relevance_terms = (
        "achievement", "application", "candidate", "certification", "company",
        "degree", "education", "employment", "experience", "job", "position",
        "project", "qualification", "requirement", "responsibility", "role",
        "salary", "skill", "summary", "technology", "tool", "work",
    )
    records = get_files(user_id)
    context_parts = []
    files_searched = []
    file_references = []
    files_skipped = []
    total_bytes = 0
    max_chars_per_file = max(900, MAX_SWARM_LIBRARY_CONTEXT_CHARS // max(1, len(records)) - 300)

    for record in records[:30]:
        filename = os.path.basename(str(record.get("filename") or "Library file").replace("\\", "/"))[:180]
        try:
            reported_size = max(0, int(record.get("size_bytes") or 0))
        except (TypeError, ValueError):
            reported_size = 0
        if reported_size > MAX_SWARM_LIBRARY_SEARCH_BYTES or total_bytes + reported_size > MAX_SWARM_LIBRARY_SEARCH_BYTES:
            files_skipped.append(filename)
            continue
        try:
            storage_key = str(record.get("storage_key") or "")
            if storage_key:
                data = storage_router.read_file(user_id, storage_key)
            else:
                full_record = get_file(record.get("id"), user_id)
                data = (full_record or {}).get("file_data")
                if isinstance(data, memoryview):
                    data = data.tobytes()
            if not isinstance(data, bytes) or not data:
                files_skipped.append(filename)
                continue
            if len(data) > MAX_SWARM_LIBRARY_SEARCH_BYTES or total_bytes + len(data) > MAX_SWARM_LIBRARY_SEARCH_BYTES:
                files_skipped.append(filename)
                continue
            total_bytes += len(data)
            mime = str(record.get("mime_type") or "application/octet-stream")
            extracted = read_attachments([SwarmFile(filename, mime, data)])
            document = str(extracted.get("documents") or "").strip()
            if not document:
                files_skipped.append(filename)
                continue
            files_searched.append(filename)
            if record.get("id") is not None:
                file_references.append({
                    "type": "file",
                    "id": str(record.get("id")),
                    "name": filename,
                })

            document_lines = [part.strip() for part in re.split(r"\n+", document) if part.strip()]
            blocks = [
                "\n".join(document_lines[index:index + 4])
                for index in range(0, len(document_lines), 4)
            ] or [document]
            ranked = []
            for index, block in enumerate(blocks):
                lowered = block.casefold()
                score = sum(min(3, lowered.count(term)) for term in query_terms)
                score += sum(
                    min(2, lowered.count(term))
                    for term in document_relevance_terms
                )
                if filename.casefold() in lowered:
                    score += 2
                if score:
                    ranked.append((score, index, block))
            if ranked:
                passage_limit = max(500, min(2200, max_chars_per_file // 5))
                chosen = []
                for _, _, block in sorted(ranked, key=lambda item: (-item[0], item[1]))[:5]:
                    if len(block) > passage_limit:
                        lowered = block.casefold()
                        positions = [lowered.find(term) for term in query_terms if term in lowered]
                        start = max(0, min(positions, default=0) - passage_limit // 3)
                        end = min(len(block), start + passage_limit)
                        excerpt_block = ("..." if start else "") + block[start:end] + ("..." if end < len(block) else "")
                    else:
                        excerpt_block = block
                    chosen.append(excerpt_block)
                excerpt = "\n\n--- Relevant passage ---\n\n".join(chosen)
            else:
                excerpt = "No exact query terms matched this file; opening context follows.\n\n" + "\n\n".join(blocks[:2])
            excerpt = excerpt[:max_chars_per_file].strip()
            context_parts.append(f"===== LIBRARY FILE: {filename} =====\n{excerpt}")
        except Exception as error:
            logger.warning("Swarm Library search skipped one file for user %s: %s", user_id, type(error).__name__)
            files_skipped.append(filename)

    context = "\n\n".join(context_parts)[:MAX_SWARM_LIBRARY_CONTEXT_CHARS]
    return {
        "context": context,
        "files_searched": files_searched,
        "file_references": file_references,
        "files_skipped": files_skipped,
    }


def _execute_swarm_run(user_id, run_id, identity_email, memory_enabled, max_concurrency, cancel_event, seed_reports=None):
    try:
        record = _swarm_find_run(user_id, run_id)
        if not record:
            return
        file_blobs = _swarm_load_run_files(user_id, record.get("attachments", []))
        identity = create_identity(user_id=str(user_id), email=str(identity_email or ""))
        outcome = run_swarm_task(
            record["task"],
            list(record.get("agent_ids") or []),
            file_blobs,
            str(user_id),
            identity,
            bool(memory_enabled),
            web_search=bool(record.get("web_search")),
            output_format=str(record.get("output_format") or "answer"),
            max_concurrency=max_concurrency,
            cancel_event=cancel_event,
            on_task_update=lambda task: _swarm_updated_task(user_id, run_id, task),
            seed_reports=seed_reports,
            library_file_loader=lambda query: _swarm_search_library(user_id, query),
        )
        reports = []
        artifact_save_errors = []
        for raw_report in outcome.get("reports", []):
            report = {key: value for key, value in raw_report.items() if not key.startswith("_")}
            artifact_bytes = raw_report.get("_artifact_bytes")
            if artifact_bytes:
                try:
                    stored = storage_router.save_file(user_id, artifact_bytes, raw_report.get("_artifact_name") or "swarm-artifact")
                    artifact = {
                        "id": uuid.uuid4().hex,
                        "filename": os.path.basename(str(raw_report.get("_artifact_name") or "swarm-artifact"))[:180],
                        "media_type": str(raw_report.get("_artifact_mime_type") or "application/octet-stream"),
                        "storage_key": stored["storage_key"],
                        "size_bytes": int(stored.get("size_bytes") or len(artifact_bytes)),
                    }
                    report["artifacts"] = [artifact]
                except Exception as error:
                    if report.get("id") == "writing":
                        report.setdefault("data", {})["word_export"] = "unavailable"
                        report["error"] = "The draft completed, but its Word file could not be saved. The editable Markdown is still available."
                    else:
                        report["status"] = "failed"
                        report["error"] = f"The generated file could not be saved to your private file storage: {type(error).__name__}."
                        report["output"] = "The specialist output was created, but its generated file could not be saved."
                        artifact_save_errors.append(report["error"])
            reports.append(report)
        result_text = str(outcome.get("result") or "")
        status = str(outcome.get("status") or "failed")
        if artifact_save_errors and status == "completed":
            status = "partial"
            result_text = (result_text + "\n\nA generated artifact could not be saved to private storage.").strip()
        now = datetime.now(timezone.utc).isoformat()

        def finish(run):
            run.update(
                status=status,
                reports=reports,
                agents=[{"id": item["id"], "name": item.get("name", item["id"]), "status": item.get("status", "failed")} for item in reports],
                plan=outcome.get("plan", run.get("plan", [])),
                result=result_text,
                error=str(outcome.get("error") or ""),
                artifacts=[artifact for report in reports for artifact in report.get("artifacts", [])],
                current_task="",
                cancel_requested=False,
                finished_at=now,
            )
        _swarm_mutate_run(user_id, run_id, finish)
    except Exception as error:
        logger.warning("Swarm execution failed for user %s (%s).", user_id, type(error).__name__)
        try:
            _swarm_mutate_run(user_id, run_id, lambda run: run.update(
                status="failed",
                error="ODDI could not complete this Swarm run. You can retry its failed specialists.",
                current_task="",
                finished_at=datetime.now(timezone.utc).isoformat(),
            ))
        except Exception as save_error:
            logger.warning("Swarm failure state could not be saved for user %s (%s).", user_id, type(save_error).__name__)
    finally:
        with _SWARM_RUNS_LOCK:
            active = _SWARM_ACTIVE_RUNS.get(str(run_id))
            if active and active.get("user_id") == str(user_id):
                _SWARM_ACTIVE_RUNS.pop(str(run_id), None)


def _swarm_prepare_run(user_id, *, task, agent_ids, web_search, output_format, automatic, attachments, background_tasks, request_email, max_concurrency=None, seed_reports=None, retry_of="", idempotency_key=""):
    now = datetime.now(timezone.utc).isoformat()
    plan = create_swarm_execution_plan(agent_ids)
    run_id = uuid.uuid4().hex
    record = {
        "id": run_id,
        "title": task[:88].strip() or "Swarm task",
        "task": task,
        "created_at": now,
        "updated_at": now,
        "status": "pending",
        "auto_route": bool(automatic),
        "web_search": bool(web_search),
        "output_format": output_format,
        "agent_ids": list(agent_ids),
        "agents": [{"id": agent_id, "name": SWARM_AGENT_REGISTRY[agent_id].name, "status": "pending"} for agent_id in agent_ids],
        "plan": plan,
        "reports": [copy.deepcopy(item) for item in (seed_reports or [])],
        "result": "",
        "error": "",
        "artifacts": [],
        "attachments": copy.deepcopy(attachments or []),
        "retry_of": retry_of,
        "current_task": "",
        "cancel_requested": False,
        "idempotency_key": idempotency_key,
    }
    try:
        _swarm_save_run(user_id, record)
    except Exception as error:
        raise HTTPException(status_code=503, detail="Swarm run history could not be saved. Please try again.") from error
    cancel_event = threading.Event()
    with _SWARM_RUNS_LOCK:
        _SWARM_ACTIVE_RUNS[run_id] = {"user_id": str(user_id), "cancel_event": cancel_event}
    try:
        settings = _swarm_settings(user_id)
        memory_config = normalize_memory_settings(settings.get("memory") if isinstance(settings, dict) else {})
        memory_enabled = bool(memory_config["enabled"] and not memory_user_disabled(user_id) and not memory_user_opted_out(user_id))
        concurrency = max(1, min(SWARM_MAX_AGENT_CONCURRENCY, int(max_concurrency or SWARM_MAX_AGENT_CONCURRENCY)))
        background_tasks.add_task(
            _execute_swarm_run,
            str(user_id), run_id, request_email, memory_enabled, concurrency, cancel_event, seed_reports,
        )
    except Exception:
        with _SWARM_RUNS_LOCK:
            _SWARM_ACTIVE_RUNS.pop(run_id, None)
        _swarm_mutate_run(user_id, run_id, lambda run: run.update(status="failed", error="The run could not be queued."))
        raise
    return record


@app.get("/api/swarm/agents", name="api_swarm_agents")
async def api_swarm_agents(request: Request):
    await require_user_id_async(request)
    return JSONResponse({"agents": get_swarm_agent_registry(), "max_agents": SWARM_MAX_AGENT_COUNT}, headers=_swarm_cache_headers())


@app.get("/api/swarm/settings", name="api_swarm_settings")
async def api_swarm_settings(request: Request):
    user_id = await require_user_id_async(request)
    settings = await run_in_threadpool(_swarm_settings, user_id)
    preferences = settings.get("swarm_preferences", {}) if isinstance(settings, dict) else {}
    merged = {**_SWARM_SETTINGS_DEFAULTS, **(preferences if isinstance(preferences, dict) else {})}
    valid_agent_ids = set(SWARM_AGENT_REGISTRY)
    stored_agents = merged.get("agents")
    merged["agents"] = [str(agent_id).strip().casefold() for agent_id in stored_agents if str(agent_id).strip().casefold() in valid_agent_ids] if isinstance(stored_agents, list) else []
    if merged.get("output_format") not in {"answer", "markdown", "report"}:
        merged["output_format"] = "answer"
    return JSONResponse({"settings": merged}, headers=_swarm_cache_headers())


@app.post("/api/swarm/settings", name="api_update_swarm_settings")
async def api_update_swarm_settings(request: Request):
    user_id = await require_user_id_async(request)
    try:
        payload = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="Settings must be sent as JSON.")
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="Settings must be a JSON object.")
    current = await run_in_threadpool(_swarm_settings, user_id)
    preferences = {**_SWARM_SETTINGS_DEFAULTS, **(current.get("swarm_preferences", {}) if isinstance(current.get("swarm_preferences"), dict) else {})}
    valid_agent_ids = set(SWARM_AGENT_REGISTRY)
    stored_agents = preferences.get("agents")
    preferences["agents"] = [str(agent_id).strip().casefold() for agent_id in stored_agents if str(agent_id).strip().casefold() in valid_agent_ids] if isinstance(stored_agents, list) else []
    if preferences.get("output_format") not in {"answer", "markdown", "report"}:
        preferences["output_format"] = "answer"
    if "auto_route" in payload:
        preferences["auto_route"] = _swarm_flag(payload["auto_route"])
    if "agents" in payload:
        try:
            preferences["agents"] = validate_swarm_agent_ids(payload["agents"], max_count=SWARM_MAX_AGENT_COUNT) if payload["agents"] else []
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
    if "output_format" in payload:
        output_format = str(payload["output_format"] or "answer").strip().casefold()
        if output_format not in {"answer", "markdown", "report"}:
            raise HTTPException(status_code=400, detail="Choose a supported text output format.")
        preferences["output_format"] = output_format
    if "max_concurrency" in payload:
        try:
            preferences["max_concurrency"] = max(1, min(SWARM_MAX_AGENT_CONCURRENCY, int(payload["max_concurrency"])))
        except (TypeError, ValueError):
            raise HTTPException(status_code=400, detail="Concurrency must be between one and three.")
    current["swarm_preferences"] = preferences
    try:
        await run_in_threadpool(save_user_settings, user_id, current)
    except Exception as error:
        raise HTTPException(status_code=503, detail="Swarm settings could not be saved.") from error
    return JSONResponse({"settings": preferences}, headers=_swarm_cache_headers())


@app.post("/api/swarm/plan", name="api_swarm_plan")
async def api_swarm_plan(request: Request):
    user_id = await require_user_id_async(request)
    try:
        payload = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="The Swarm plan request must be JSON.")
    if not isinstance(payload, dict):
        raise HTTPException(status_code=400, detail="The Swarm plan request must be a JSON object.")
    task = str(payload.get("task") or "").strip()
    if not task:
        raise HTTPException(status_code=400, detail="Describe the task you want the Swarm to handle.")
    if len(task) > 6000:
        raise HTTPException(status_code=413, detail="Keep the task under 6,000 characters.")
    settings = await run_in_threadpool(_swarm_settings, user_id)
    preferences = settings.get("swarm_preferences", {}) if isinstance(settings, dict) else {}
    auto_mode = _swarm_flag(payload.get("auto_route"), preferences.get("auto_route", True))
    web_search = _swarm_flag(payload.get("web_search"))
    file_names = payload.get("file_names") or []
    if not isinstance(file_names, list):
        file_names = []
    if auto_mode:
        agent_ids = swarm_auto_route(task, _swarm_flag(payload.get("has_files")), web_search, str(payload.get("output_format") or "answer"), [str(name)[:180] for name in file_names[:5]])
    else:
        try:
            agent_ids = validate_swarm_agent_ids(payload.get("agents") or [], max_count=SWARM_MAX_AGENT_COUNT)
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
    return JSONResponse({"agent_ids": agent_ids, "plan": create_swarm_execution_plan(agent_ids), "auto_route": auto_mode}, headers=_swarm_cache_headers())


@app.get("/api/swarm/runs", name="api_swarm_runs")
async def api_swarm_runs(request: Request):
    user_id = await require_user_id_async(request)
    settings = await run_in_threadpool(_swarm_settings, user_id)
    runs = sorted((item for item in _swarm_runs_from(settings) if isinstance(item, dict)), key=lambda item: str(item.get("created_at") or ""), reverse=True)
    recovered = []
    for item in runs[:_SWARM_RUN_HISTORY_LIMIT]:
        recovered.append(await run_in_threadpool(_recover_interrupted_swarm_run, user_id, item))
    return JSONResponse({"runs": [_swarm_public_run(item, summary=True) for item in recovered]}, headers=_swarm_cache_headers())


def _recover_interrupted_swarm_run(user_id, run):
    active_states = {"pending", "running", "retrying", "cancelling"}
    recovery_prefix = "The Swarm stopped responding before this run finished."
    is_active_run = run.get("status") in active_states
    is_recovered_run = run.get("status") == "failed" and str(run.get("error") or "").startswith(recovery_prefix)
    unfinished_snapshot = any(
        isinstance(item, dict) and item.get("status") in active_states
        for key in ("reports", "agents", "plan")
        for item in (run.get(key) or [])
    )
    if not is_active_run and not (is_recovered_run and unfinished_snapshot):
        return run
    if is_active_run:
        with _SWARM_RUNS_LOCK:
            active = _SWARM_ACTIVE_RUNS.get(str(run.get("id")))
        if active and active.get("user_id") == str(user_id):
            return run
        try:
            updated = datetime.fromisoformat(str(run.get("updated_at") or run.get("created_at") or "").replace("Z", "+00:00"))
        except (TypeError, ValueError):
            updated = datetime.min.replace(tzinfo=timezone.utc)
        if (datetime.now(timezone.utc) - updated).total_seconds() < 900:
            return run
    def mark_interrupted(value):
        reports = value.get("reports") if isinstance(value.get("reports"), list) else []
        value["reports"] = reports
        report_by_id = {report.get("id"): report for report in reports if isinstance(report, dict)}
        agent_by_id = {agent.get("id"): agent for agent in value.get("agents", []) if isinstance(agent, dict)}
        plan_by_agent = {task.get("agent_id"): task for task in value.get("plan", []) if isinstance(task, dict)}
        agent_ids = list(dict.fromkeys(str(agent_id) for agent_id in (value.get("agent_ids") or [])))
        for agent_id in agent_ids:
            report = report_by_id.get(agent_id)
            summary = agent_by_id.get(agent_id, {})
            task = plan_by_agent.get(agent_id, {})
            unfinished = (
                summary.get("status") in active_states
                or task.get("status") in active_states
                or (report and report.get("status") in active_states)
            )
            if not unfinished:
                continue
            specialist = SWARM_AGENT_REGISTRY.get(agent_id)
            retryable = bool(specialist and specialist.available)
            specialist_error = (
                "The server stopped before this specialist returned a response."
                if retryable else "This specialist is no longer available. Start a new task with a current specialist."
            )
            if report is None:
                report = {"id": agent_id, "name": getattr(specialist, "name", summary.get("name") or agent_id),
                          "output": "", "sources": [], "artifacts": [], "data": {}, "attempts": 1}
                reports.append(report)
                report_by_id[agent_id] = report
            report.update(status="failed" if retryable else "unavailable", error=specialist_error, retryable=retryable)
            if summary:
                summary["status"] = report["status"]
            if task:
                task.update(status=report["status"], error=specialist_error, retryable=retryable)
        for report in reports:
            if isinstance(report, dict) and report.get("status") in active_states:
                specialist = SWARM_AGENT_REGISTRY.get(str(report.get("id") or ""))
                retryable = bool(specialist and specialist.available)
                report.update(status="failed" if retryable else "unavailable",
                              error="The server stopped before this specialist returned a response." if retryable else "This specialist is no longer available. Start a new task with a current specialist.",
                              retryable=retryable)
        retryable_reports = [report for report in reports if isinstance(report, dict) and report.get("retryable")]
        error = (
            "The Swarm stopped responding before this run finished. Retry the incomplete specialists to continue."
            if retryable_reports else
            "The Swarm stopped responding before this run finished. Start a new task with a current specialist."
        )
        for task in value.get("plan", []):
            if isinstance(task, dict) and task.get("agent_id") == "central" and task.get("status") in active_states:
                task.update(status="failed", error=error, retryable=False)
        report_by_id = {report.get("id"): report for report in value.get("reports", []) if isinstance(report, dict)}
        for agent in value.get("agents", []):
            if isinstance(agent, dict) and agent.get("id") in report_by_id:
                agent["status"] = report_by_id[agent["id"]].get("status", agent.get("status"))
        value.update(status="failed", error=error, current_task="", cancel_requested=False,
                     finished_at=datetime.now(timezone.utc).isoformat())
    return _swarm_mutate_run(user_id, run["id"], mark_interrupted) or run


@app.get("/api/swarm/runs/{run_id}", name="api_swarm_run_detail")
async def api_swarm_run_detail(request: Request, run_id: str):
    user_id = await require_user_id_async(request)
    run = await run_in_threadpool(_swarm_find_run, user_id, run_id)
    if not run:
        raise HTTPException(status_code=404, detail="Swarm run not found.")
    run = await run_in_threadpool(_recover_interrupted_swarm_run, user_id, run)
    return JSONResponse({"run": _swarm_public_run(run)}, headers=_swarm_cache_headers())


@app.patch("/api/swarm/runs/{run_id}/writing", name="api_save_swarm_writing_draft")
async def api_save_swarm_writing_draft(request: Request, run_id: str):
    user_id = await require_user_id_async(request)
    try:
        payload = await request.json()
    except Exception as error:
        raise HTTPException(status_code=400, detail="The writing draft could not be read.") from error
    markdown = payload.get("markdown") if isinstance(payload, dict) else None
    if not isinstance(markdown, str) or not markdown.strip():
        raise HTTPException(status_code=400, detail="The writing draft cannot be empty.")
    if len(markdown) > 50_000:
        raise HTTPException(status_code=413, detail="Keep the writing draft under 50,000 characters.")

    def save_draft(record):
        report = next((
            item for item in record.get("reports", [])
            if isinstance(item, dict) and item.get("id") == "writing"
        ), None)
        if not report or report.get("status") != "completed":
            raise ValueError("This run has no completed Writing Agent draft to edit.")
        report["output"] = markdown.strip()
        data = report.get("data")
        if not isinstance(data, dict):
            data = {}
            report["data"] = data
        data["edited_at"] = datetime.now(timezone.utc).isoformat()

    try:
        run = await run_in_threadpool(_swarm_mutate_run, user_id, run_id, save_draft)
    except ValueError as error:
        raise HTTPException(status_code=409, detail=str(error)) from error
    if not run:
        raise HTTPException(status_code=404, detail="Swarm run not found.")
    return JSONResponse({"run": _swarm_public_run(run)}, headers=_swarm_cache_headers())


@app.get("/api/swarm/runs/{run_id}/artifacts/{artifact_id}", name="api_swarm_artifact")
async def api_swarm_artifact(request: Request, run_id: str, artifact_id: str):
    user_id = await require_user_id_async(request)
    run = await run_in_threadpool(_swarm_find_run, user_id, run_id)
    if not run:
        raise HTTPException(status_code=404, detail="Swarm run not found.")
    artifact = next((item for item in run.get("artifacts", []) if str(item.get("id")) == str(artifact_id)), None)
    if not artifact:
        raise HTTPException(status_code=404, detail="Swarm artifact not found.")
    try:
        path = await run_in_threadpool(storage_router.get_file_path, user_id, artifact["storage_key"])
    except (StorageFileNotFound, KeyError):
        raise HTTPException(status_code=404, detail="Swarm artifact is no longer available.")
    media_type = artifact.get("media_type") or "application/octet-stream"
    inline_preview = request.query_params.get("inline") == "1"
    if inline_preview and media_type not in {"image/png", "image/jpeg", "image/webp"}:
        raise HTTPException(status_code=400, detail="Only generated images can be opened as a preview.")
    return FileResponse(
        path,
        media_type=media_type,
        filename=os.path.basename(artifact.get("filename") or "swarm-output"),
        headers={"Cache-Control": "private, no-store"},
        content_disposition_type="inline" if inline_preview else "attachment",
    )


@app.post("/api/swarm/run", name="api_create_swarm_run")
async def api_create_swarm_run(request: Request, background_tasks: BackgroundTasks):
    user_id = await require_user_id_async(request)
    try:
        form = await request.form()
    except Exception as error:
        raise HTTPException(status_code=400, detail="The Swarm request could not be read.") from error
    task = str(form.get("task") or "").strip()
    if not task:
        raise HTTPException(status_code=400, detail="Describe the task you want the Swarm to handle.")
    if len(task) > 6000:
        raise HTTPException(status_code=413, detail="Keep the task under 6,000 characters.")
    idem = str(form.get("idempotency_key") or "").strip()[:100]
    if idem:
        settings = await run_in_threadpool(_swarm_settings, user_id)
        duplicate = next((item for item in _swarm_runs_from(settings) if isinstance(item, dict) and item.get("idempotency_key") == idem), None)
        if duplicate:
            return JSONResponse({"run": _swarm_public_run(duplicate)}, headers=_swarm_cache_headers())
    raw_agents = form.get("agents")
    try:
        requested_agents = json.loads(str(raw_agents)) if raw_agents else []
    except (json.JSONDecodeError, TypeError):
        raise HTTPException(status_code=400, detail="The selected agent list could not be read.")
    web_search = str(form.get("web_search") or "").strip().casefold() in {"1", "true", "yes", "on"}
    settings = await run_in_threadpool(_swarm_settings, user_id)
    preferences = settings.get("swarm_preferences", {}) if isinstance(settings, dict) else {}
    automatic = _swarm_flag(form.get("auto_route"), preferences.get("auto_route", True))
    selection_mode = str(form.get("selection_mode") or "").strip().casefold()
    if selection_mode == "manual_chat" and automatic:
        raise HTTPException(status_code=400, detail="A selected-agent chat must use the specialists chosen by the user.")
    max_selected_agents = 3 if selection_mode == "manual_chat" else SWARM_MAX_AGENT_COUNT
    output_format = str(form.get("output_format") or preferences.get("output_format") or "answer").strip().casefold()
    if output_format not in {"answer", "markdown", "report"}:
        raise HTTPException(status_code=400, detail="Choose a supported output format.")
    preferred_output_format = str(form.get("preferred_output_format") or output_format).strip().casefold()
    if preferred_output_format not in {"answer", "markdown", "report"}:
        raise HTTPException(status_code=400, detail="Choose a supported saved output format.")
    uploads = [item for item in form.getlist("files") if hasattr(item, "read") and getattr(item, "filename", "")]
    if len(uploads) > 5:
        raise HTTPException(status_code=413, detail="Attach up to five files per Swarm run.")
    file_blobs = []
    total_bytes = 0
    for upload in uploads:
        content = await upload.read(MAX_UPLOAD_BYTES + 1)
        total_bytes += len(content)
        if len(content) > MAX_UPLOAD_BYTES or total_bytes > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="Swarm attachments have a combined 64 MB limit.")
        if not content:
            raise HTTPException(status_code=400, detail="An attached file is empty.")
        filename = os.path.basename(str(upload.filename or "attachment").replace("\\", "/"))[:180] or "attachment"
        mime = str(upload.content_type or "application/octet-stream")[:160]
        file_blobs.append((filename, mime, content))
    try:
        if file_blobs:
            await run_in_threadpool(read_attachments, [SwarmFile(*item) for item in file_blobs])
    except AttachmentReadError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    if automatic:
        agent_ids = swarm_auto_route(task, bool(file_blobs), web_search, output_format, [item[0] for item in file_blobs])
    else:
        try:
            agent_ids = validate_swarm_agent_ids(requested_agents, max_count=max_selected_agents)
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
    if not agent_ids:
        raise HTTPException(status_code=400, detail="Choose at least one specialist agent.")
    stored_attachments = []
    try:
        for filename, mime, content in file_blobs:
            stored = await run_in_threadpool(storage_router.save_file, user_id, content, filename)
            stored_attachments.append({"name": filename, "mime_type": mime, "storage_key": stored["storage_key"], "size_bytes": stored["size_bytes"]})
        settings["swarm_preferences"] = {
            **_SWARM_SETTINGS_DEFAULTS,
            **(settings.get("swarm_preferences", {}) if isinstance(settings.get("swarm_preferences"), dict) else {}),
            "auto_route": automatic,
            "agents": [] if automatic else agent_ids,
            "output_format": preferred_output_format,
        }
        await run_in_threadpool(save_user_settings, user_id, settings)
        record = _swarm_prepare_run(
            str(user_id), task=task, agent_ids=agent_ids, web_search=web_search, output_format=output_format,
            automatic=automatic, attachments=stored_attachments, background_tasks=background_tasks,
            request_email=request.session.get("email", ""), max_concurrency=preferences.get("max_concurrency"), idempotency_key=idem,
        )
    except StorageQuotaExceeded as error:
        for item in stored_attachments:
            try:
                await run_in_threadpool(storage_router.delete_file, user_id, item["storage_key"])
            except Exception:
                pass
        raise HTTPException(status_code=413, detail="Your private file storage does not have enough space for these Swarm attachments.") from error
    except HTTPException:
        for item in stored_attachments:
            try:
                await run_in_threadpool(storage_router.delete_file, user_id, item["storage_key"])
            except Exception:
                pass
        raise
    except Exception as error:
        for item in stored_attachments:
            try:
                await run_in_threadpool(storage_router.delete_file, user_id, item["storage_key"])
            except Exception:
                pass
        logger.warning("Swarm run could not be queued for user %s (%s).", user_id, type(error).__name__)
        raise HTTPException(status_code=503, detail="The Swarm run could not be queued. Please try again.") from error
    return JSONResponse({"run": _swarm_public_run(record)}, status_code=202, headers=_swarm_cache_headers())


@app.post("/api/swarm/runs/{run_id}/cancel", name="api_cancel_swarm_run")
async def api_cancel_swarm_run(request: Request, run_id: str):
    user_id = await require_user_id_async(request)
    run = await run_in_threadpool(_swarm_find_run, user_id, run_id)
    if not run:
        raise HTTPException(status_code=404, detail="Swarm run not found.")
    if run.get("status") not in {"pending", "running", "cancelling"}:
        return JSONResponse({"run": _swarm_public_run(run), "cancelled": False}, headers=_swarm_cache_headers())
    with _SWARM_RUNS_LOCK:
        active = _SWARM_ACTIVE_RUNS.get(str(run_id))
        if active and active.get("user_id") == str(user_id):
            active["cancel_event"].set()
    updated = await run_in_threadpool(_swarm_mutate_run, user_id, run_id, lambda item: item.update(status="cancelling", cancel_requested=True))
    return JSONResponse({"run": _swarm_public_run(updated or run), "cancelled": True}, headers=_swarm_cache_headers())


@app.post("/api/swarm/runs/{run_id}/retry", name="api_retry_swarm_run")
async def api_retry_swarm_run(request: Request, run_id: str, background_tasks: BackgroundTasks):
    user_id = await require_user_id_async(request)
    prior = await run_in_threadpool(_swarm_find_run, user_id, run_id)
    if not prior:
        raise HTTPException(status_code=404, detail="Swarm run not found.")
    if prior.get("status") in {"pending", "running", "cancelling"}:
        raise HTTPException(status_code=409, detail="Cancel or wait for this run before retrying it.")
    retryable = any(item.get("status") in {"failed", "blocked"} and item.get("retryable") for item in prior.get("reports", []) if isinstance(item, dict))
    if not retryable:
        raise HTTPException(status_code=409, detail="This run has no failed specialist that can be retried.")
    try:
        record = _swarm_prepare_run(
            str(user_id), task=str(prior.get("task") or ""), agent_ids=list(prior.get("agent_ids") or []),
            web_search=bool(prior.get("web_search")), output_format=str(prior.get("output_format") or "answer"),
            automatic=False, attachments=list(prior.get("attachments") or []), background_tasks=background_tasks,
            request_email=request.session.get("email", ""), retry_of=run_id, seed_reports=list(prior.get("reports") or []),
        )
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(status_code=503, detail="The failed Swarm tasks could not be queued for retry.") from error
    return JSONResponse({"run": _swarm_public_run(record)}, status_code=202, headers=_swarm_cache_headers())


@app.post("/chat", name="chat")
async def chat(request: Request, background_tasks: BackgroundTasks):
    generation_started = time.perf_counter()
    chat_request_id = getattr(request.state, "chat_request_id", uuid.uuid4().hex[:12])
    form = await request.form()

    message = form.get("message")
    if message is None:
        return JSONResponse({"error": "Message is required."}, status_code=400)
    message = str(message)
    response_style = str(form.get("response_style") or "").strip().casefold()
    incognito_mode = str(form.get("incognito") or "").strip().casefold() in {"1", "true", "yes", "on"}
    personalization_disabled = incognito_mode or ODDI_BROWSER_LOCAL_CHATS or is_temporary_browser_account(request)
    specialist_agent = str(form.get("specialist_agent") or "").strip().casefold()
    if specialist_agent not in {"research", "search", "coding", "document_search"}:
        specialist_agent = ""
    response_style_instructions = {
        "concise": "Answer concisely, prioritizing the most useful information.",
        "step_by_step": "Explain the answer as clear, numbered steps.",
        "table": "When comparing multiple items, use a compact Markdown table; otherwise answer normally.",
        "detailed": "Give a thorough, well-structured explanation with useful examples.",
    }
    memory_message = str(form.get("memory_message") or message).strip()
    assistant_message_id = str(form.get("assistant_message_id") or "").strip()
    if not assistant_message_id or len(assistant_message_id) > 128:
        assistant_message_id = str(uuid.uuid4())
    document_workspace_requested = str(form.get("document_workspace") or "").strip().casefold() in {"1", "true", "yes", "on"}
    document_title = str(form.get("document_title") or "").strip()[:120]
    generation_id = str(form.get("generation_id") or "").strip()
    if not generation_id or len(generation_id) > 128:
        generation_id = str(uuid.uuid4())
    memory_enabled_value = form.get("memory_enabled")
    client_memory_enabled = (
        str(memory_enabled_value).strip().casefold() in {"1", "true", "yes", "on"}
        if memory_enabled_value is not None
        else None
    )
    auto_extract_value = form.get("memory_auto_extract")
    client_memory_auto_extract = (
        str(auto_extract_value).strip().casefold() in {"1", "true", "yes", "on"}
        if auto_extract_value is not None
        else None
    )
    if ODDI_BROWSER_LOCAL_CHATS or incognito_mode:
        # A cloud process may generate a response, but it must not persist
        # chat-derived memory alongside a browser-local or Incognito chat.
        client_memory_enabled = False
        client_memory_auto_extract = False

    raw_conversation_id = form.get("conversation_id")
    try:
        conversation_id = int(raw_conversation_id) if raw_conversation_id else None
    except (TypeError, ValueError):
        conversation_id = None
    if incognito_mode:
        # Incognito replies are generated without creating or consulting an
        # account conversation row, even if a stale client ID is submitted.
        conversation_id = None
    if ODDI_BROWSER_LOCAL_CHATS:
        # Client-generated IDs are intentionally not mapped to a server row.
        conversation_id = None

    # Starlette FormData supports multiple values for the same field name.
    raw_uploaded_files = form.getlist("files")
    uploaded_files = [
        item for item in raw_uploaded_files
        if hasattr(item, "file") and hasattr(item, "filename")
    ]

    # Convert FastAPI UploadFile objects into a Flask-FileStorage-compatible
    # object before passing them to the existing ODDI engine.
    engine_files = adapt_uploaded_files(uploaded_files)
    library_source_indexes = set()
    for raw_index in form.getlist("library_source_indexes"):
        try:
            library_source_indexes.add(int(raw_index))
        except (TypeError, ValueError):
            continue
    requested_source_file_ids = list(dict.fromkeys(
        str(raw_id).strip()
        for raw_id in form.getlist("source_file_ids")
        if str(raw_id).strip() and len(str(raw_id).strip()) <= 128
    ))
    requested_source_file_id_set = set(requested_source_file_ids)
    try:
        library_request_history = json.loads(str(form.get("history", "[]")))
    except (json.JSONDecodeError, TypeError):
        library_request_history = []
    if not isinstance(library_request_history, list):
        library_request_history = []

    # Chat is private. Authenticate before storing uploaded files.
    user_id = await require_user_id_async(request)

    # Persist the current user turn before generation can persist its reply.
    # The browser also writes this message for recovery; the shared message ID
    # makes that write an idempotent no-op if it arrives after this one.
    raw_user_message_record = form.get("user_message_record")
    if conversation_id is not None and raw_user_message_record:
        try:
            candidate = json.loads(str(raw_user_message_record))
        except (json.JSONDecodeError, TypeError, ValueError):
            candidate = None
        if isinstance(candidate, dict) and candidate.get("role") == "user":
            message_id = str(candidate.get("id") or "").strip()
            if message_id and len(message_id) <= 128:
                user_message_record = {
                    "id": message_id,
                    "role": "user",
                    "text": str(candidate.get("text") or memory_message),
                    "files": [str(name)[:255] for name in candidate.get("files", [])[:30]]
                    if isinstance(candidate.get("files"), list) else [],
                    "fileRefs": [
                        {
                            "id": str(file_ref.get("id") or "")[:128],
                            "name": str(file_ref.get("name") or "")[:255],
                            "type": str(file_ref.get("type") or "")[:128],
                        }
                        for file_ref in candidate.get("fileRefs", [])[:30]
                        if isinstance(file_ref, dict) and file_ref.get("id")
                    ] if isinstance(candidate.get("fileRefs"), list) else [],
                    "specialist_agent": str(candidate.get("specialist_agent") or "")[:64],
                    "pinned": False,
                }
                try:
                    saved_user_conversation = await run_in_threadpool(
                        append_conversation_message,
                        conversation_id,
                        user_id,
                        user_message_record,
                    )
                    saved_user_ids = {
                        str(item.get("id") or "")
                        for item in (saved_user_conversation or {}).get("messages", [])
                        if isinstance(item, dict)
                    }
                    if message_id not in saved_user_ids:
                        # A revision conflict can return the latest snapshot
                        # without applying this append; retry once from it.
                        saved_user_conversation = await run_in_threadpool(
                            append_conversation_message,
                            conversation_id,
                            user_id,
                            user_message_record,
                        )
                        saved_user_ids = {
                            str(item.get("id") or "")
                            for item in (saved_user_conversation or {}).get("messages", [])
                            if isinstance(item, dict)
                        }
                    if message_id not in saved_user_ids:
                        logger.warning(
                            "User chat message could not be confirmed before generation request_id=%s conversation=%s user=%s",
                            chat_request_id,
                            conversation_id,
                            user_id,
                        )
                except Exception:
                    logger.exception(
                        "User chat message persistence failed before generation request_id=%s conversation=%s user=%s",
                        chat_request_id,
                        conversation_id,
                        user_id,
                    )

    external_document_context = ""
    external_document_sources = []
    search_library_for_turn = (
        specialist_agent == "document_search"
        or _chat_requests_library_content(memory_message, library_request_history)
    )
    if search_library_for_turn:
        try:
            library_search = await run_in_threadpool(
                _swarm_search_library,
                user_id,
                memory_message,
            )
            external_document_context = str(library_search.get("context") or "")
            external_document_sources = list(library_search.get("file_references") or [])
            if not external_document_context:
                searched_names = library_search.get("files_searched") or []
                skipped_names = library_search.get("files_skipped") or []
                if searched_names:
                    external_document_context = (
                        "[ODDI LIBRARY CONTENT SEARCH: NO READABLE CONTENT EXTRACTED]\n"
                        "The Library entries were found, but no readable text could be extracted. "
                        "Tell the user that file contents could not be read and name any skipped files; "
                        "do not infer their contents.\n"
                        + "\n".join(f"- {name}" for name in [*searched_names, *skipped_names][:30])
                    )
                else:
                    external_document_context = (
                        "[ODDI LIBRARY CONTENT SEARCH: NO READABLE FILES FOUND]\n"
                        "No readable Library files were available for this request. Tell the user clearly; "
                        "do not answer from filenames alone or claim to have read files."
                    )
            logger.info(
                "Library content lookup request_id=%s user=%s agent=%s files=%s skipped=%s context_chars=%s",
                chat_request_id,
                user_id,
                specialist_agent or "none",
                len(library_search.get("files_searched") or []),
                len(library_search.get("files_skipped") or []),
                len(external_document_context),
            )
        except Exception as library_search_error:
            logger.warning(
                "Document Search Agent Library lookup failed request_id=%s user=%s (%s)",
                chat_request_id,
                user_id,
                type(library_search_error).__name__,
            )
    persist_library_uploads = not is_temporary_browser_account(request) and not incognito_mode
    verified_library_records = {}
    if requested_source_file_ids:
        try:
            verified_library_records = {
                str(record.get("id")): record
                for record in get_files(user_id)
                if record.get("id") is not None
                and str(record.get("id")) in requested_source_file_id_set
            }
        except Exception as source_lookup_error:
            logger.warning("Could not verify cited Library files for user %s: %s", user_id, type(source_lookup_error).__name__)
    fast_reply = fast_response(memory_message) if not engine_files and not specialist_agent and not search_library_for_turn else None
    settings_loaded = fast_reply is None
    settings_fetched = False
    user_settings = {}
    if settings_loaded:
        try:
            # Memory consent is server-authoritative, including for Drive-backed
            # accounts. A failed lookup fails closed for memory while chat itself
            # remains available.
            user_settings = await run_in_threadpool(get_user_settings, user_id)
            settings_fetched = True
        except Exception as error:
            logger.warning("User settings lookup failed; memory is disabled for this turn (%s).", type(error).__name__)
            user_settings = {}
            settings_loaded = False
    elif not personalization_disabled:
        try:
            # Use the short account-settings cache on canned-answer candidates.
            # This lets saved response preferences disable canned replies without
            # adding a Drive/database read to every routine fast response.
            user_settings = await run_in_threadpool(_cached_user_settings, user_id)
            settings_fetched = True
        except Exception as error:
            logger.warning("Personalization lookup failed; fast response remains available (%s).", type(error).__name__)
            user_settings = {}
    if settings_fetched and not personalization_disabled:
        _cache_user_settings(user_id, user_settings)
    personalization = user_settings.get("personalization") if isinstance(user_settings, dict) else {}
    personal_instructions = ""
    if not personalization_disabled and isinstance(personalization, dict) and personalization.get("enabled") is not False:
        personal_instructions = str(personalization.get("instructions") or "").strip()[:3000]
    if personal_instructions and fast_reply is not None:
        # A canned answer cannot honor account-wide response preferences.
        fast_reply = None
        settings_loaded = settings_fetched
    server_memory = normalize_memory_settings(user_settings.get("memory"))
    memory_enabled = (
        settings_loaded
        and server_memory["enabled"]
        and client_memory_enabled is not False
        and not memory_user_disabled(user_id)
    )
    memory_auto_extract = (
        memory_enabled
        and server_memory["auto_extract"]
        and client_memory_auto_extract is not False
        and not memory_user_opted_out(user_id)
    )
    privacy_settings = user_settings.get("privacy", {}) if isinstance(user_settings, dict) else {}
    request_logging_enabled = (
        not ODDI_BROWSER_LOCAL_CHATS
        and not incognito_mode
        and fast_reply is None
        and privacy_settings.get("request_logging", True) is not False
    )

    # Build the authoritative identity from the authenticated session.
    # The message itself must never be allowed to determine host/admin status.
    identity = create_identity(
        user_id=str(user_id),
        email=str(request.session.get("email", "")),
    )

    # If the browser lost the successful response, its Retry action repeats
    # this assistant message ID. Return the already-saved reply instead of
    # generating a second answer for the same turn.
    if conversation_id is not None and not ODDI_DRIVE_CHAT_STORAGE:
        try:
            existing_conversation = await run_in_threadpool(
                get_conversation, conversation_id, user_id
            )
        except DriveStorageError:
            logger.warning("Google Drive chat lookup failed; generating without server-side deduplication.")
            existing_conversation = None
        except Exception:
            logger.exception(
                "Conversation lookup failed request_id=%s conversation=%s; continuing without server-side deduplication",
                chat_request_id,
                conversation_id,
            )
            existing_conversation = None
        existing_reply = next(
            (
                item for item in (existing_conversation or {}).get("messages", [])
                if str(item.get("id") or "") == assistant_message_id
                and item.get("role") == "assistant"
            ),
            None,
        )
        if existing_reply is not None:
            response = response_from_engine(str(existing_reply.get("text") or existing_reply.get("content") or ""))
            response.headers["X-Oddi-Assistant-Message-Id"] = assistant_message_id
            response.headers["X-Oddi-Completed-At"] = str(existing_reply.get("completed_at") or datetime.now(timezone.utc).isoformat())
            response.headers["X-Oddi-Conversation-Persisted"] = "1"
            response.headers["X-Oddi-Conversation-Revision"] = str((existing_conversation or {}).get("revision") or 0)
            response_meta = existing_reply.get("response_meta") if isinstance(existing_reply.get("response_meta"), dict) else {}
            response.headers["X-Oddi-Provider"] = str(response_meta.get("provider") or "Auto routing")
            response.headers["X-Oddi-Latency-Ms"] = str(response_meta.get("latency_ms") or 0)
            response.headers["X-Oddi-Token-Estimate"] = str(response_meta.get("token_estimate") or 0)
            existing_sources = existing_reply.get("sources") if isinstance(existing_reply.get("sources"), list) else []
            if existing_sources:
                response.headers["X-Oddi-Sources"] = json.dumps(existing_sources[:48], separators=(",", ":"))
            return response

    # Plan Library persistence before generation, but commit only after the
    # AI successfully processes the uploads. This keeps replacements reversible.
    file_record_ids = []
    upload_fingerprint = None
    is_upload_retry = False
    upload_library_plan = {
        "uploads": [],
        "replace_file_ids": [],
        "replace_storage_keys": [],
        "quota_credit_bytes": 0,
        "already_exists_count": 0,
    }

    if engine_files and persist_library_uploads:
        file_payloads = []

        try:
            library_upload_payloads = []
            for upload_index, uploaded_file in enumerate(engine_files):
                uploaded_file.stream.seek(0)
                file_bytes = uploaded_file.stream.read()
                uploaded_file.stream.seek(0)

                payload = (uploaded_file, file_bytes)
                file_payloads.append(payload)
                if upload_index not in library_source_indexes:
                    library_upload_payloads.append(payload)

            upload_fingerprint = _uploaded_batch_fingerprint(file_payloads)
            previous_fingerprint = _get_chat_upload_retry_fingerprint(
                user_id,
                generation_id,
            )
            is_upload_retry = previous_fingerprint is not None
            if is_upload_retry and previous_fingerprint != upload_fingerprint:
                return JSONResponse(
                    {"error": "This retry does not match the original attachments. Start a new message to upload different files."},
                    status_code=409,
                )

            if not is_upload_retry:
                upload_library_plan = _plan_chat_library_uploads(user_id, library_upload_payloads)
                current_file_count = count_files(user_id)
                replacement_count = len(upload_library_plan["replace_file_ids"])
                planned_count = len(upload_library_plan["uploads"])
                projected_count = current_file_count - replacement_count + planned_count
                if projected_count > USER_FILE_LIMIT:
                    return JSONResponse(
                        {
                            "error": "Your Library is full (30 files). Delete an older file from Library before uploading more.",
                            "file_limit": USER_FILE_LIMIT,
                            "file_count": current_file_count,
                        },
                        status_code=413,
                    )

                # Validate the projected size with credit for replaced objects.
                # Exact duplicates contribute zero bytes and do not consume a
                # Library slot.
                planned_bytes = sum(
                    len(entry["file_bytes"])
                    for entry in upload_library_plan["uploads"]
                )
                if planned_bytes:
                    quota = storage_router.get_file_quota(user_id)
                    remaining_bytes = int(quota.get("remaining_bytes", 0) or 0)
                    quota_credit_bytes = int(upload_library_plan["quota_credit_bytes"] or 0)
                else:
                    quota = None
                    remaining_bytes = 0
                    quota_credit_bytes = 0

                if planned_bytes > remaining_bytes + quota_credit_bytes:
                    return JSONResponse(
                        {
                            "error": "File storage quota exceeded.",
                            "storage": quota,
                        },
                        status_code=413,
                    )

            for uploaded_file, _ in file_payloads:
                uploaded_file.stream.seek(0)

        except FileLimitExceeded as file_limit_error:
            logger.info(
                "File count limit reached for user %s: %s",
                user_id,
                file_limit_error,
            )
            return JSONResponse(
                {
                    "error": str(file_limit_error),
                    "file_limit": USER_FILE_LIMIT,
                    "file_count": count_files(user_id),
                },
                status_code=413,
            )

        except StorageQuotaExceeded as quota_error:
            logger.warning(
                "File storage quota exceeded for user %s: %s",
                user_id,
                quota_error,
            )
            return JSONResponse(
                {
                    "error": "File storage quota exceeded.",
                    "storage": storage_router.get_file_quota(user_id),
                },
                status_code=413,
            )

        except Exception as storage_error:
            logger.exception(
                "File storage failed for user %s: %s",
                user_id,
                storage_error,
            )
            return JSONResponse(
                {"error": "Uploaded file could not be stored."},
                status_code=500,
            )

    history_json = str(form.get("history", "[]"))

    try:
        conversation_history = json.loads(history_json)
    except (json.JSONDecodeError, TypeError):
        conversation_history = []

    conversation_history = conversation_history[-100:]

    if request_logging_enabled:
        if engine_files:
            logger.info("Files received: %s", [f.filename for f in engine_files])
        else:
            logger.info("No files uploaded")

    cancel_event = threading.Event()
    with _chat_generation_lock:
        _active_chat_generations[generation_id] = {
            "user_id": str(user_id),
            "cancelled": cancel_event,
        }

    try:
        # Match the original text the user typed before the frontend's
        # Library-recall instructions or other internal context can alter it.
        # Requests with attachments continue through the document-aware engine.
        if fast_reply is not None:
            reply = fast_reply
        else:
            engine_message = message
            style_instruction = response_style_instructions.get(response_style)
            if style_instruction:
                engine_message += f"\n\n[Response format guidance: {style_instruction}]"
            reply = await run_in_threadpool(
                process_message,
                engine_message,
                engine_files,
                conversation_history,
                user_id=user_id,
                identity=identity,
                long_term_memory_enabled=memory_enabled,
                specialist_agent=specialist_agent,
                personal_instructions=personal_instructions,
                external_document_context=external_document_context,
                external_document_sources=external_document_sources,
            )
        if cancel_event.is_set():
            return PlainTextResponse("Generation stopped.", status_code=409)
    except Exception:
        logger.exception(
            "Chat processing failed request_id=%s user=%s; uploaded Library files were left unchanged.",
            chat_request_id,
            user_id,
        )
        return JSONResponse(
            {
                "error": "The chat service could not complete this reply. Please retry.",
                "request_id": chat_request_id,
            },
            status_code=503,
        )
    finally:
        with _chat_generation_lock:
            _active_chat_generations.pop(generation_id, None)

    if str(reply or "").strip() == str(PROVIDER_UNAVAILABLE_RESPONSE).strip():
        logger.warning(
            "No configured provider completed chat request_id=%s user=%s",
            chat_request_id,
            user_id,
        )
        return JSONResponse(
            {
                "error": "ODDI-AI providers are unavailable right now. Please try again in a moment.",
                "code": "providers_unavailable",
                "request_id": chat_request_id,
            },
            status_code=503,
            headers={"Cache-Control": "no-store"},
        )

    if upload_library_plan.get("uploads"):
        try:
            file_record_ids = await run_in_threadpool(
                _persist_chat_library_uploads,
                user_id,
                conversation_id,
                upload_library_plan,
            )
        except FileLimitExceeded as file_limit_error:
            logger.info(
                "File count changed before Library replacement for user %s: %s",
                user_id,
                file_limit_error,
            )
            return JSONResponse(
                {
                    "error": str(file_limit_error),
                    "file_limit": USER_FILE_LIMIT,
                    "file_count": count_files(user_id),
                },
                status_code=413,
            )
        except StorageQuotaExceeded as quota_error:
            logger.warning(
                "File storage quota changed before Library replacement for user %s: %s",
                user_id,
                quota_error,
            )
            return JSONResponse(
                {
                    "error": "File storage quota exceeded.",
                    "storage": storage_router.get_file_quota(user_id),
                },
                status_code=413,
            )
        except Exception as storage_error:
            logger.exception(
                "Could not update Library files for user %s: %s",
                user_id,
                storage_error,
            )
            return JSONResponse(
                {"error": "Uploaded file could not be stored."},
                status_code=500,
            )

    if engine_files and persist_library_uploads:
        _remember_chat_upload_retry_fingerprint(
            user_id,
            generation_id,
            upload_fingerprint,
        )

    # Keep the exact source set with the assistant message so retry and chat
    # history can render the same citations after a reload.
    library_file_mappings = []
    planned_file_ids_by_upload_index = {}
    for index, (entry, file_record_id) in enumerate(
        zip(upload_library_plan.get("uploads") or [], file_record_ids)
    ):
        upload_index = int(entry.get("upload_index", index))
        file_id = int(file_record_id)
        planned_file_ids_by_upload_index[upload_index] = file_id
        library_file_mappings.append([upload_index, file_id])
    for entry in upload_library_plan.get("duplicate_uploads") or []:
        upload_index = int(entry.get("upload_index", -1))
        file_id = planned_file_ids_by_upload_index.get(
            int(entry.get("planned_upload_index", -1))
        )
        if upload_index >= 0 and file_id is not None:
            library_file_mappings.append([upload_index, file_id])
    library_file_mappings.extend(
        [int(entry.get("upload_index", 0)), int(entry["file_id"])]
        for entry in upload_library_plan.get("already_existing_files") or []
        if entry.get("file_id") is not None
    )
    library_file_ids_by_index = {
        int(upload_index): int(file_id)
        for upload_index, file_id in library_file_mappings
    }
    direct_engine_files = [
        uploaded_file for upload_index, uploaded_file in enumerate(engine_files)
        if upload_index not in library_source_indexes
    ]
    capability_sources = getattr(reply, "sources", []) or []
    has_file_search_citations = any(
        isinstance(source, dict) and source.get("type") in {"file", "file_reference"}
        for source in capability_sources
    )
    if file_record_ids or library_file_ids_by_index or has_file_search_citations:
        try:
            all_library_records = get_files(user_id)
            verified_library_records.update({
                str(record.get("id")): record
                for record in all_library_records
                if record.get("id") is not None
            })
        except Exception as source_lookup_error:
            logger.warning("Could not load Library filenames for citations for user %s: %s", user_id, type(source_lookup_error).__name__)

    response_sources = []
    seen_source_keys = set()

    def add_response_source(source):
        if not isinstance(source, dict):
            return
        if source.get("type") == "web":
            url = str(source.get("url") or "").strip()
            if not url.startswith(("https://", "http://")):
                return
            key = ("web", url.rstrip("/").casefold())
            normalized = {"type": "web", "title": str(source.get("title") or url)[:240], "url": url[:1800]}
        elif source.get("type") == "file":
            file_id = str(source.get("id") or "").strip()
            filename = str(source.get("name") or source.get("filename") or "").strip()
            if not file_id or not filename:
                return
            key = ("file", file_id)
            normalized = {"type": "file", "storage": "server", "id": file_id, "name": filename[:240]}
        else:
            return
        if key in seen_source_keys:
            return
        seen_source_keys.add(key)
        if len(response_sources) < 48:
            response_sources.append(normalized)

    for file_id in requested_source_file_ids:
        record = verified_library_records.get(file_id)
        if record:
            add_response_source({
                "type": "file",
                "id": file_id,
                "name": record.get("filename") or record.get("name") or "Library file",
            })

    for upload_index, file_id in library_file_ids_by_index.items():
        record = verified_library_records.get(str(file_id))
        uploaded_name = ""
        if 0 <= upload_index < len(direct_engine_files):
            uploaded_name = str(direct_engine_files[upload_index].filename or "")
        add_response_source({
            "type": "file",
            "id": str(file_id),
            "name": (record or {}).get("filename") or (record or {}).get("name") or uploaded_name or "Library file",
        })

    for source in capability_sources:
        if isinstance(source, dict) and source.get("type") == "file_reference":
            cited_name = str(source.get("name") or "").strip().casefold()
            matched_record = next((
                record for record in verified_library_records.values()
                if str(record.get("filename") or record.get("name") or "").strip().casefold() == cited_name
            ), None)
            if matched_record:
                add_response_source({
                    "type": "file",
                    "id": str(matched_record.get("id")),
                    "name": matched_record.get("filename") or matched_record.get("name") or source.get("name"),
                })
            continue
        add_response_source(source)

    completed_at = datetime.now(timezone.utc).isoformat()
    latency_ms = max(0, int((time.perf_counter() - generation_started) * 1000))
    reply_text = str(reply)
    token_estimate = max(1, (len(reply_text) + 3) // 4) if isinstance(reply, str) else 0
    assistant_saved_to_conversation = False
    saved_conversation_revision = 0
    if conversation_id is not None and isinstance(reply, str):
        assistant_message = {
            "id": assistant_message_id,
            "role": "assistant",
            "text": reply_text,
            "sources": response_sources,
            "pinned": False,
            "feedback": None,
            "background_completion": True,
            "completed_at": completed_at,
            "response_meta": {
                "latency_ms": latency_ms,
                "token_estimate": token_estimate,
                "provider": "Auto routing",
            },
        }
        if document_workspace_requested:
            assistant_message["document_workspace"] = True
            if document_title:
                assistant_message["document_title"] = document_title
        try:
            # Persist before sending the completed answer to the browser. A
            # background-only Drive write allowed a reload to win this race.
            saved_conversation = await run_in_threadpool(
                append_conversation_message,
                conversation_id,
                user_id,
                assistant_message,
            )
        except DriveStorageError:
            logger.warning(
                "Google Drive save failed before reply delivery; keeping local recovery and scheduling a retry."
            )
            saved_conversation = None
        except Exception:
            logger.exception(
                "Completed reply persistence failed request_id=%s conversation=%s user=%s; returning reply for local recovery",
                chat_request_id,
                conversation_id,
                user_id,
            )
            saved_conversation = None

        if not saved_conversation:
            logger.warning(
                "Completed chat response could not be saved conversation=%s user=%s",
                conversation_id,
                user_id,
            )
            if ODDI_DRIVE_CHAT_STORAGE:
                background_tasks.add_task(
                    _persist_completed_chat_message,
                    conversation_id,
                    user_id,
                    assistant_message,
                )
        else:
            assistant_saved_to_conversation = True
            saved_conversation_revision = int(saved_conversation.get("revision", 0) or 0)

    if (
        fast_reply is not None
        and client_memory_enabled is not False
        and client_memory_auto_extract is not False
    ):
        background_tasks.add_task(
            _enqueue_memory_message_if_enabled,
            str(user_id),
            assistant_message_id,
            memory_message,
        )
    elif (
        not ODDI_BROWSER_LOCAL_CHATS
        and memory_provider_keys_configured()
        and memory_enabled
        and memory_auto_extract
    ):
        background_tasks.add_task(
            enqueue_memory_message,
            str(user_id),
            assistant_message_id,
            memory_message if is_safe_memory_message(memory_message) else "",
        )

    if file_record_ids:
        for file_record_id in file_record_ids:
            try:
                update_file_record(
                    file_record_id,
                    user_id,
                    status="processed",
                )
            except Exception as file_db_error:
                logger.warning(
                    "File metadata status update failed: %s",
                    file_db_error,
                )

    response = response_from_engine(reply)
    response.headers["X-Oddi-Assistant-Message-Id"] = assistant_message_id
    response.headers["X-Oddi-Completed-At"] = completed_at
    response.headers["X-Oddi-Provider"] = "Auto routing"
    response.headers["X-Oddi-Latency-Ms"] = str(latency_ms)
    response.headers["X-Oddi-Token-Estimate"] = str(token_estimate)
    response.headers["X-Oddi-Conversation-Persisted"] = "1" if assistant_saved_to_conversation else "0"
    if assistant_saved_to_conversation:
        response.headers["X-Oddi-Conversation-Revision"] = str(saved_conversation_revision)
    if response_sources:
        response.headers["X-Oddi-Sources"] = json.dumps(response_sources, separators=(",", ":"))
    planned_count = len(upload_library_plan.get("uploads") or [])
    replaced_count = min(planned_count, len(upload_library_plan.get("replace_file_ids") or []))
    response.headers["X-Oddi-Library-Added-Count"] = str(max(0, planned_count - replaced_count))
    response.headers["X-Oddi-Library-Updated-Count"] = str(replaced_count)
    response.headers["X-Oddi-Library-Existing-Count"] = str(int(upload_library_plan.get("already_exists_count") or 0))
    if library_file_mappings:
        response.headers["X-Oddi-Library-File-Indexes"] = json.dumps(
            library_file_mappings,
            separators=(",", ":"),
        )
    return response


# =========================================================
# DEVELOPMENT ENTRYPOINT
# =========================================================

if __name__ == "__main__":
    import socket
    import uvicorn

    host = "127.0.0.1" if ODDI_LOCAL_ONLY_STORAGE else "0.0.0.0"
    port = int(os.getenv("PORT", "8000"))

    # Avoid starting a second local copy and turning a normal duplicate launch
    # into Uvicorn's noisy WinError 10048 bind failure.
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as port_probe:
        port_probe.settimeout(0.35)
        port_is_live = port_probe.connect_ex(("127.0.0.1", port)) == 0
    if port_is_live:
        print(
            f"A server is already accepting connections at http://127.0.0.1:{port}. "
            "Use the running ODDI instance, or stop the process holding this port before restarting."
        )
        raise SystemExit(0)

    uvicorn.run(
        "server:app",
        host=host,
        port=port,
        reload=not ODDI_LOCAL_ONLY_STORAGE,
    )
