from dataclasses import replace
import io

import pytest

from app.swarm import (
    AGENT_REGISTRY,
    AgentResult,
    analyze_data_files,
    auto_route,
    create_execution_plan,
    create_presentation_pptx,
    get_agent_registry,
    run_task,
    validate_agent_ids,
    _central_synthesis,
    _research_focus_query,
    _research_handler,
)


def test_registry_has_seven_unique_specialists_with_real_handlers():
    assert len(AGENT_REGISTRY) == 7
    assert len({agent.id for agent in AGENT_REGISTRY.values()}) == 7
    assert "image" not in AGENT_REGISTRY
    assert "review" not in AGENT_REGISTRY
    assert "presentation" not in AGENT_REGISTRY
    assert all(callable(agent.handler) for agent in AGENT_REGISTRY.values())
    assert len(get_agent_registry()) == 7


def test_auto_router_selects_relevant_agents_without_image_generation_agent():
    assert auto_route("Research current sources and write a report") == ["search", "research", "writing"]
    assert "image" not in auto_route("Generate an image of a mountain")
    assert "presentation" not in auto_route("Make a PowerPoint deck", output_format="pptx")


def test_manual_agent_validation_does_not_add_hidden_specialists():
    assert validate_agent_ids(["coding", "coding"]) == ["coding"]
    with pytest.raises(ValueError):
        validate_agent_ids(["coding", "review"])
    with pytest.raises(ValueError):
        validate_agent_ids(["coding", "unknown"])
    with pytest.raises(ValueError):
        validate_agent_ids(list(AGENT_REGISTRY)[:6])


def test_selected_agent_chat_accepts_at_most_three_specialists():
    selected = ["search", "research", "writing"]
    assert validate_agent_ids(selected, max_count=3) == selected
    with pytest.raises(ValueError, match="no more than 3"):
        validate_agent_ids([*selected, "coding"], max_count=3)


def test_dependency_plan_orders_research_and_writing():
    plan = {row["agent_id"]: row for row in create_execution_plan(["search", "research", "writing"])}
    assert plan["research"]["depends_on"] == ["agent:search"]
    assert plan["writing"]["depends_on"] == ["agent:search", "agent:research"]
    assert plan["central"]["depends_on"] == ["agent:search", "agent:research", "agent:writing"]


def test_research_agent_searches_task_when_selected_alone(monkeypatch):
    task = "Does Tesla have a flying car?"
    searched = []
    source = {
        "type": "web",
        "title": "Tesla Roadster",
        "url": "https://www.tesla.com/roadster",
        "snippet": "Roadster specifications",
        "search_provider": "Bing",
    }
    observed = {}

    def search(query, **kwargs):
        searched.append(query)
        return [dict(source)] if query == "Tesla flying car?" else []

    def provider(ctx, instructions, **kwargs):
        observed.update(ctx=ctx, instructions=instructions, sources=kwargs.get("sources"))
        return AgentResult(status="completed", output="Source-grounded response.")

    monkeypatch.setattr("app.swarm._plan_research_queries", lambda ctx: [])
    monkeypatch.setattr("app.swarm.search_research_sources", search)
    monkeypatch.setattr("app.swarm._fetch_research_candidates", lambda sources: ([{**source, "retrieval_status": "page_text", "page_text": "Verified source text."} for source in sources], 1))
    monkeypatch.setattr("app.swarm._run_provider_agent", provider)

    result = _research_handler({
        "task": task,
        "files": [],
        "user_id": "test-user",
        "agent": AGENT_REGISTRY["research"],
        "identity": None,
        "memory_enabled": False,
        "web_search": False,
        "upstream": [],
    })

    assert searched == ["Tesla flying car?"]
    assert result.status == "completed"
    assert result.sources[0]["url"] == source["url"]
    assert observed["sources"][0]["url"] == source["url"]


def test_research_agent_does_not_turn_snippets_into_unverified_conclusions(monkeypatch):
    source = {
        "type": "web",
        "title": "Tesla unveils what it says is a flying car - CNN",
        "url": "https://news.google.com/rss/articles/example",
        "snippet": "Recent article headline about a flying Roadster demonstration.",
        "published_at": "Thu, 15 Oct 2026 23:42:00 GMT",
        "search_provider": "Google News",
    }
    monkeypatch.setattr("app.swarm._plan_research_queries", lambda ctx: [])
    monkeypatch.setattr("app.swarm.search_research_sources", lambda query, **kwargs: [dict(source)])
    monkeypatch.setattr("app.swarm._fetch_research_candidates", lambda sources: ([{**source, "retrieval_status": "snippet_only"}], 0))
    monkeypatch.setattr(
        "app.swarm._run_provider_agent",
        lambda *args, **kwargs: pytest.fail("snippet-only evidence must not be presented as a verified model report"),
    )

    result = _research_handler({
        "task": "Did Tesla demonstrate a flying Roadster?",
        "files": [],
        "user_id": "test-user",
        "agent": AGENT_REGISTRY["research"],
        "identity": None,
        "memory_enabled": False,
        "web_search": False,
        "upstream": [],
    })

    assert result.status == "completed"
    assert result.data["evidence_level"] == "search_snippets_only"
    assert result.data["pages_read"] == 0
    assert "cannot independently verify" in result.output
    assert "[Tesla unveils what it says is a flying car - CNN]" in result.output


def test_central_synthesis_preserves_snippet_only_limit_instead_of_guessing(monkeypatch):
    monkeypatch.setattr("app.swarm._call_engine", lambda *args, **kwargs: pytest.fail("unverified snippet-only research must not be resynthesized"))
    report = {
        "id": "research",
        "name": "Research Agent",
        "status": "completed",
        "data": {"evidence_level": "search_snippets_only"},
        "sources": [{
            "title": "Tesla unveils what it says is a flying car - CNN",
            "url": "https://news.google.com/rss/articles/example",
            "published_at": "Thu, 15 Oct 2026 23:42:00 GMT",
            "retrieval_status": "snippet_only",
        }],
    }

    answer = _central_synthesis("Did Tesla demonstrate a flying Roadster?", [report], "test-user", None, False, "answer")

    assert "cannot independently verify" in answer
    assert "[Tesla unveils what it says is a flying car - CNN]" in answer


def test_research_query_focus_keeps_topic_and_named_entities_without_report_instructions():
    task = (
        "Deep research: Do flying cars exist today? Distinguish actual demonstrated vehicles from concepts, "
        "and verify whether Tesla has demonstrated a flying Roadster. Cite evidence and limitations."
    )

    assert _research_focus_query(task) == "flying cars today? Tesla Roadster"


def test_executor_appends_source_links_when_synthesis_omits_them(monkeypatch):
    source = {"title": "Tesla Roadster", "url": "https://www.tesla.com/roadster"}
    monkeypatch.setitem(
        AGENT_REGISTRY,
        "research",
        replace(AGENT_REGISTRY["research"], handler=lambda ctx: AgentResult(
            status="completed", output="Plain answer without citations.", sources=[source]
        )),
    )
    monkeypatch.setattr("app.swarm._central_synthesis", lambda *args: "Plain answer without citations.")

    outcome = run_task("Research flying cars", ["research"], [], "test-user", None, False)

    assert outcome["status"] == "completed"
    assert "## Sources" in outcome["result"]
    assert "[Tesla Roadster](https://www.tesla.com/roadster)" in outcome["result"]


def test_data_agent_computes_bounded_summary_from_real_csv_values():
    data = b"name,score,group\nA,10,x\nB,20,y\nB,20,y\nC,,x\n"
    result = analyze_data_files([("scores.csv", "text/csv", data)])[0]
    assert result["rows_analyzed"] == 4
    assert result["duplicate_rows_detected"] == 1
    assert result["columns"][1] == {"name": "score", "type": "numeric", "missing": 1, "nonempty": 3}
    score = result["numeric_summaries"][0]
    assert (score["count"], score["sum"], score["mean"], score["median"], score["min"], score["max"]) == (3, 50.0, 16.666667, 20.0, 10.0, 20.0)
    assert (score["range"], score["q1"], score["q3"]) == (10.0, 15.0, 20.0)
    group = next(item for item in result["grouped_summaries"] if item["group_by"] == "group")
    y_group = next(item for item in group["groups"] if item["value"] == "y")
    assert y_group["measures"][0]["sum"] == 40.0
    assert group["top_by_sum"][0]["groups"][0]["value"] == "y"


def test_data_agent_rejects_empty_table():
    with pytest.raises(ValueError, match="empty"):
        analyze_data_files([("empty.csv", "text/csv", b"")])


def test_presentation_builder_returns_a_valid_editable_pptx():
    blob = create_presentation_pptx({"title": "Market Review", "slides": [{"title": "Evidence", "bullets": ["Source-backed finding"], "speaker_notes": "Caveat"}]})
    from pptx import Presentation
    deck = Presentation(io.BytesIO(blob))
    assert len(deck.slides) == 2
    assert deck.slides[1].shapes.title.text == "Evidence"


def test_data_agent_answers_from_pasted_table_and_includes_calculated_values(monkeypatch):
    from app.swarm import _data_handler

    task = (
        "Analysis question: Which region had the highest sales total?\n\n"
        "ODDI_DATA_TABLE_BEGIN\nRegion,Sales\nEast,5\nWest,7\nWest,3\nEast,4\nODDI_DATA_TABLE_END"
    )
    seen = {}

    def fake_provider(ctx, instructions, **kwargs):
        seen.update(ctx)
        return AgentResult(status="completed", output="West had the highest sales total.")

    monkeypatch.setattr("app.swarm._run_provider_agent", fake_provider)
    result = _data_handler({
        "task": task,
        "files": [],
        "user_id": "test-user",
        "agent": AGENT_REGISTRY["data"],
        "identity": None,
        "memory_enabled": False,
        "upstream": [],
    })

    assert result.status == "completed"
    assert "Highest Sales total by Region: West = 10.0" in result.output
    assert result.data["datasets"][0]["filename"] == "Pasted dataset.csv"
    assert "ODDI_DATA_TABLE_END" not in seen["task"]
    assert "DATASET SUMMARY COMPUTED LOCALLY" in seen["upstream"][-1]["output"]


def test_executor_passes_dependencies_and_synthesizes(monkeypatch):
    seen = []
    def handler(agent_id):
        def execute(ctx):
            seen.append(agent_id)
            upstream = {row["id"] for row in ctx["upstream"]}
            if agent_id == "research":
                assert "search" in upstream
            if agent_id == "writing":
                assert {"search", "research"}.issubset(upstream)
            return AgentResult(status="completed", output=f"{agent_id} findings", sources=[{"url": "https://example.test/source", "title": "Source"}] if agent_id == "search" else [])
        return execute
    for agent_id in ("search", "research", "writing"):
        monkeypatch.setitem(AGENT_REGISTRY, agent_id, replace(AGENT_REGISTRY[agent_id], handler=handler(agent_id)))
    monkeypatch.setattr("app.swarm._central_synthesis", lambda *args: "Unified answer")
    outcome = run_task("Research and write a report", ["search", "research", "writing"], [], "test-user", None, False, max_concurrency=3)
    assert outcome["status"] == "completed"
    assert outcome["result"].startswith("Unified answer")
    assert "## Sources" in outcome["result"]
    assert "https://example.test/source" in outcome["result"]
    assert set(seen) == {"search", "research", "writing"}
    assert all(row["status"] == "completed" for row in outcome["reports"])


def test_executor_returns_specialist_answer_when_central_synthesis_fails(monkeypatch):
    for agent_id, output in (("coding", "print('ready')"), ("writing", "The answer is ready.")):
        monkeypatch.setitem(
            AGENT_REGISTRY,
            agent_id,
            replace(AGENT_REGISTRY[agent_id], handler=lambda ctx, value=output: AgentResult(status="completed", output=value)),
        )
    monkeypatch.setattr("app.swarm._central_synthesis", lambda *args: "")

    outcome = run_task("Give me a short answer", ["coding", "writing"], [], "test-user", None, False)

    assert outcome["status"] == "completed"
    assert "Central Intelligence could not combine" in outcome["result"]
    assert "Coding Agent" in outcome["result"]
    assert "print('ready')" in outcome["result"]
    assert "Writing Agent" in outcome["result"]
    assert "The answer is ready." in outcome["result"]
    assert next(row for row in outcome["plan"] if row["agent_id"] == "central")["status"] == "completed"


def test_executor_preserves_single_specialist_response_without_synthesis(monkeypatch):
    monkeypatch.setitem(
        AGENT_REGISTRY,
        "writing",
        replace(AGENT_REGISTRY["writing"], handler=lambda ctx: AgentResult(status="completed", output="Direct specialist answer.")),
    )
    monkeypatch.setattr("app.swarm._central_synthesis", lambda *args: "")

    outcome = run_task("Answer briefly", ["writing"], [], "test-user", None, False)

    assert outcome["status"] == "completed"
    assert outcome["result"] == "Direct specialist answer."


def test_executor_retries_a_retryable_failure_once(monkeypatch):
    attempts = {"writing": 0}
    def fail_then_succeed(ctx):
        attempts["writing"] += 1
        if attempts["writing"] == 1:
            return AgentResult(status="failed", error="temporary provider error", retryable=True)
        return AgentResult(status="completed", output="Draft")
    monkeypatch.setitem(AGENT_REGISTRY, "writing", replace(AGENT_REGISTRY["writing"], handler=fail_then_succeed))
    monkeypatch.setattr("app.swarm._central_synthesis", lambda *args: "Draft result")
    outcome = run_task("Write a short report", ["writing"], [], "test-user", None, False)
    assert attempts["writing"] == 2
    assert outcome["reports"][0]["attempts"] == 2
    assert outcome["status"] == "completed"


def test_executor_cancels_before_start(monkeypatch):
    import threading
    called = []
    monkeypatch.setitem(AGENT_REGISTRY, "writing", replace(AGENT_REGISTRY["writing"], handler=lambda ctx: called.append(True) or AgentResult(status="completed", output="unexpected")))
    event = threading.Event(); event.set()
    outcome = run_task("Write", ["writing"], [], "test-user", None, False, cancel_event=event)
    assert outcome["status"] == "cancelled"
    assert called == []
